"""Excel template preview and auditable exports, independent of storage and HTTP.

Templates provide design metadata only. No example row becomes a paper, evidence,
or final value. Exports use the current item IDs/order and never evaluate input as
an Excel formula.
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timezone
from io import BytesIO
import json
import math
import re
from urllib.parse import urlsplit
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from .locations import source_location as _verification_location

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .keywords import suggest_keywords


DATA_TYPES = {"text", "number", "integer", "boolean", "date", "narrative"}
STATUS_LABELS = {
    "unreviewed": "Not yet reviewed",
    "approved": "Approved value",
    "not_reported": "Not reported",
    "not_applicable": "Not applicable",
    "unclear": "Unclear",
    "conflicting": "Conflicting evidence",
}
FINAL_STATUSES = {"approved", "not_reported", "not_applicable"}
HEADER_FILL = "FFDCE6F1"  # Exact pale blue from the supplied reference workbook.
MAX_CELL_LENGTH = 32767


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _header_rows(value):
    try:
        rows = sorted({int(v.strip()) for v in str(value or "1").split(",")})
    except ValueError as exc:
        raise ValueError("Header rows must be comma-separated positive row numbers, such as 1,2.") from exc
    if not rows or rows[0] < 1 or rows[-1] > 10000:
        raise ValueError("Header rows must be between 1 and 10000.")
    return rows


def _merged_value(sheet, row, column):
    value = sheet.cell(row, column).value
    if value is not None:
        return value
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return sheet.cell(merged.min_row, merged.min_col).value
    return None


def _list_value(value):
    if isinstance(value, list):
        return value
    value = _text(value).strip()
    if not value:
        return []
    if value.startswith("["):
        try:
            parsed = json.loads(value)
            if isinstance(parsed, list):
                return parsed
        except ValueError:
            pass
    return [part.strip() for part in re.split(r"[;\n,|]", value) if part.strip()]


def _boolean(value):
    if isinstance(value, bool):
        return value
    return _text(value).strip().lower() in {"true", "yes", "1", "required", "y"}


def _item(heading, position, **extra):
    return {
        "name": heading, "export_heading": heading, "definition": "", "instructions": "",
        "data_type": "text", "allowed_values": [], "units": "", "required": False,
        "position": position, "context": {}, "keywords": [], "rules": [], "section": "study", **extra,
    }


_DICTIONARY_ALIASES = {
    "name": {"name", "item", "item name", "field", "field name", "variable", "variable name", "display name"},
    "export_heading": {"export heading", "export column heading", "column heading", "heading", "label"},
    "definition": {"definition", "description"},
    "instructions": {"instructions", "extraction instructions", "guidance"},
    "data_type": {"data type", "type", "datatype"},
    "allowed_values": {"allowed values", "values", "choices", "categories"},
    "units": {"units", "unit"},
    "required": {"required", "mandatory"},
    "keywords": {"keywords", "keyword", "synonyms", "phrases", "keywords synonyms"},
    "rules": {"rules", "matching rules"},
    "context": {"context"},
    "group": {"group", "treatment group"},
    "population": {"population"},
    "outcome": {"outcome"},
    "time_point": {"time point", "timepoint"},
}


def preview_template(content: bytes, sheet: str | None = None, header_rows: str = "1",
                     mode: str = "columns", mapping: dict | None = None, *, all_sheets: bool = False) -> dict:
    """Preview headings or explicitly selected dictionary metadata, without import.

    Dictionary mappings use property names and 1-based Excel column numbers.
    Multiple column header rows join with `` / ``; actual merged cells supply
    their shared parent heading. Returned item metadata remains user-editable.
    """
    if mode not in {"columns", "dictionary"}:
        raise ValueError("Template mode must be columns or dictionary.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(member.file_size for member in archive.infolist()) > 100_000_000:
                raise ValueError("Workbook expands beyond the 100 MB template limit.")
        book = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    except (BadZipFile, OSError, KeyError, ParseError) as exc:
        raise ValueError("The template must be a valid .xlsx workbook.") from exc
    try:
        rows = _header_rows(header_rows)
        selected_names = book.sheetnames if all_sheets else [sheet or book.sheetnames[0]]
        previews = []
        for name in selected_names:
            if name not in book.sheetnames:
                raise ValueError("Select a worksheet included in this workbook.")
            try:
                previews.append(_preview_worksheet(book, name, rows, mode, mapping))
            except ValueError as exc:
                raise ValueError(f"Worksheet {name!r}: {exc}") from exc
        if not all_sheets:
            return previews[0]
        items = [item for preview in previews for item in preview['items']]
        for position, item in enumerate(items):
            item['position'] = position
        return dict(sheets=book.sheetnames, sheet=None, all_sheets=True,
                    worksheets=previews, items=items, header_rows=rows, mode=mode,
                    warnings=[f"{preview['sheet']}: {warning}" for preview in previews for warning in preview['warnings']],
                    preview_rows=[], mapping={}, columns=[], example_included=False)
    finally:
        book.close()


def _preview_worksheet(book, selected, rows, mode, mapping):
    ws = book[selected]
    if ws.max_column > 1000 or ws.max_row > 20000:
        raise ValueError("Template preview supports up to 1000 columns and 20000 rows.")
    if rows[-1] > ws.max_row:
        raise ValueError("A selected header row is outside the worksheet.")
    warnings = ["Template example rows are preview only. No study records or evidence will be imported."]
    headers = []
    for column in range(1, ws.max_column + 1):
        parts = []
        for row in rows:
            value = _text(_merged_value(ws, row, column)).strip()
            if value and (not parts or value != parts[-1]):
                parts.append(value)
        headers.append(" / ".join(parts))
    result_items = []
    resolved_mapping = {}
    if mode == "columns":
        for column, heading in enumerate(headers, 1):
            if not heading:
                warnings.append(f"Column {get_column_letter(column)} has no selected heading and was skipped.")
                continue
            item = _item(heading, len(result_items), source_column=column)
            item["keywords"] = suggest_keywords(item)
            result_items.append(item)
    else:
        if mapping is None:
            for column, heading in enumerate(headers, 1):
                normalized = re.sub(r"[_/\-]+", " ", heading).strip().lower()
                for property_name, aliases in _DICTIONARY_ALIASES.items():
                    if normalized in aliases and property_name not in resolved_mapping:
                        resolved_mapping[property_name] = column
        else:
            if not isinstance(mapping, dict):
                raise ValueError("Dictionary mapping must map property names to 1-based column numbers.")
            for key, column in mapping.items():
                if key not in _DICTIONARY_ALIASES:
                    raise ValueError(f"Unknown dictionary property: {key}.")
                if column in (None, ""):
                    continue
                try:
                    number = int(column)
                except (ValueError, TypeError) as exc:
                    raise ValueError(f"The mapping for {key} must be a 1-based column number.") from exc
                if isinstance(column, bool) or str(column).strip() != str(number) or not 1 <= number <= ws.max_column:
                    raise ValueError(f"The mapping for {key} is outside the worksheet columns.")
                resolved_mapping[key] = number
        if "name" not in resolved_mapping and "export_heading" not in resolved_mapping:
            raise ValueError("Map an item name or export heading column for this data dictionary.")
        for row in range(rows[-1] + 1, ws.max_row + 1):
            raw = {key: ws.cell(row, column).value for key, column in resolved_mapping.items()}
            heading = _text(raw.get("export_heading") or raw.get("name")).strip()
            if not heading:
                if any(value is not None for value in raw.values()):
                    warnings.append(f"Dictionary row {row} has no name or heading and was skipped.")
                continue
            item = _item(heading, len(result_items), source_row=row)
            for key in ("name", "export_heading", "definition", "instructions", "units"):
                if raw.get(key) is not None:
                    item[key] = _text(raw[key]).strip()
            if not item["name"]:
                item["name"] = heading
            item["required"] = _boolean(raw.get("required"))
            item["keywords"] = [_text(value) for value in _list_value(raw.get("keywords"))]
            if not item["keywords"] and _text(raw.get("keywords")).strip() != "[]":
                item["keywords"] = suggest_keywords(item)
            item["allowed_values"] = _list_value(raw.get("allowed_values"))
            type_name = _text(raw.get("data_type") or "text").strip().lower()
            type_name = {"numeric": "number", "int": "integer", "bool": "boolean", "string": "text"}.get(type_name, type_name)
            if type_name not in DATA_TYPES:
                warnings.append(f"Dictionary row {row}: unrecognized type {type_name!r}; preview uses text. Correct before import.")
                type_name = "text"
            item["data_type"] = type_name
            if raw.get("context"):
                try:
                    context = json.loads(_text(raw["context"]))
                    if not isinstance(context, dict):
                        raise ValueError("Context must be an object.")
                    item["context"] = context
                except ValueError:
                    warnings.append(f"Dictionary row {row}: context is not a JSON object; retained as a context note.")
                    item["context"] = {"note": _text(raw["context"])}
            for key in ("group", "population", "outcome", "time_point"):
                if raw.get(key) is not None:
                    item["context"][key] = _text(raw[key])
            if raw.get("rules"):
                try:
                    rules = json.loads(_text(raw["rules"]))
                    if not isinstance(rules, list) or any(not isinstance(rule, dict) for rule in rules):
                        raise ValueError("Rules must be an array of objects.")
                    item["rules"] = rules
                except ValueError:
                    warnings.append(f"Dictionary row {row}: invalid matching rules JSON was omitted. Correct before import.")
            result_items.append(item)
    duplicates = [heading for heading, count in Counter(item["export_heading"] for item in result_items).items() if count > 1]
    if duplicates:
        warnings.append("Repeated headings remain separate items. Rename them if needed: " + ", ".join(duplicates))
    if not result_items:
        warnings.append("No extraction items were found for the selected mapping.")
    preview_rows = [[_text(ws.cell(row, column).value) for column in range(1, ws.max_column + 1)]
                    for row in range(rows[-1] + 1, min(ws.max_row, rows[-1] + 5) + 1)]
    answer = {"sheets": book.sheetnames, "sheet": selected, "header_rows": rows, "mode": mode,
              "items": result_items, "preview_rows": preview_rows, "warnings": warnings,
              "mapping": resolved_mapping, "columns": [{"column": i, "heading": h} for i, h in enumerate(headers, 1)],
              "example_included": False}
    for item in result_items:
        item["worksheet_name"] = selected
        item["section"] = "meta_analysis" if re.sub(r"[^a-z]", "", selected.lower()) == "metaanalysis" else "study"
        if mode == "columns":
            column = item["source_column"]
            header, body = ws.cell(rows[-1], column), ws.cell(rows[-1] + 1, column)
            item["template_style"] = _template_style(ws, header, body)
            if item["section"] == "meta_analysis":
                normalized = re.sub(r"[^a-z]", "", item["export_heading"].lower())
                item["data_type"] = "integer" if normalized == "n" else "number" if normalized in {"prescore", "presd", "postscore", "postsd", "r"} else "text"
    return answer


def _template_style(sheet, header, body):
    """Copy only presentational metadata; never template examples or formulas."""
    def cell_style(cell):
        def color(value):
            return value.rgb if value and value.type == "rgb" and isinstance(value.rgb, str) else None
        return {"font": {"name": cell.font.name, "size": cell.font.sz, "bold": cell.font.bold,
                         "italic": cell.font.italic, "color": color(cell.font.color)},
                "fill": color(cell.fill.fgColor) if cell.fill.patternType == "solid" else None,
                "alignment": {"horizontal": cell.alignment.horizontal, "vertical": cell.alignment.vertical,
                              "wrap_text": cell.alignment.wrap_text, "text_rotation": cell.alignment.text_rotation},
                "border": {edge: {"style": getattr(cell.border, edge).style if getattr(cell.border, edge) else None,
                                  "color": color(getattr(cell.border, edge).color) if getattr(cell.border, edge) else None}
                           for edge in ("left", "right", "top", "bottom")},
                "number_format": cell.number_format}
    return {"header": cell_style(header), "body": cell_style(body),
            "column_width": sheet.column_dimensions[header.column_letter].width,
            "header_height": sheet.row_dimensions[header.row].height,
            "body_height": sheet.row_dimensions[body.row].height}


def _is_outcome_item(item):
    return item.get("section", "study") == "meta_analysis"


def _active_outcomes(snapshot):
    paper_order = {paper["id"]: index for index, paper in enumerate(_included_papers(snapshot))}
    return sorted((row for row in snapshot.get("outcomes", []) if row.get("paper_id") in paper_order
                   and not row.get("deleted") and row.get("included", True)),
                  key=lambda row: (paper_order[row["paper_id"]], row.get("position", 0), row.get("id", "")))


def _field_key(record):
    return record["paper_id"], record["item_id"], record.get("outcome_id") or None


def _active_items(snapshot):
    review_id = snapshot.get("review", {}).get("id")
    return sorted((item for item in snapshot.get("items", []) if not item.get("deleted")
                   and item.get("review_id", review_id) == review_id), key=lambda item: item.get("position", 0))


def _included_papers(snapshot):
    review_id = snapshot.get("review", {}).get("id")
    return [paper for paper in snapshot.get("papers", []) if paper.get("included", True)
            and paper.get("review_id", review_id) == review_id]


def _value_error(value, item):
    kind = item.get("data_type", "text")
    if value is None or (isinstance(value, str) and not value.strip()):
        return "An approved value cannot be empty."
    if kind in {"text", "narrative"} and not isinstance(value, str):
        return "The approved text or narrative must be a string."
    if kind == "number" and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)):
        return "The approved number must be a finite numeric value."
    if kind == "integer" and (isinstance(value, bool) or not isinstance(value, int)):
        return "The approved integer must be a whole number."
    if kind == "boolean" and not isinstance(value, bool):
        return "The approved boolean must be true or false."
    if kind == "date":
        try:
            if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                raise ValueError("Date must be YYYY-MM-DD.")
            date.fromisoformat(value)
        except (ValueError, TypeError):
            return "The approved date must use YYYY-MM-DD."
    if kind not in DATA_TYPES:
        return "The extraction item's data type is unknown."
    if item.get("allowed_values") and not any(type(value) is type(option) and value == option for option in item["allowed_values"]):
        return "The approved value is outside the item's allowed values."
    if _is_outcome_item(item) and isinstance(value, (int, float)) and not isinstance(value, bool):
        heading = re.sub(r"[^a-z]", "", (item.get("export_heading") or item.get("name", "")).lower())
        if heading == "n" and (value <= 0 or int(value) != value):
            return "Outcome sample size must be a positive whole number."
        if heading in {"presd", "postsd"} and value < 0:
            return "Standard deviations cannot be negative."
        if heading == "r" and not -1 <= value <= 1:
            return "The pre/post correlation must be between -1 and 1."
    return None


def _scoped_snapshot(snapshot):
    """Defend the export boundary even if a caller passes unrelated records."""
    review_id = snapshot.get("review", {}).get("id")
    papers = [paper for paper in snapshot.get("papers", []) if paper.get("review_id", review_id) == review_id]
    items = [item for item in snapshot.get("items", []) if item.get("review_id", review_id) == review_id]
    paper_ids, item_ids = {paper["id"] for paper in papers}, {item["id"] for item in items}
    scoped = {**snapshot, "papers": papers, "items": items}
    for key in ("fields", "evidence", "suggestions", "feedback", "learning_examples"):
        scoped[key] = [record for record in snapshot.get(key, []) if record.get("review_id", review_id) == review_id and record.get("paper_id") in paper_ids and record.get("item_id") in item_ids]
    scoped["pages"] = [page for page in snapshot.get("pages", []) if page.get("paper_id") in paper_ids]
    scoped["history"] = [record for record in snapshot.get("history", []) if record.get("review_id", review_id) == review_id]
    scoped["outcomes"] = [row for row in snapshot.get("outcomes", []) if row.get("review_id", review_id) == review_id and row.get("paper_id") in paper_ids]
    outcomes = {row["id"]: row for row in scoped["outcomes"]}
    scoped["outcome_fields"] = [record for record in snapshot.get("outcome_fields", [])
                                if record.get("review_id", review_id) == review_id and record.get("paper_id") in paper_ids
                                and record.get("item_id") in item_ids and record.get("outcome_id") in outcomes
                                and outcomes[record["outcome_id"]]["paper_id"] == record["paper_id"]]
    scoped["fields"] = [field for field in scoped["fields"] if not field.get("outcome_id")]
    scoped["evidence"] = [record for record in scoped["evidence"] if not record.get("outcome_id") or
                          (record["outcome_id"] in outcomes and outcomes[record["outcome_id"]]["paper_id"] == record["paper_id"])]
    return scoped


def validate_snapshot(snapshot: dict) -> dict:
    """Validate finalization without changing any stored decisions or values."""
    snapshot = _scoped_snapshot(snapshot)
    items, papers = _active_items(snapshot), _included_papers(snapshot)
    fields = {_field_key(field): field for field in [*snapshot.get("fields", []), *snapshot.get("outcome_fields", [])]}
    evidence = {record["id"]: record for record in snapshot.get("evidence", [])}
    pages = {(page.get("paper_id"), page.get("page_number")): page for page in snapshot.get("pages", [])}
    issues = []
    counts = Counter({status: 0 for status in STATUS_LABELS})
    outcome_id = None

    def issue(severity, code, message, paper=None, item=None):
        entry = {"severity": severity, "code": code, "message": message}
        if paper is not None:
            entry["paper_id"] = paper["id"]
        if item is not None:
            entry["item_id"] = item["id"]
            if outcome_id:
                entry["outcome_id"] = outcome_id
        issues.append(entry)

    if not items:
        issue("error", "no_items", "Add at least one active extraction item before a final export.")
    if not papers:
        issue("error", "no_papers", "Include at least one paper before a final export.")
    if not snapshot.get("review", {}).get("reviewer"):
        issue("warning", "reviewer_missing", "Set the reviewer name to identify the decisions in this review.")
    for item in items:
        if len(_text(item.get("export_heading") or item.get("name"))) > MAX_CELL_LENGTH:
            issue("error", "excel_heading_length", "An export heading exceeds Excel's 32767-character cell limit. Shorten this heading.", item=item)
    for paper in papers:
        if paper.get("processing_status", "ready") != "ready":
            issue("error", "processing_incomplete", "PDF processing is incomplete or failed.", paper)
        if paper.get("search_status") == "partial":
            issue("error", "search_incomplete", "Some searches did not finish. Review the listed pages and rules, then scan this paper again before finalizing.", paper)
        elif paper.get("search_status", "current") != "current":
            issue("error", "search_stale", "Refresh this paper's search using the current extraction rules.", paper)
        unsearchable = [page for page in snapshot.get("pages", []) if page.get("paper_id") == paper["id"] and page.get("status", "ok") != "ok"]
        if unsearchable:
            numbers = ", ".join(str(page.get("page_number")) for page in unsearchable)
            issue("warning" if paper.get("manual_review_ack") else "error", "unsearchable_pages",
                  f"Physical PDF pages {numbers} could not be searched. " +
                  ("Manual review was acknowledged." if paper.get("manual_review_ack") else "Review them manually and acknowledge completion."), paper)
        contexts = [(item, None) for item in items if not _is_outcome_item(item)]
        contexts.extend((item, row["id"]) for row in _active_outcomes(snapshot) if row["paper_id"] == paper["id"]
                        for item in items if _is_outcome_item(item))
        for item, outcome_id in contexts:
            field = fields.get((paper["id"], item["id"], outcome_id), {})
            status = field.get("status", "unreviewed")
            counts[status] += 1
            if status not in STATUS_LABELS:
                issue("error", "invalid_status", "This field has an unknown review status.", paper, item)
            elif status == "unreviewed":
                issue("error" if item.get("required") else "warning", "unreviewed",
                      "Required item is not yet reviewed." if item.get("required") else "Optional item is not yet reviewed; final data cell will be empty.", paper, item)
            elif status in {"unclear", "conflicting"}:
                issue("error", status, "Resolve this field or approve an evidence-supported narrative preserving the uncertainty.", paper, item)
            if status in {"not_reported", "not_applicable"} and not (_text(field.get("rationale")).strip() or _text(field.get("notes")).strip()):
                issue("error", "missing_status_rationale", "Record a rationale or note for this explicit missingness decision.", paper, item)
            if field.get("needs_reconsideration"):
                issue("error", "reconsideration", "Reconsider this field after changes to its design or supporting evidence.", paper, item)
            if status == "approved":
                if isinstance(field.get("value"), str) and len(field["value"]) > MAX_CELL_LENGTH:
                    issue("error", "excel_value_length", "This approved value exceeds Excel's 32767-character cell limit. Split it into additional items before export.", paper, item)
                error = _value_error(field.get("value"), item)
                if error:
                    issue("error", "invalid_value", error, paper, item)
                linked = field.get("evidence_ids", [])
                if not linked:
                    issue("error", "missing_evidence", "Approved values require explicitly linked accepted source evidence.", paper, item)
                used_ids = list(dict.fromkeys([*linked, *(evidence_id for observation in field.get("observations", []) for evidence_id in observation.get("evidence_ids", []))]))
                for evidence_id in used_ids:
                    record = evidence.get(evidence_id)
                    if not record or _field_key(record) != (paper["id"], item["id"], outcome_id) or record.get("decision") != "accepted" or not record.get("quote"):
                        issue("error", "invalid_evidence", "Linked evidence must be accepted and belong to this paper and extraction item.", paper, item)
                        break
                    if record.get("source_stale") or record.get("document_version", 0) != paper.get("document_version", 0):
                        issue("error", "stale_evidence_source", "This evidence predates changes to the extracted document. Re-select its source before finalizing.", paper, item)
                    page = pages.get((paper["id"], record.get("page_number")))
                    if page is None:
                        issue("error", "missing_source_page", "A linked source page is no longer present in the processed document.", paper, item)
                    elif (page.get("status", "ok") == "ok" and record.get("manual") and record.get("start") is None):
                        if not _text(record.get("rationale")).strip() or not _text(record.get("source_location")).strip():
                            issue("error", "image_source_location", "Image transcription requires a rationale and a precise source location.", paper, item)
                    elif page.get("status", "ok") != "ok":
                        if not record.get("manual") or not _text(record.get("rationale")).strip():
                            issue("error", "manual_source_rationale", "Evidence from an unsearchable page requires a labeled manual transcription and rationale.", paper, item)
                    else:
                        start, end = record.get("start"), record.get("end")
                        if (not isinstance(start, int) or not isinstance(end, int) or start < 0 or end <= start or
                                not isinstance(page.get("text"), str) or page["text"][start:end] != record["quote"]):
                            issue("error", "source_mismatch", "Linked evidence no longer matches its exact source text and location. Select the source again.", paper, item)
                if field.get("value_kind", "reported") not in {"reported", "interpretation", "calculated"}:
                    issue("error", "invalid_value_kind", "Choose reported, interpretation, or calculated for this approved value.", paper, item)
                if field.get("value_kind") in {"interpretation", "calculated"} and not _text(field.get("rationale")).strip():
                    issue("error", "missing_rationale", "Interpretations and calculated values need a rationale.", paper, item)
                if field.get("value_kind") == "calculated" and not field.get("calculation_inputs"):
                    issue("error", "missing_calculation_inputs", "Calculated values need their calculation inputs.", paper, item)
                for observation in field.get("observations", []):
                    value = observation.get("value")
                    if value is None or isinstance(value, (dict, list)) or (isinstance(value, float) and not math.isfinite(value)):
                        issue("error", "invalid_observation_value", "Approved observations require a finite scalar value.", paper, item)
                    if not any(_text(observation.get(label)).strip() for label in ("group", "population", "outcome", "time_point")):
                        issue("error", "unlabeled_observation", "Each observation needs a group, population, outcome, or time-point label.", paper, item)
                    if not observation.get("evidence_ids"):
                        issue("error", "missing_observation_evidence", "Each approved observation requires its own linked supporting evidence.", paper, item)
                    for evidence_id in observation.get("evidence_ids", []):
                        record = evidence.get(evidence_id)
                        if not record or _field_key(record) != (paper["id"], item["id"], outcome_id) or record.get("decision") != "accepted":
                            issue("error", "invalid_observation_evidence", "Observation evidence must be accepted and belong to this paper and item.", paper, item)
                            break
    attention_fields = {(entry["paper_id"], entry["item_id"], entry.get("outcome_id")) for entry in issues
                        if entry.get("paper_id") and entry.get("item_id")}
    counts.update({"attention_fields": len(attention_fields),
                   "papers": len(papers), "items": len(items), "fields": sum(counts.values()), "outcomes": len(_active_outcomes(snapshot)),
                   "errors": sum(issue["severity"] == "error" for issue in issues),
                   "warnings": sum(issue["severity"] == "warning" for issue in issues)})
    return {"issues": issues, "can_finalize": not counts["errors"], "counts": dict(counts)}


def _accepted_passages(snapshot):
    """Accepted, current exact sources, without promoting them to approved values."""
    pages = {(p['paper_id'], p['page_number']): p for p in snapshot.get('pages', [])}
    papers = {p['id']: p for p in snapshot.get('papers', [])}
    result = {}
    for record in snapshot.get('evidence', []):
        if record.get('decision') != 'accepted' or not record.get('quote') or record.get('source_stale'):
            continue
        page = pages.get((record['paper_id'], record.get('page_number')))
        paper = papers.get(record['paper_id'], {})
        if page is None or record.get('document_version', 0) != paper.get('document_version', 0):
            continue
        if record.get('manual') and record.get('start') is None:
            valid = bool(_text(record.get('rationale')).strip()) and (page.get('status') != 'ok' or bool(_text(record.get('source_location')).strip()))
        else:
            start, end = record.get('start'), record.get('end')
            valid = (page.get('status', 'ok') == 'ok' and isinstance(start, int) and isinstance(end, int)
                     and 0 <= start < end and page.get('text', '')[start:end] == record['quote'])
        if valid:
            result.setdefault(_field_key(record), []).append(record)
    return result


def _outcome_metadata_value(item, outcome, paper, study_values):
    heading = re.sub(r"[^a-z]", "", (item.get("export_heading") or item.get("name", "")).lower())
    if heading == "study":
        return study_values.get(paper["id"]) or paper.get("citation") or paper.get("title") or paper.get("filename")
    if heading == "outcome":
        name = outcome.get("outcome") or outcome.get("label")
        return " | ".join(str(part) for part in [name, outcome.get("group"), outcome.get("time_point"), outcome.get("units")] if part) or None
    return None


def export_preview(snapshot: dict, draft: bool = True, *, mark_unfinished: bool = True, include_accepted: bool = False) -> dict:
    snapshot = _scoped_snapshot(snapshot)
    validation = validate_snapshot(snapshot)
    if not draft and not validation["can_finalize"]:
        raise ValueError("Final export is blocked by unresolved validation issues.")
    fields = {_field_key(field): field for field in [*snapshot.get("fields", []), *snapshot.get("outcome_fields", [])]}
    accepted = _accepted_passages(snapshot) if include_accepted else {}
    items = _active_items(snapshot)
    invalid = {(issue.get("paper_id"), issue.get("item_id"), issue.get("outcome_id")) for issue in validation["issues"]
               if issue["severity"] == "error" and issue.get("item_id")}
    study_values = {}
    for item in items:
        if not _is_outcome_item(item) and (item.get("export_heading") or item.get("name", "")).strip().casefold() == "citation":
            for paper in _included_papers(snapshot):
                key = (paper["id"], item["id"], None)
                field = fields.get(key, {})
                if field.get("status") == "approved" and key not in invalid:
                    study_values[paper["id"]] = field.get("value")

    def make_row(paper, selected_items, outcome=None):
        outcome_id = outcome["id"] if outcome else None
        values, statuses, sources = [], [], []
        for item in selected_items:
            key = (paper["id"], item["id"], outcome_id)
            field = fields.get(key, {})
            status = field.get("status", "unreviewed")
            statuses.append(status)
            sources.append('saved_value' if status == 'approved' and key not in invalid else 'missingness' if status in {'not_reported', 'not_applicable'} and not field.get('needs_reconsideration') else 'empty')
            metadata = _outcome_metadata_value(item, outcome, paper, study_values) if outcome else None
            if status == "approved" and key not in invalid:
                values.append(field.get("value"))
            elif status in {"not_reported", "not_applicable"} and not field.get("needs_reconsideration"):
                values.append(STATUS_LABELS[status])
            elif status != 'approved' and not field.get('needs_reconsideration') and accepted.get(key):
                values.append('\n\n'.join(dict.fromkeys(record['quote'] for record in accepted[key])))
                sources[-1] = 'accepted_passage'
            elif status == 'unreviewed' and not field.get('needs_reconsideration') and metadata:
                values.append(metadata)
                sources[-1] = 'row_metadata'
            elif draft and mark_unfinished:
                label = "Reconsideration required" if field.get("needs_reconsideration") else STATUS_LABELS.get(status, status)
                if status == "approved" and not field.get("needs_reconsideration"):
                    label = "Approval requires validation"
                values.append(f"[DRAFT: {label}]")
            else:
                values.append(None)
        row = {"paper_id": paper["id"], "title": paper.get("title", ""), "values": values, "statuses": statuses, "sources": sources}
        if outcome:
            row.update(outcome_id=outcome_id, outcome=outcome.get("outcome"), group=outcome.get("group"), time_point=outcome.get("time_point"), units=outcome.get("units"), label=outcome.get("label"))
        return row

    def column(item):
        return {"id": item["id"], "heading": item.get("export_heading") or item["name"], "worksheet_name": item.get("worksheet_name") or ("Meta Analysis" if _is_outcome_item(item) else "Data"), "section": item.get("section", "study")}
    papers = _included_papers(snapshot)
    paper_map = {paper["id"]: paper for paper in papers}
    study_items = [item for item in items if not _is_outcome_item(item)]
    rows = [make_row(paper, study_items) for paper in papers]
    groups = {}
    for item in items:
        name = column(item)["worksheet_name"]
        key = (name.casefold(), _is_outcome_item(item))
        groups.setdefault(key, {"name": name, "items": [], "row_scope": "outcome" if _is_outcome_item(item) else "study"})["items"].append(item)
    worksheets = []
    for group in groups.values():
        selected = group["items"]
        sheet_rows = ([make_row(paper_map[outcome["paper_id"]], selected, outcome) for outcome in _active_outcomes(snapshot)]
                      if group["row_scope"] == "outcome" else [make_row(paper, selected) for paper in papers])
        worksheets.append({"name": group["name"], "row_scope": group["row_scope"], "columns": [column(item) for item in selected], "rows": sheet_rows})
    if not worksheets:
        worksheets = [{"name": "Data", "row_scope": "study", "columns": [], "rows": rows}]
    return {"columns": [column(item) for item in study_items], "rows": rows, "worksheets": worksheets, "draft": draft, "validation": validation}


def _set_cell(cell, value):
    if isinstance(value, (list, dict)):
        value = _text(value)
    if isinstance(value, str):
        if len(value) > MAX_CELL_LENGTH:
            raise ValueError("Excel cannot preserve a cell longer than 32767 characters. Shorten the value or split the narrative into additional items.")
        # Explicit string typing prevents =, +, -, @, hyperlinks and imported
        # formulas from becoming executable spreadsheet content, preserving text.
        cell.value = value
        cell.data_type = "s"
        cell.number_format = "@"
    else:
        cell.value = value
    cell.font = Font(name="Arial", size=10, color="172B3A")
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if isinstance(value, str):
        try:
            parsed = urlsplit(value)
            if parsed.scheme.lower() in {"http", "https"} and parsed.netloc and not re.search(r"\s", value):
                cell.hyperlink = value
                cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        except ValueError:
            pass




def verification_preview(snapshot: dict, draft: bool = True, *, include_accepted: bool = False) -> dict:
    """Match every exported cell to sources selected for that exact study/outcome."""
    snapshot = _scoped_snapshot(snapshot)
    preview = export_preview(snapshot, draft=draft, mark_unfinished=False, include_accepted=include_accepted)
    accepted = _accepted_passages(snapshot) if include_accepted else {}
    fields = {_field_key(field): field for field in [*snapshot.get("fields", []), *snapshot.get("outcome_fields", [])]}
    evidence = {record["id"]: record for record in snapshot.get("evidence", [])}
    pages = {(page.get("paper_id"), page.get("page_number")): page for page in snapshot.get("pages", [])}
    for table in [preview, *preview["worksheets"]]:
        for row in table["rows"]:
            for index, column in enumerate(table["columns"]):
                locations, linked = [], []
                key = (row["paper_id"], column["id"], row.get("outcome_id"))
                if row["statuses"][index] == "approved" and row["values"][index] is not None:
                    field = fields[key]
                    linked = list(dict.fromkeys([*field.get("evidence_ids", []),
                        *(eid for observation in field.get("observations", []) for eid in observation.get("evidence_ids", []))]))
                elif row['sources'][index] == 'accepted_passage':
                    linked = [record['id'] for record in accepted[key]]
                elif row['sources'][index] == 'row_metadata':
                    locations.append("Row label (review metadata; no source passage selected)")
                for eid in linked:
                    record = evidence[eid]
                    page = pages.get((row["paper_id"], record.get("page_number")), {})
                    location = _verification_location(record, page)
                    if location not in locations:
                        locations.append(location)
                row["values"][index] = "\n".join(locations) if locations else None
    return preview


def _table(book, title, headings, rows, *, allow_continuations=True):
    rows = list(rows)
    # Excel limits strings to 32767 characters. Supporting records may exceed
    # that (notably before/after audit JSON). Continuation rows retain the full
    # strings, linked by stable IDs and explicitly numbered record parts.
    if allow_continuations and title not in {"Data", "Verification"} and not title.startswith("Extraction") and any(isinstance(value, (str, dict, list)) and len(_text(value)) > MAX_CELL_LENGTH for row in rows for value in row):
        expanded = []
        for record_number, row in enumerate(rows, 1):
            values = [_text(value) if isinstance(value, (dict, list)) else value for value in row]
            parts = max([1, *(math.ceil(len(value) / MAX_CELL_LENGTH) for value in values if isinstance(value, str))])
            for part in range(parts):
                output = []
                for column, value in enumerate(values):
                    if headings[column].endswith(" ID") or headings[column] in {"Property", "Schema version", "Rule version"}:
                        output.append(value)
                    elif isinstance(value, str):
                        output.append(value[part * MAX_CELL_LENGTH:(part + 1) * MAX_CELL_LENGTH])
                    else:
                        output.append(value if part == 0 else None)
                expanded.append([*output, record_number, part + 1, parts])
        headings = [*headings, "Record row", "Record part", "Total parts"]
        rows = expanded
    ws = book.create_sheet(title)
    for column, heading in enumerate(headings, 1):
        cell = ws.cell(1, column)
        _set_cell(cell, heading)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name="Arial", size=10, bold=True, color="172B3A")
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="9CB9D2"))
    for row_number, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            _set_cell(ws.cell(row_number, column), value)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    if headings:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headings))}{max(1, ws.max_row)}"
    for column, heading in enumerate(headings, 1):
        observed = [_text(ws.cell(row, column).value) for row in range(2, min(ws.max_row, 101) + 1)]
        longest = max([len(_text(heading)), *(min(len(value), 65) for value in observed)])
        width = max(16, min(52, max(len(_text(heading)) * 0.7, longest * 0.8)))
        ws.column_dimensions[get_column_letter(column)].width = round(width, 1)
    for row in ws.iter_rows():
        lines = 1
        for cell in row:
            width = ws.column_dimensions[cell.column_letter].width
            lines = max(lines, sum(max(1, math.ceil(len(part) / max(1, width - 2))) for part in _text(cell.value).split("\n")))
        ws.row_dimensions[row[0].row].height = min(409, max(32 if row[0].row == 1 else 28, lines * 14 + 10))
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    return ws


def _safe_template_style(raw):
    """Imported/user-edited style metadata must never make an export fail."""
    if not isinstance(raw, dict):
        return {}

    def number(value, minimum, maximum):
        return value if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and minimum <= value <= maximum else None

    def color(value):
        return value if isinstance(value, str) and re.fullmatch(r"(?:[0-9a-fA-F]{6}|[0-9a-fA-F]{8})", value) else None

    def mapping(value):
        return value if isinstance(value, dict) else {}

    result = {key: number(raw.get(key), 1, maximum) for key, maximum in [('column_width', 255), ('header_height', 409), ('body_height', 409)]}
    for kind in ('header', 'body'):
        details = mapping(raw.get(kind))
        font, alignment = mapping(details.get('font')), mapping(details.get('alignment'))
        name = font.get('name')
        result[kind] = {'font': {'name': name if isinstance(name, str) and 0 < len(name) <= 255 and not re.search(r'[\x00-\x1f]', name) else 'Arial',
                                'size': number(font.get('size'), 1, 409) or 10,
                                'bold': font.get('bold') is True, 'italic': font.get('italic') is True,
                                'color': color(font.get('color'))},
                        'fill': color(details.get('fill')),
                        'alignment': {'horizontal': alignment.get('horizontal') if alignment.get('horizontal') in ('left', 'center', 'right', 'fill', 'justify', 'centerContinuous', 'distributed', 'general') else None,
                                      'vertical': alignment.get('vertical') if alignment.get('vertical') in ('top', 'center', 'bottom', 'justify', 'distributed') else 'top',
                                      'text_rotation': int(number(alignment.get('text_rotation'), 0, 180) or 0)},
                        'border': {}}
        for edge in ('left', 'right', 'top', 'bottom'):
            spec = mapping(mapping(details.get('border')).get(edge))
            border_style = spec.get('style')
            if border_style in ('hair', 'medium', 'dashed', 'dotted', 'dashDot', 'thin', 'thick', 'double', 'mediumDashDotDot', 'mediumDashed', 'mediumDashDot', 'slantDashDot', 'dashDotDot'):
                result[kind]['border'][edge] = {'style': border_style, 'color': color(spec.get('color'))}
    return result


def _apply_template_style(sheet, items, *, verification=False):
    for column, item in enumerate(items, 1):
        raw_style = item.get('template_style')
        if not isinstance(raw_style, dict) or not raw_style:
            continue
        style = _safe_template_style(raw_style)
        width = style.get('column_width')
        if isinstance(width, (int, float)) and 0 < width <= 255:
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in range(1, sheet.max_row + 1):
            details = style.get('header' if row == 1 else 'body') or {}
            cell = sheet.cell(row, column)
            font = details.get('font') or {}
            if font:
                cell.font = Font(name=font.get('name') or 'Arial', size=font.get('size') or 10,
                                 bold=bool(font.get('bold')), italic=bool(font.get('italic')), color=font.get('color'))
            cell.fill = PatternFill('solid', fgColor=details['fill']) if details.get('fill') else PatternFill()
            alignment = details.get('alignment') or {}
            if alignment:
                cell.alignment = Alignment(horizontal=alignment.get('horizontal'), vertical=alignment.get('vertical') or 'top',
                                           wrap_text=True, text_rotation=alignment.get('text_rotation') or 0)
            border = details.get('border') or {}
            cell.border = Border(**{edge: Side(style=spec.get('style'), color=spec.get('color')) for edge, spec in border.items() if edge in {'left', 'right', 'top', 'bottom'}})
            if row > 1 and isinstance(cell.value, str):
                cell.number_format = '@'
        if isinstance(style.get('header_height'), (int, float)):
            sheet.row_dimensions[1].height = max(sheet.row_dimensions[1].height or 0, style['header_height'])
        # Body rows grow to accommodate actual extraction text and locators.
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, column)
            width = sheet.column_dimensions[get_column_letter(column)].width
            lines = sum(max(1, math.ceil(len(part) / max(1, width - 2))) for part in _text(cell.value).split('\n'))
            height = style.get('body_height') or 28
            sheet.row_dimensions[row].height = min(409, max(sheet.row_dimensions[row].height or 0, height, lines * 14 + 10))


def _write_design_sheets(book, snapshot, data, locations, view):
    """Write all data sheets first, then their exact row/column location copies."""
    items = {item['id']: item for item in _active_items(snapshot)}
    groups = [dict(table) for table in data['worksheets']]
    used = {name.casefold() for name in book.sheetnames}
    def sheet_title(name):
        base = re.sub(r'[\[\]:*?/\\\x00-\x1f]', '_', name).strip("'") or 'Sheet'
        title, suffix = base[:31], 1
        while title.casefold() in used:
            suffix += 1
            tail = f' ({suffix})'
            title = base[:31-len(tail)] + tail
        used.add(title.casefold())
        return title
    for group in groups:
        group['data_title'] = sheet_title(group['name'])
    for group in groups:
        name = 'Verification' if len(groups) == 1 and group['name'] == 'Data' else f"Verify - {group['name']}"
        group['verification_title'] = sheet_title(name)
    papers = {p['id']: p for p in snapshot['papers']}
    review = snapshot['review']
    for kind, preview in [('data', data), ('verification', locations)]:
        if view != 'combined' and view != kind:
            continue
        for group_index, group in enumerate(groups):
            table = preview['worksheets'][group_index]
            group_items = [items[column['id']] for column in table['columns']]
            sheet = _table(book, group[f'{kind}_title'], [column['heading'] for column in table['columns']],
                           [row['values'] for row in table['rows']], allow_continuations=False)
            sheet.oddHeader.center.text = 'Extraction data' if kind == 'data' else 'Source locations'
            sheet.sheet_properties.tabColor = '426B8D' if kind == 'data' else '548235'
            for column, item in enumerate(group_items, 1):
                sheet.cell(1, column).comment = Comment(f"Worksheet: {group['name']}\nItem ID: {item['id']}\nReview ID: {review.get('id')}\nRow scope: {group['row_scope']}", 'Systematic')
            for row_number, row in enumerate(table['rows'], 2):
                paper = papers[row['paper_id']]
                for column, item in enumerate(group_items, 1):
                    index, cell = column - 1, sheet.cell(row_number, column)
                    location = locations['worksheets'][group_index]['rows'][row_number - 2]['values'][index] or ''
                    context = f"\nOutcome ID: {row['outcome_id']}\nOutcome: {row.get('outcome') or row.get('label') or ''}\nGroup: {row.get('group') or ''}\nTime point: {row.get('time_point') or ''}\nUnits: {row.get('units') or ''}" if row.get('outcome_id') else ''
                    cell.comment = Comment(f"Paper: {row['title']}\nFilename: {paper.get('filename', '')}\nPaper ID: {row['paper_id']}\nItem ID: {item['id']}{context}\nContent: {row['sources'][index]}\nField status: {row['statuses'][index]}\nSource: {location}", 'Systematic')
                    if kind == 'data' and row['sources'][index] == 'saved_value':
                        if item.get('data_type') == 'date' and isinstance(cell.value, str):
                            cell.value = date.fromisoformat(cell.value)
                            cell.number_format = 'yyyy-mm-dd'
                        elif item.get('data_type') == 'integer':
                            cell.number_format = '0'
                        elif item.get('data_type') == 'number':
                            cell.number_format = 'General'
            _apply_template_style(sheet, group_items, verification=kind == 'verification')


def export_download_workbook(snapshot: dict, view: str = 'data') -> bytes:
    """Focused Excel downloads that match the preview, without candidate audits."""
    if view not in {'data', 'verification', 'combined'}:
        raise ValueError('Choose data, verification, or combined.')
    snapshot = _scoped_snapshot(snapshot)
    data = export_preview(snapshot, mark_unfinished=False, include_accepted=True)
    locations = verification_preview(snapshot, include_accepted=True)
    book = Workbook()
    book.remove(book.active)
    _write_design_sheets(book, snapshot, data, locations, view)
    review = snapshot['review']
    book.properties.title = f"{review.get('name', 'Review')} — {view}"
    book.properties.creator = review.get('reviewer') or 'Systematic'
    book.properties.description = (
        f"Review ID: {review.get('id')}. Items needing attention: {data['validation']['counts']['attention_fields']}. "
        'Saved final values take priority; otherwise current accepted passages appear verbatim. '
        'Blank cells do not imply missingness. Cell notes retain source identities and field status. '
        'Locations use physical PDF pages and 1-based extracted-text lines, which may differ from printed lines. '
        'Image sources use recorded figure/table/region descriptions; unknown details are flagged. '
        'All data worksheets precede matching verification worksheets. No extraction decisions are changed by export.')
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def export_workbook(snapshot: dict, draft: bool = True, *, verification: bool = False) -> bytes:
    """Return an XLSX with active-design data and stable-ID supporting records."""
    snapshot = _scoped_snapshot(snapshot)
    preview = verification_preview(snapshot, draft=draft) if verification else export_preview(snapshot, draft=draft)
    book = Workbook()
    book.remove(book.active)
    review = snapshot.get("review", {})
    items = _active_items(snapshot)
    papers = _included_papers(snapshot)
    paper_ids = {paper["id"] for paper in papers}
    all_items = {item["id"]: item for item in snapshot.get("items", [])}
    item_ids = {item["id"] for item in items}
    all_papers = {paper["id"]: paper for paper in snapshot.get("papers", [])}
    fields = {(field["paper_id"], field["item_id"]): field for field in snapshot.get("fields", [])}
    if any(_is_outcome_item(item) for item in items):
        data = export_preview(snapshot, draft=draft)
        locations = verification_preview(snapshot, draft=draft)
        _write_design_sheets(book, snapshot, data, locations, 'verification' if verification else 'combined')
    else:
        main = _table(book, "Extraction DRAFT" if draft else "Extraction", [column["heading"] for column in preview["columns"]],
                      [row["values"] for row in preview["rows"]])
        main.sheet_properties.tabColor = "F0AD4E" if draft else "426B8D"
        main.oddHeader.center.text = ("DRAFT source locations" if draft else "Source locations") if verification else ("DRAFT — unfinished fields are marked" if draft else "Final extraction")
        main.oddFooter.right.text = "Page &P of &N"
        for row_number, row in enumerate(preview["rows"], 2):
            for column, item in enumerate(items, 1):
                value = row["values"][column - 1]
                cell = main.cell(row_number, column)
                if not verification and row["statuses"][column - 1] == "approved" and not (isinstance(value, str) and value.startswith("[DRAFT:")):
                    if item.get("data_type") == "date" and isinstance(value, str):
                        cell.value = date.fromisoformat(value)
                        cell.number_format = "yyyy-mm-dd"
                    elif item.get("data_type") == "integer":
                        cell.number_format = "0"
                    elif item.get("data_type") == "number":
                        cell.number_format = "0.###############"
                if isinstance(value, str) and value.startswith("[DRAFT:"):
                    cell.fill = PatternFill("solid", fgColor="FFF3CD")

    evidence_rows = []
    for record in snapshot.get("evidence", []):
        paper, item = all_papers.get(record.get("paper_id"), {}), all_items.get(record.get("item_id"), {})
        evidence_rows.append([record.get("id"), record.get("paper_id"), record.get("item_id"), paper.get("title"), paper.get("sha256"),
                              paper.get("filename"), item.get("export_heading", item.get("name")), record.get("decision"), record.get("quote"),
                              record.get("page_number"), record.get("start"), record.get("end"), record.get("boxes", []), record.get("ocr", False),
                              record.get("manual", False), record.get("rationale"), record.get("notes"), record.get("reviewer"), record.get("created_at"), record.get("updated_at"),
                              record.get("suggestion_id"), record.get("paper_id") in paper_ids, not item.get("deleted", False),
                              record.get("schema_version", review.get("schema_version")), record.get("rule_version", review.get("rule_version")),
                              record.get("source_stale", False), record.get("document_version"), paper.get("document_version"),
                              record.get("source_location"), record.get("table_title"), record.get("row_header"), record.get("column_header"), record.get("outcome_id")])
    _table(book, "Evidence", ["Evidence ID", "Paper ID", "Item ID", "Paper title", "Document SHA-256", "Filename", "Item heading", "Evidence decision", "Exact source quote",
                              "Physical PDF page", "Start offset", "End offset", "Highlight coordinates", "OCR derived", "Manual transcription", "Manual transcription rationale", "Notes", "Reviewer", "Created UTC", "Updated UTC",
                              "Suggestion ID", "Paper included", "Item active", "Schema version", "Rule version", "Source stale", "Evidence document version", "Current document version", "Source location descriptor", "Table title", "Table row label", "Table column label", "Outcome ID"], evidence_rows)

    status_rows, observation_rows = [], []
    for paper in snapshot.get("papers", []):
        for item in snapshot.get("items", []):
            if _is_outcome_item(item):
                continue
            field = fields.get((paper["id"], item["id"]), {})
            status_rows.append([paper["id"], item["id"], paper.get("title"), item.get("export_heading", item.get("name")),
                                STATUS_LABELS.get(field.get("status", "unreviewed"), field.get("status")), field.get("status", "unreviewed"),
                                field.get("value"), field.get("value_kind", "reported"), field.get("rationale"), field.get("calculation_inputs", []),
                                field.get("notes"), field.get("evidence_ids", []), field.get("needs_reconsideration", False), field.get("reviewer"), field.get("updated_at"),
                                paper.get("included", True), not item.get("deleted", False), paper.get("processing_status"), paper.get("search_status"), paper.get("manual_review_ack", False),
                                field.get("schema_version", review.get("schema_version")), field.get("rule_version", review.get("rule_version"))])
            for position, observation in enumerate(field.get("observations", []), 1):
                observation_rows.append([observation.get("id"), paper["id"], item["id"], position, observation.get("group"), observation.get("population"),
                                         observation.get("outcome"), observation.get("time_point"), observation.get("value"), observation.get("units"), observation.get("evidence_ids", []),
                                         field.get("status", "unreviewed"), field.get("value_kind", "reported"), field.get("reviewer"), field.get("updated_at"),
                                         paper.get("included", True), not item.get("deleted", False), observation.get("statistic")])
    _table(book, "Review status", ["Paper ID", "Item ID", "Paper title", "Item heading", "Status", "Status code", "Stored value (not necessarily approved)", "Value kind", "Rationale",
                                   "Calculation inputs", "Notes", "Linked evidence IDs", "Needs reconsideration", "Reviewer", "Updated UTC", "Paper included", "Item active", "PDF processing", "Search status",
                                   "Manual review acknowledged", "Schema version", "Rule version"], status_rows)
    _table(book, "Data dictionary", ["Item ID", "Export position", "Display name", "Export heading", "Definition", "Instructions", "Data type", "Allowed values", "Units", "Required", "Context", "Schema version", "Rule version", "Worksheet", "Section"],
           [[item["id"], index, item.get("name"), item.get("export_heading"), item.get("definition"), item.get("instructions"), item.get("data_type"), item.get("allowed_values", []),
             item.get("units"), item.get("required", False), item.get("context", {}), item.get("schema_version"), item.get("rule_version"), item.get("worksheet_name"), item.get("section", "study")] for index, item in enumerate(items, 1)])
    keyword_rows = []
    for item in items:
        for keyword in item.get("keywords", []):
            keyword_rows.append([item["id"], item.get("name"), "keyword", keyword, "literal", "", None, item.get("rule_version")])
        for rule in item.get("rules", []):
            keyword_rows.append([item["id"], item.get("name"), "rule", rule.get("pattern"), rule.get("kind", "literal"), rule.get("label"), rule.get("capture_group"), item.get("rule_version")])
    _table(book, "Keyword mappings", ["Item ID", "Item name", "Mapping type", "Keyword or pattern", "Match kind", "Rule label", "Capture group", "Rule version"], keyword_rows)
    _table(book, "Decision history", ["History ID", "Review ID", "Paper ID", "Item ID", "Action", "Before", "After", "Reviewer", "Created UTC", "Schema version", "Rule version", "Outcome ID"],
           [[event.get(key) for key in ("id", "review_id", "paper_id", "item_id", "action", "before", "after", "reviewer", "created_at", "schema_version", "rule_version", "outcome_id")]
            for event in snapshot.get("history", [])])
    _table(book, "Observations", ["Observation ID", "Paper ID", "Item ID", "Observation position", "Treatment group", "Population", "Outcome", "Time point", "Stored value", "Units", "Linked evidence IDs",
                                  "Field status", "Value kind", "Reviewer", "Updated UTC", "Paper included", "Item active", "Statistic type"], observation_rows)
    if snapshot.get("outcomes") or any(_is_outcome_item(item) for item in items):
        _table(book, "Outcome rows", ["Outcome ID", "Review ID", "Paper ID", "Label", "Outcome", "Treatment group", "Time point", "Units", "Included", "Deleted", "Position", "Created UTC", "Updated UTC"],
               [[row.get(key) for key in ("id", "review_id", "paper_id", "label", "outcome", "group", "time_point", "units", "included", "deleted", "position", "created_at", "updated_at")]
                for row in snapshot.get("outcomes", [])])
        _table(book, "Outcome values", ["Outcome ID", "Review ID", "Paper ID", "Item ID", "Field status", "Stored value (not necessarily approved)", "Value kind", "Rationale", "Calculation inputs", "Notes", "Linked evidence IDs", "Needs reconsideration", "Reviewer", "Updated UTC", "Schema version", "Rule version"],
               [[field.get(key) for key in ("outcome_id", "review_id", "paper_id", "item_id", "status", "value", "value_kind", "rationale", "calculation_inputs", "notes", "evidence_ids", "needs_reconsideration", "reviewer", "updated_at", "schema_version", "rule_version")]
                for field in snapshot.get("outcome_fields", [])])
    _table(book, "Search decisions", ["Suggestion ID", "Paper ID", "Item ID", "Physical PDF page", "Start offset", "End offset", "Exact passage", "Context", "Matched rules", "Proposed value (not approved)",
                                      "Decision", "Active", "OCR derived", "Rule version", "Created UTC"],
           [[suggestion.get(key) for key in ("id", "paper_id", "item_id", "page_number", "start", "end", "quote", "context", "matched_rules", "proposed_value", "decision", "active", "ocr", "rule_version", "created_at")]
            for suggestion in snapshot.get("suggestions", [])])
    _table(book, "Validation", ["Severity", "Code", "Issue", "Paper ID", "Item ID", "Outcome ID"],
           [[issue.get(key) for key in ("severity", "code", "message", "paper_id", "item_id", "outcome_id")] for issue in preview["validation"]["issues"]])
    _table(book, "Retrieval audit", ["Suggestion ID", "Paper ID", "Item ID", "Rank", "Relevance score", "Component scores", "Weights", "Section", "Table ID", "Row", "Column", "Extracted proposal", "Validation", "Chunk ID", "Origin", "Reviewer feedback adjustment"],
           [[row.get(key) for key in ("id", "paper_id", "item_id", "rank", "rank_score", "scores", "ranking_weights", "section", "table_id", "table_row", "table_column", "extracted_value", "value_validation", "chunk_id", "origin", "learning")]
            for row in snapshot.get("suggestions", [])])
    _table(book, "Retrieval feedback", ["Feedback ID", "Paper ID", "Item ID", "Suggestion ID", "Accepted", "Rank", "Scores", "Candidate text", "Item definition snapshot", "Candidate snapshot", "Engine", "Created UTC", "Example ID", "Evidence ID", "Decision", "Rejection reason", "Learning epoch", "Feedback sequence", "Item meaning version"],
           [[row.get(key) for key in ("id", "paper_id", "item_id", "suggestion_id", "accepted", "rank", "scores", "candidate_text", "item", "candidate", "engine_version", "created_at", "example_id", "evidence_id", "decision", "rejection_reason", "epoch", "sequence", "item_meaning_version")]
            for row in snapshot.get("feedback", [])])
    _table(book, "Learning examples", ["Example ID", "Feedback ID", "Paper ID", "Item ID", "Quote", "Decision", "Rejection reason", "Eligible", "Exclusion reason", "Excluded by reviewer"],
           [[row.get(key) for key in ("example_id", "feedback_id", "paper_id", "item_id", "quote", "decision", "rejection_reason", "eligible", "excluded_reason", "excluded")]
            for row in snapshot.get("learning_examples", [])])
    now = datetime.now(timezone.utc).isoformat()
    _table(book, "Export metadata", ["Property", "Value"], [
        ["Export status", "DRAFT" if draft else "FINAL"], ["Exported UTC", now], ["Review ID", review.get("id")], ["Review name", review.get("name")],
        ["Reviewer", review.get("reviewer")], ["Schema version", review.get("schema_version")], ["Rule version", review.get("rule_version")],
        ["Review created UTC", review.get("created_at")], ["Review updated UTC", review.get("updated_at")],
        ["Learning configuration", review.get("learning", {"enabled":False,"version":1,"epoch":0,"revision":0})],
        ["Included papers", len(papers)], ["Active extraction items", len(item_ids)],
        ["Included outcome rows", len(_active_outcomes(snapshot))],
        ["Workbook view", "Verification source locations" if verification else "Extraction values"],
        ["Main table policy", "Only valid approved values and explicit missingness decisions appear as finalized data. Supporting sheets retain historical and unfinished records with their statuses."],
        ["Optional unfinished fields", "Empty cells in FINAL exports remain Not yet reviewed in Review status. They are not Not reported."],
        ["Evidence locations", "Page numbers are physical PDF pages. Coordinates are top-left PDF page units. Exact source quotations are separate from normalized values."],
        ["Verification locations", "Each locator identifies linked accepted evidence for a saved final value. Lines are 1-based in Systematic's extracted page text, counting blank lines, and may differ from the PDF's printed line numbers. Multiline passages show a line range. Image/OCR and manual sources use recorded figure/table/region descriptions and selected PDF coordinates instead of invented line numbers. Internal table IDs never stand in for printed table numbers. Missing precise regions are explicitly flagged. All linked locations are retained. Missingness decisions and unfinished values have blank location cells."],
        ["Long text", "Excel row heights are limited. Open the formula bar to read any long cell that exceeds the visible row height."],
        ["Long supporting records", "When a supporting cell exceeds 32767 characters, numbered continuation rows preserve all text. Rejoin each non-ID text column in Record part order within Record row. Stable IDs are repeated for linkage."],
    ])
    book.properties.title = f"{'Verification ' if verification else ''}{'DRAFT' if draft else 'FINAL'} — {review.get('name', 'Review extraction')}"
    book.properties.creator = review.get("reviewer") or "Review Extractor"
    book.properties.description = "Local research extraction with evidence decisions, stable identifiers, and audit history."
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()
