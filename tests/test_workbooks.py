"""Template and Excel safety/traceability checks use synthetic source records."""

from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from review_extractor.workbooks import export_preview, export_workbook, export_download_workbook, preview_template, validate_snapshot, verification_preview


def test_accepted_passages_fill_only_unsaved_cells_without_changing_decisions():
    data = snapshot()
    data['fields'][0].update(status='unreviewed', value=999, evidence_ids=[])
    before = deepcopy(data)
    preview = export_preview(data, mark_unfinished=False, include_accepted=True)
    assert preview['rows'][0]['values'] == [data['evidence'][0]['quote']]
    assert preview['rows'][0]['sources'] == ['accepted_passage']
    assert verification_preview(data, include_accepted=True)['rows'][0]['values'] == ['Page 1, line 1']
    assert data == before
    data['fields'][0].update(status='approved', value=0, evidence_ids=['evidence-1'])
    assert export_preview(data, include_accepted=True)['rows'][0]['values'] == [0]


@pytest.mark.parametrize('change', [dict(decision='pending'), dict(decision='rejected'), dict(source_stale=True), dict(start=10), dict(document_version=99), dict(review_id='other'), dict(paper_id='other')])
def test_accepted_fallback_never_includes_invalid_or_foreign_sources(change):
    data = snapshot()
    data['fields'][0].update(status='unreviewed', value=None)
    data['evidence'][0].update(change)
    assert export_preview(data, mark_unfinished=False, include_accepted=True)['rows'][0]['values'] == [None]


@pytest.mark.parametrize('view', ['data', 'verification', 'combined'])
def test_focused_downloads_match_and_keep_untrusted_accepted_text_literal(view):
    data = snapshot()
    data['fields'][0].update(status='unreviewed', value=None)
    quote = '=HYPERLINK("https://example.invalid", "Do not execute")'
    data['pages'][0]['text'] = quote
    data['evidence'][0].update(quote=quote, start=0, end=len(quote))
    book = load_workbook(BytesIO(export_download_workbook(data, view)))
    expected = ['Data'] if view == 'data' else ['Verification'] if view == 'verification' else ['Data', 'Verification']
    assert book.sheetnames == expected
    if view != 'verification':
        assert book['Data']['A2'].value == quote
        assert book['Data']['A2'].data_type == 's'
    if view != 'data':
        assert book['Verification']['A2'].value == 'Page 1, line 1'
    assert 'Paper ID: paper-1' in book.worksheets[0]['A2'].comment.text
    assert 'accepted_passage' in book.worksheets[0]['A2'].comment.text


def test_multiworksheet_template_preserves_each_design_and_excludes_examples():
    book = Workbook()
    book.active.title = 'Study details'
    for sheet in [book.active, book.create_sheet('Outcomes'), book.create_sheet('Risk of bias')]:
        sheet.append(['Result', 'Notes'])
        sheet.append(['NEVER IMPORT EXAMPLE', 999])
    book.create_sheet('Blank')
    result = preview_template(workbook_bytes(book), all_sheets=True)
    assert [item['worksheet_name'] for item in result['items']] == ['Study details']*2 + ['Outcomes']*2 + ['Risk of bias']*2
    assert [item['position'] for item in result['items']] == list(range(6))
    assert all('value' not in item for item in result['items'])
    assert any('Blank' in warning and 'No extraction items' in warning for warning in result['warnings'])
    selected = preview_template(workbook_bytes(book), sheet='Outcomes')
    assert {item['worksheet_name'] for item in selected['items']} == {'Outcomes'}


def test_combined_exports_all_data_sheets_before_matching_verification_sheets():
    data = snapshot()
    original = deepcopy(data)
    data['items'] = []
    data['fields'] = []
    data['evidence'] = []
    names = ['Study details', 'Outcomes', 'Risk of bias']
    for index, name in enumerate(names):
        item = dict(original['items'][0], id=f'item-{index}', worksheet_name=name, position=index)
        field = dict(original['fields'][0], item_id=item['id'], value=index, evidence_ids=[f'e-{index}'], observations=[])
        evidence = dict(original['evidence'][0], id=f'e-{index}', item_id=item['id'])
        data['items'].append(item); data['fields'].append(field); data['evidence'].append(evidence)
    for view in ['data', 'verification', 'combined']:
        book = load_workbook(BytesIO(export_download_workbook(data, view)))
        expected = (names if view != 'verification' else []) + ([f'Verify - {name}' for name in names] if view != 'data' else [])
        assert book.sheetnames == expected
        for index, name in enumerate(names):
            if view != 'verification':
                assert book[name]['A2'].value == index
                assert book[name].max_column == 1
            if view != 'data':
                assert book[f'Verify - {name}']['A2'].value == 'Page 1, line 1'
    data['items'][1]['deleted'] = True
    book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
    assert book.sheetnames == ['Study details', 'Risk of bias', 'Verify - Study details', 'Verify - Risk of bias']


def test_long_and_colliding_verification_names_are_unique_and_excel_safe():
    data = snapshot()
    names = ['A'*31, 'A'*30+'B', 'Verify - '+('A'*22)]
    data['items'] = [dict(data['items'][0], id=f'item-{i}', worksheet_name=name, position=i) for i, name in enumerate(names)]
    data['fields'] = []; data['evidence'] = []
    book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
    assert book.sheetnames[:3] == names
    assert len(book.sheetnames) == len(set(name.casefold() for name in book.sheetnames)) == 6
    assert all(len(name) <= 31 for name in book.sheetnames)


def workbook_bytes(book):
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def snapshot():
    item = {"id": "item-1", "review_id": "review-1", "name": "Result", "export_heading": "Result", "data_type": "integer",
            "position": 0, "required": True, "keywords": ["events"], "rules": [{"pattern": r"(\d+) events", "kind": "regex", "label": "Event count", "capture_group": 1}],
            "schema_version": 2, "rule_version": 3}
    paper = {"id": "paper-1", "review_id": "review-1", "title": "Synthetic trial", "filename": "synthetic.pdf", "sha256": "synthetic-document-hash",
             "included": True, "processing_status": "ready", "search_status": "current", "manual_review_ack": False}
    evidence = {"id": "evidence-1", "paper_id": "paper-1", "item_id": "item-1", "page_number": 1, "quote": "Treatment: 0 adverse events at week 3.",
                "boxes": [{"x0": 10, "top": 20, "x1": 150, "bottom": 32}], "start": 0, "end": len("Treatment: 0 adverse events at week 3."), "ocr": False, "decision": "accepted", "reviewer": "Researcher",
                "created_at": "2026-01-01T00:00:00+00:00", "updated_at": "2026-01-01T00:00:00+00:00"}
    field = {"paper_id": "paper-1", "item_id": "item-1", "status": "approved", "value": 0, "value_kind": "reported", "evidence_ids": ["evidence-1"],
             "reviewer": "Researcher", "schema_version": 2, "rule_version": 3,
             "observations": [{"id": "obs-1", "group": "Treatment", "time_point": "Week 3", "value": 0, "units": "events", "evidence_ids": ["evidence-1"]}]}
    return {"review": {"id": "review-1", "name": "Synthetic review", "reviewer": "Researcher", "schema_version": 2, "rule_version": 3},
            "items": [item], "papers": [paper], "pages": [{"paper_id": "paper-1", "page_number": 1, "status": "ok", "text": evidence["quote"]}], "fields": [field], "evidence": [evidence],
            "suggestions": [{"id": "suggestion-1", "paper_id": "paper-1", "item_id": "item-1", "quote": evidence["quote"], "decision": "rejected"}],
            "history": [{"id": "history-1", "review_id": "review-1", "paper_id": "paper-1", "item_id": "item-1", "action": "field_approved", "after": field, "reviewer": "Researcher", "schema_version": 2, "rule_version": 3}]}


def test_reference_headings_match_actual_workbook_when_available():
    reference = json.loads((Path(__file__).parents[1] / "review_extractor/resources/reference_items.json").read_text(encoding="utf-8"))
    expected = [item["export_heading"] for item in reference]
    assert len(expected) == 29
    assert expected[:5] == ["Citation", "Link", "Country / setting", "Study design", "Mean age"]
    assert expected[10] == "Baseline / reported GMS outcome scores"
    assert expected[15] == "Hz / burst pattern"
    assert expected[-1] == "Follow-up timing"
    assert all("value" not in item for item in reference)
    actual = Path("C:/Users/diego/Downloads/Systematic Example.xlsx")
    if actual.exists():
        preview = preview_template(actual.read_bytes(), sheet="Sheet1")
        assert [item["export_heading"] for item in preview["items"]] == expected
        assert len(preview["preview_rows"]) == 1
        assert preview["example_included"] is False


def test_columns_sheet_choice_merged_multirow_and_preview_only():
    book = Workbook()
    book.active.title = "Instructions"
    book.active.append(["Template instructions"])
    data = book.create_sheet("Design")
    data.append(["Treatment", None, "Control", None, None])
    data.append(["Value", "Time point", "Value", "Time point", None])
    data.merge_cells("A1:B1")
    data.merge_cells("C1:D1")
    data.append([999, "EXAMPLE ONLY", 500, "EXAMPLE ONLY", "unmapped"])
    preview = preview_template(workbook_bytes(book), sheet="Design", header_rows="1,2")
    assert preview["sheets"] == ["Instructions", "Design"]
    assert [i["export_heading"] for i in preview["items"]] == ["Treatment / Value", "Treatment / Time point", "Control / Value", "Control / Time point"]
    assert len(preview["items"]) == 4 and preview["example_included"] is False
    assert preview["items"][0]["keywords"] == ["Treatment", "Value"]
    assert any("skipped" in message for message in preview["warnings"])
    assert not any("value" in item for item in preview["items"])


def test_dictionary_mapping_types_keywords_context_and_order():
    book = Workbook()
    book.active.append(["Variable", "Description", "Type", "Required", "Keywords", "Units", "Group", "Choices"])
    book.active.append(["Event count", "Count at week 3", "integer", "yes", "adverse events;side effects", "events", "Treatment", "[0, 1, 2]"])
    book.active.append(["Interpretation", "Retain discrepancy", "narrative", "no", "flow diagram\ntext", None, "Control", None])
    preview = preview_template(workbook_bytes(book), mode="dictionary")
    first, second = preview["items"]
    assert first["required"] and first["data_type"] == "integer"
    assert first["keywords"] == ["adverse events", "side effects"]
    assert first["allowed_values"] == [0, 1, 2]
    assert first["context"] == {"group": "Treatment"}
    assert second["required"] is False
    explicit = preview_template(workbook_bytes(book), mode="dictionary", mapping={"name": 2, "export_heading": 1, "keywords": 5})
    assert explicit["items"][0]["name"] == "Count at week 3"
    assert explicit["items"][0]["export_heading"] == "Event count"
    assert explicit["mapping"]["name"] == 2


@pytest.mark.parametrize("options", [{"sheet": "Missing"}, {"header_rows": "0"}, {"header_rows": "2"}, {"mode": "bad"}, {"mode": "dictionary", "mapping": {"name": 0}}])
def test_invalid_template_configuration_is_actionable(options):
    book = Workbook()
    book.active.append(["Heading"])
    with pytest.raises(ValueError):
        preview_template(workbook_bytes(book), **options)


def test_arbitrary_template_column_count_and_duplicate_headings():
    book = Workbook()
    names = [f"Column {number}" for number in range(35)] + ["Column 0"]
    book.active.append(names)
    preview = preview_template(workbook_bytes(book))
    assert [item["name"] for item in preview["items"]] == names
    assert any("Repeated headings" in warning for warning in preview["warnings"])


def test_numeric_zero_final_traceability_and_reference_styling():
    data = snapshot()
    assert validate_snapshot(data)["can_finalize"]
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    main = book.worksheets[0]
    assert main.title == "Extraction"
    assert main["A2"].value == 0 and main["A2"].data_type == "n"
    assert main["A1"].fill.fgColor.rgb == "FFDCE6F1"
    assert main["A1"].alignment.wrap_text and main["A2"].alignment.vertical == "top"
    assert main.freeze_panes == "A2" and main.auto_filter.ref == "A1:A2"
    assert main.column_dimensions["A"].width >= 16 and main.row_dimensions[2].height >= 28
    assert {"Evidence", "Review status", "Data dictionary", "Keyword mappings", "Decision history", "Observations", "Export metadata"}.issubset(book.sheetnames)
    evidence_rows = list(book["Evidence"].values)
    assert "Exact source quote" in evidence_rows[0]
    assert data["evidence"][0]["quote"] in evidence_rows[1]
    assert "synthetic-document-hash" in evidence_rows[1]
    observations = list(book["Observations"].values)
    assert "Treatment" in observations[1] and "Week 3" in observations[1] and 0 in observations[1]
    assert "field_approved" in str(list(book["Decision history"].values))
    assert "Researcher" in str(list(book["Export metadata"].values))


def test_current_design_rename_reorder_deleted_new_papers_and_no_mutation():
    data = snapshot()
    before = deepcopy(data)
    data["items"][0].update(name="Renamed", export_heading="Renamed result", position=2)
    data["items"].append({"id": "item-2", "review_id": "review-1", "name": "New narrative", "export_heading": "New narrative", "data_type": "narrative", "position": 0})
    data["items"].append({"id": "item-3", "review_id": "review-1", "name": "Deleted", "position": 1, "deleted": True})
    data["papers"].append({**data["papers"][0], "id": "paper-2", "title": "Additional synthetic study"})
    data["papers"].append({**data["papers"][0], "id": "excluded", "included": False})
    preview = export_preview(data)
    assert [column["heading"] for column in preview["columns"]] == ["New narrative", "Renamed result"]
    assert len(preview["rows"]) == 2
    assert preview["rows"][0]["values"] == ["[DRAFT: Not yet reviewed]", 0]
    assert preview["rows"][1]["values"][1] == "[DRAFT: Not yet reviewed]"
    assert data["fields"] == before["fields"]


@pytest.mark.parametrize("status, expected", [("unreviewed", "Not yet reviewed"), ("unclear", "Unclear"), ("conflicting", "Conflicting evidence")])
def test_draft_unfinished_statuses_never_export_proposed_values(status, expected):
    data = snapshot()
    data["fields"][0].update(status=status, value=999)
    assert not validate_snapshot(data)["can_finalize"]
    assert export_preview(data)["rows"][0]["values"] == [f"[DRAFT: {expected}]"]
    with pytest.raises(ValueError, match="blocked"):
        export_workbook(data, draft=False)


@pytest.mark.parametrize("status, expected", [("not_reported", "Not reported"), ("not_applicable", "Not applicable")])
def test_explicit_missingness_is_separate_from_zero_and_no_matches(status, expected):
    data = snapshot()
    data["fields"][0].update(status=status, value=None, evidence_ids=[], rationale="Explicitly reviewed the full paper.")
    assert validate_snapshot(data)["can_finalize"]
    assert export_preview(data, draft=False)["rows"][0]["values"] == [expected]
    data["fields"] = []
    data["suggestions"] = []
    assert export_preview(data)["rows"][0]["values"] == ["[DRAFT: Not yet reviewed]"]


@pytest.mark.parametrize("change, code", [
    (lambda data: data["fields"][0].update(evidence_ids=[]), "missing_evidence"),
    (lambda data: data["evidence"][0].update(decision="rejected"), "invalid_evidence"),
    (lambda data: data["evidence"][0].update(paper_id="other-paper"), "invalid_evidence"),
    (lambda data: data["evidence"][0].update(source_stale=True), "stale_evidence_source"),
    (lambda data: data["fields"][0].update(needs_reconsideration=True), "reconsideration"),
    (lambda data: data["fields"][0].update(value="zero"), "invalid_value"),
    (lambda data: data["fields"][0].update(value_kind="interpretation"), "missing_rationale"),
    (lambda data: data["fields"][0].update(value_kind="calculated", rationale="Count derived"), "missing_calculation_inputs"),
    (lambda data: data["papers"][0].update(search_status="stale"), "search_stale"),
    (lambda data: data["papers"][0].update(search_status="partial"), "search_incomplete"),
    (lambda data: data["papers"][0].update(processing_status="error"), "processing_incomplete"),
])
def test_final_gates_invalid_approvals_and_unresolved_processing(change, code):
    data = snapshot()
    change(data)
    validation = validate_snapshot(data)
    assert not validation["can_finalize"]
    assert code in {issue["code"] for issue in validation["issues"]}


def test_ocr_pages_require_explicit_manual_review_acknowledgment():
    data = snapshot()
    data["pages"].append({"paper_id": "paper-1", "page_number": 2, "status": "ocr_required"})
    assert not validate_snapshot(data)["can_finalize"]
    data["papers"][0]["manual_review_ack"] = True
    validation = validate_snapshot(data)
    assert validation["can_finalize"]
    assert any(issue["code"] == "unsearchable_pages" and issue["severity"] == "warning" for issue in validation["issues"])


def test_formula_injection_is_literal_in_every_sheet_and_safe_urls_clickable():
    data = snapshot()
    attack = '=HYPERLINK("https://example.invalid","unsafe")'
    data["items"][0].update(data_type="narrative", export_heading="=COLUMN()", keywords=["+SUM(1,2)"])
    data["fields"][0]["value"] = attack
    data["evidence"][0]["quote"] = "@SUM(1,2)"
    data["evidence"][0]["end"] = len("@SUM(1,2)")
    data["pages"][0]["text"] = "@SUM(1,2)"
    data["fields"][0]["notes"] = "-2+3"
    data["history"][0]["action"] = "=NOW()"
    book = load_workbook(BytesIO(export_workbook(data, draft=False)), data_only=False)
    assert book.worksheets[0]["A1"].value == "=COLUMN()"
    assert book.worksheets[0]["A2"].value == attack
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    assert book.worksheets[0]["A2"].hyperlink is None
    data["fields"][0]["value"] = "https://example.org/article/123"
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert book.worksheets[0]["A2"].hyperlink.target == "https://example.org/article/123"
    data["fields"][0]["value"] = "file:///C:/Windows/calc.exe"
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert book.worksheets[0]["A2"].hyperlink is None


def test_groups_narrative_boolean_dates_and_long_cells_preserve_types():
    data = snapshot()
    narrative = "Treatment: 0 events at week 3.\nControl: 2 events at week 5. Counts disagree across sources; no average calculated."
    data["items"][0]["data_type"] = "narrative"
    data["fields"][0]["value"] = narrative
    data["fields"][0]["observations"].append({"id": "obs-2", "group": "Control", "time_point": "Week 5", "value": 2, "units": "events", "evidence_ids": ["evidence-1"]})
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert book.worksheets[0]["A2"].value == narrative
    assert book["Observations"].max_row == 3
    data["items"][0]["data_type"] = "boolean"
    data["fields"][0]["value"] = False
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert book.worksheets[0]["A2"].value is False and book.worksheets[0]["A2"].data_type == "b"
    data["items"][0]["data_type"] = "date"
    data["fields"][0]["value"] = "2026-01-12"
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert book.worksheets[0]["A2"].value == datetime(2026, 1, 12)
    data["items"][0]["data_type"] = "narrative"
    data["fields"][0]["value"] = "x" * 32768
    assert "excel_value_length" in {issue["code"] for issue in validate_snapshot(data)["issues"]}
    with pytest.raises(ValueError, match="blocked"):
        export_workbook(data, draft=False)


def test_optional_unreviewed_is_blank_final_and_status_retained():
    data = snapshot()
    data["items"][0]["required"] = False
    data["fields"] = []
    validation = validate_snapshot(data)
    assert validation["can_finalize"] and validation["counts"]["unreviewed"] == 1
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert book.worksheets[0]["A2"].value is None
    assert "Not yet reviewed" in str(list(book["Review status"].values))


@pytest.mark.parametrize("status", ["unreviewed", "unclear", "conflicting"])
def test_clean_preview_hides_working_values_even_with_accepted_passages(status):
    data = snapshot()
    data["fields"][0].update(status=status, value=999)
    before = deepcopy(data)
    preview = export_preview(data, mark_unfinished=False)
    assert preview["rows"][0]["values"] == [None]
    assert preview["rows"][0]["statuses"] == [status]
    assert data == before
    # Downloadable drafts still retain their unfinished-field labels.
    assert export_preview(data)["rows"][0]["values"][0].startswith("[DRAFT:")


def test_clean_preview_keeps_saved_zero_and_hides_withdrawn_approval():
    data = snapshot()
    assert export_preview(data, mark_unfinished=False)["rows"][0]["values"] == [0]
    data["evidence"][0]["decision"] = "rejected"
    assert export_preview(data, mark_unfinished=False)["rows"][0]["values"] == [None]


def test_attention_count_counts_fields_once_when_multiple_checks_fail():
    data = snapshot()
    data["fields"][0].update(needs_reconsideration=True, evidence_ids=[], value=None)
    result = validate_snapshot(data)
    assert result["counts"]["errors"] >= 3
    assert result["counts"]["attention_fields"] == 1
    data["fields"] = []
    assert validate_snapshot(data)["counts"]["attention_fields"] == 1


def test_verification_workbook_copies_design_and_links_saved_zero_to_all_sources():
    data = snapshot()
    quote = data["evidence"][0]["quote"]
    text = "Heading\n\n" + quote + "\nOther source\nTreatment: 0 events."
    data["pages"][0]["text"] = text
    data["evidence"][0].update(start=text.index(quote), end=text.index(quote) + len(quote))
    second = {**data["evidence"][0], "id": "evidence-2", "quote": "Treatment: 0 events.",
              "start": text.index("Treatment: 0 events."), "end": len(text)}
    data["evidence"].append(second)
    data["fields"][0]["evidence_ids"].append(second["id"])
    data["items"].append({**data["items"][0], "id": "item-2", "name": "Unfinished", "export_heading": "Unfinished", "required": False, "position": -1})
    before = deepcopy(data)
    values = load_workbook(BytesIO(export_workbook(data, draft=False)))
    locations = load_workbook(BytesIO(export_workbook(data, draft=False, verification=True)))
    assert locations.sheetnames == values.sheetnames
    assert list(locations.worksheets[0].values)[0] == list(values.worksheets[0].values)[0]
    assert values.worksheets[0]["B2"].value == 0
    assert locations.worksheets[0]["A2"].value is None
    assert locations.worksheets[0]["B2"].value == "Page 1, line 3\nPage 1, line 5"
    assert locations.worksheets[0]["B2"].data_type == 's'
    for sheet in values.sheetnames[1:]:
        if sheet != 'Export metadata':
            assert list(locations[sheet].values) == list(values[sheet].values)
    assert "extracted page text" in str(list(locations['Export metadata'].values))
    assert data == before


def test_verification_multiline_manual_and_unaccepted_values():
    data = snapshot()
    quote = "Treatment:\n0 adverse events.\n"
    data["pages"][0]["text"] = "Heading\n" + quote
    data["evidence"][0].update(quote=quote, start=len("Heading\n"), end=len("Heading\n") + len(quote))
    assert verification_preview(data)["rows"][0]["values"] == ["Page 1, lines 2-3"]
    data["pages"][0]["status"] = "ocr_required"
    data["papers"][0]["manual_review_ack"] = True
    data["evidence"][0].update(manual=True, rationale="Manually transcribed from the image.")
    assert verification_preview(data)["rows"][0]["values"] == ["Page 1, selected region x=10-150, y=20-32 PDF units from top-left, manual transcription; no reliable line number"]
    data["fields"][0]["status"] = "unreviewed"
    assert verification_preview(data)["rows"][0]["values"] == [None]
    data["fields"][0]["status"] = "approved"
    data["evidence"][0]["decision"] = "rejected"
    assert verification_preview(data)["rows"][0]["values"] == [None]


def test_verification_date_locations_remain_text():
    data = snapshot()
    data["items"][0]["data_type"] = 'date'
    data["fields"][0]["value"] = '2026-01-12'
    book = load_workbook(BytesIO(export_workbook(data, draft=False, verification=True)))
    assert book.worksheets[0]['A2'].value == 'Page 1, line 1'
    assert book.worksheets[0]['A2'].number_format == '@'


def test_verification_image_uses_specific_descriptor_and_never_invents_a_line():
    data = snapshot()
    data['pages'][0].update(status='ocr_required', text='')
    data['papers'][0]['manual_review_ack'] = True
    data['evidence'][0].update(manual=True, rationale='Read from the plotted image.', start=None, end=None,
                               source_location='within Figure 1, panel B, lower-right outcome plot', boxes=[])
    location = verification_preview(data)['rows'][0]['values'][0]
    assert 'Page 1, within Figure 1, panel B, lower-right outcome plot' in location
    assert 'no reliable line number' in location and 'line 1' not in location
    data['evidence'][0]['source_location'] = ''
    assert 'precise region not recorded' in verification_preview(data)['rows'][0]['values'][0]


def test_verification_ocr_table_uses_row_column_and_region_not_internal_table_number():
    data = snapshot()
    data['evidence'][0].update(ocr=True, table_id='internal-table-42', row_header='Treatment', column_header='Week 6')
    location = verification_preview(data)['rows'][0]['values'][0]
    assert 'table region, row: Treatment, column: Week 6' in location
    assert 'selected region x=10-150, y=20-32' in location
    assert 'internal-table-42' not in location and 'line 1' not in location


def test_verification_supports_manual_image_on_a_searchable_page():
    data = snapshot()
    data['evidence'][0].update(quote='Image-only result: 0.', manual=True, start=None, end=None,
                              rationale='Transcribed from the image.', source_location='within Figure 1, panel C')
    assert validate_snapshot(data)['can_finalize']
    assert 'within Figure 1, panel C' in verification_preview(data)['rows'][0]['values'][0]
    data['evidence'][0]['source_location'] = ''
    assert not validate_snapshot(data)['can_finalize']


def test_long_audit_and_exact_source_quotes_split_with_reconstructable_parts():
    data = snapshot()
    quote = "Exact synthetic source; " * 2500
    data["evidence"][0].update(quote=quote, manual=True, rationale="Transcribed from image-only physical page 1.")
    data["pages"][0]["status"] = "ocr_required"
    data["papers"][0]["manual_review_ack"] = True
    data["history"][0]["before"] = {"notes": "long historical note " * 2500}
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    rows = list(book["Evidence"].values)
    columns = {heading: index for index, heading in enumerate(rows[0])}
    assert rows[1][columns["Evidence ID"]] == rows[2][columns["Evidence ID"]] == "evidence-1"
    assert "".join(row[columns["Exact source quote"]] or "" for row in rows[1:]) == quote
    assert rows[1][columns["Manual transcription rationale"]] == data["evidence"][0]["rationale"]
    history = list(book["Decision history"].values)
    hcols = {heading: index for index, heading in enumerate(history[0])}
    before = json.loads("".join(row[hcols["Before"]] or "" for row in history[1:]))
    assert before == data["history"][0]["before"]


def test_review_scope_protects_main_and_supporting_sheets():
    data = snapshot()
    data["items"].append({"id": "other-item", "review_id": "other-review", "name": "Private item"})
    data["papers"].append({"id": "other-paper", "review_id": "other-review", "title": "PRIVATE_MARKER"})
    data["evidence"].append({"id": "other-evidence", "paper_id": "other-paper", "item_id": "other-item", "quote": "PRIVATE_MARKER"})
    data["fields"].append({"paper_id": "other-paper", "item_id": "other-item", "value": "PRIVATE_MARKER"})
    data["history"].append({"id": "other-history", "review_id": "other-review", "action": "PRIVATE_MARKER"})
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert "PRIVATE_MARKER" not in str([[list(row) for row in sheet.values] for sheet in book])


def test_invalid_date_nan_and_empty_value_are_not_approved():
    data = snapshot()
    for kind, value in [("date", "20260112"), ("number", float("nan")), ("number", float("inf")), ("integer", True), ("narrative", "  ")]:
        data["items"][0]["data_type"] = kind
        data["fields"][0]["value"] = value
        assert "invalid_value" in {issue["code"] for issue in validate_snapshot(data)["issues"]}


def test_validation_matches_service_source_and_observation_requirements():
    data = snapshot()
    data["pages"][0]["text"] = "Changed source content."
    assert "source_mismatch" in {issue["code"] for issue in validate_snapshot(data)["issues"]}
    data = snapshot()
    data["fields"][0]["observations"][0].update(value=None, evidence_ids=[], group="", time_point="")
    assert {"invalid_observation_value", "unlabeled_observation", "missing_observation_evidence"}.issubset({issue["code"] for issue in validate_snapshot(data)["issues"]})
    data = snapshot()
    data["items"][0]["allowed_values"] = [False]
    assert "invalid_value" in {issue["code"] for issue in validate_snapshot(data)["issues"]}
