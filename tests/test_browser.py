"""Real-browser checks. Opt in with RUN_BROWSER_TESTS=1; no browser at runtime."""

import os
from pathlib import Path
import socket
import threading
import time

import httpx
import pytest
import uvicorn

from review_extractor.app import create_app

pytestmark = [pytest.mark.browser, pytest.mark.skipif(
    os.environ.get("RUN_BROWSER_TESTS") != "1", reason="Set RUN_BROWSER_TESTS=1 for real-browser verification")]

DEMO = Path(__file__).resolve().parents[1] / "demo"


@pytest.fixture
def live_app(tmp_path):
    app = create_app(tmp_path / "data", run_background=False)
    with socket.socket() as reservation:
        reservation.bind(("127.0.0.1", 0))
        port = reservation.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="error"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    url = f"http://127.0.0.1:{port}"
    try:
        with httpx.Client(base_url=url, timeout=20) as api:
            for _ in range(150):
                try:
                    if api.get("/api/health").status_code == 200:
                        break
                except httpx.TransportError:
                    time.sleep(0.05)
            else:
                raise AssertionError("Local test server did not start")
            yield url, api, app.state.service
    finally:
        server.should_exit = True
        thread.join(timeout=15)


@pytest.fixture
def browser_page():
    playwright = pytest.importorskip("playwright.sync_api")
    with playwright.sync_playwright() as driver:
        browser = driver.chromium.launch(channel=os.environ.get("PLAYWRIGHT_CHANNEL") or None)
        page = browser.new_page(viewport={"width": 1600, "height": 1050}, device_scale_factor=1)
        page.set_default_timeout(15000)
        yield page
        browser.close()


def seed(api, *, geometry=False):
    r = api.post("/api/reviews", json={"name": "Browser verification", "reviewer": "Browser reviewer"}).json()["id"]
    item = api.post(f"/api/reviews/{r}/items", json={"name": "Frequency", "data_type": "text", "keywords": ["20 Hz"]}).json()
    filename = "geometry-cropped-rotated.pdf" if geometry else "study-A.pdf"
    p = api.post(f"/api/reviews/{r}/papers", files={"files": (filename, (DEMO / filename).read_bytes(), "application/pdf")}).json()["papers"][0]["id"]
    api.post(f"/api/reviews/{r}/refresh")
    api.patch(f"/api/reviews/{r}", json={"settings": {"last_paper_id": p, "last_item_id": item["id"]}})
    return r, p, item["id"]


def test_browser_shell_is_local_and_has_three_sections(live_app, browser_page):
    url, api, _service = live_app
    seed(api)
    requests, errors = [], []
    browser_page.on("request", lambda request: requests.append(request.url))
    browser_page.on("pageerror", lambda error: errors.append(str(error)))
    browser_page.goto(url)
    browser_page.wait_for_load_state("networkidle")
    assert "Review" in browser_page.title()
    assert browser_page.get_by_text("Browser verification", exact=True).count() > 0
    assert browser_page.locator("[data-section]").count() == 3
    assert all(address.startswith(url) or address.startswith(("data:", "blob:")) for address in requests)
    assert not errors


def test_accepted_sources_first_and_approve_value_carries_locations(live_app, browser_page):
    from playwright.sync_api import expect
    from io import BytesIO
    from openpyxl import load_workbook
    from test_retrieval import page as source_page
    url, api, service = live_app
    r, p, i = seed(api)
    service.set_pages(r, p, [source_page('Methods\nStimulation was 20 Hz.\n\nResults\nParticipants received 20 Hz.')])
    service.refresh_search(r, p)
    suggestions = service.get_suggestions(r, p, i)
    target = next(s for s in reversed(suggestions) if s.get('page_number'))
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    card = page.locator(f'[data-suggestion-id="{target["id"]}"]')
    card.get_by_role('button', name='Accept passage', exact=True).click()
    expect(page.locator('#suggestions .suggestion-card').first).to_have_attribute('data-suggestion-id', target['id'])
    expect(page.locator('#suggestions .suggestion-card').first).to_have_class('suggestion-card accepted')
    page.locator('#value-tab').click()
    expect(page.locator('#evidence-links blockquote')).to_have_text(target['quote'])
    expect(page.locator('#evidence-links input')).to_be_checked()
    location = service.evidence_with_locations(r, p, i)[0]['verification_location']
    expect(page.locator('#evidence-links .hint')).to_have_text(location)
    assert 'line' in location
    assert page.locator('#evidence-links').bounding_box()['y'] < page.locator('#field-value').bounding_box()['y']
    expect(page.locator('#field-value')).to_have_value('')
    page.locator('#field-value').fill('20 Hz')
    page.locator('#save-field').click()
    expect(page.locator('#field-progress')).to_contain_text('Approved value')
    assert service.get_field(r, p, i)['value'] == '20 Hz'
    exported = api.get(f'/api/reviews/{r}/export?view=combined')
    book = load_workbook(BytesIO(exported.content))
    assert book['Data']['A2'].value == '20 Hz'
    assert book['Verification']['A2'].value == location
    # Reopening preserves the saved value and linked source selection.
    page.reload()
    page.locator('[data-section="extract"]').click()
    page.locator('#value-tab').click()
    expect(page.locator('#field-value')).to_have_value('20 Hz')
    expect(page.locator('#evidence-links input')).to_be_checked()


def test_multiworksheet_import_preview_and_six_tab_combined_download(live_app, browser_page, tmp_path):
    from openpyxl import Workbook, load_workbook
    from io import BytesIO
    from playwright.sync_api import expect
    url, api, service = live_app
    r = api.post('/api/reviews', json={'name': 'Three-sheet review'}).json()['id']
    template = Workbook()
    template.active.title = 'Study details'
    names = ['Study details', 'Outcomes', 'Risk of bias']
    for name in names:
        sheet = template[name] if name in template.sheetnames else template.create_sheet(name)
        sheet.append(['Result'])
        sheet.append(['EXAMPLE ONLY'])
    path = tmp_path / 'three-sheets.xlsx'; template.save(path)
    page = browser_page
    page.goto(url)
    page.locator('#template-file').set_input_files(path)
    page.locator('#template-form button[type=submit]').click()
    expect(page.locator('#template-items tr')).to_have_count(3)
    assert page.locator('#template-items tr td:nth-child(2)').all_text_contents() == names
    page.locator('#import-items').click()
    expect(page.locator('#template-preview')).to_be_hidden()
    page.locator('[data-section="export"]').click()
    expect(page.locator('#export-table [role=tab]')).to_have_count(3)
    page.get_by_role('tab', name='Outcomes', exact=True).click()
    expect(page.get_by_role('tab', name='Outcomes', exact=True)).to_have_attribute('aria-selected', 'true')
    with page.expect_download() as event:
        page.locator('#combined-export').click()
    book = load_workbook(BytesIO(event.value.path().read_bytes()))
    assert book.sheetnames == names + [f'Verify - {name}' for name in names]
    assert service.export_snapshot(r)['fields'] == []


def test_export_preview_has_blanks_and_one_attention_count(live_app, browser_page):
    url, api, service = live_app
    r, p, i = seed(api)
    suggestion = next(row for row in service.get_suggestions(r, p, i) if row.get('page_number'))
    service.decide_suggestion(r, p, suggestion['id'], {'decision': 'accepted'})
    page = browser_page
    page.goto(url)
    page.locator('[data-section="export"]').click()
    page.locator('#validation-summary h2').wait_for()
    assert page.locator('#validation-summary h2').inner_text() == '1 item needs attention'
    assert page.locator('#validation-summary button, #validation-summary .validation-issue').count() == 0
    assert page.locator('#export-table tbody td').inner_text() == suggestion['quote']
    # Accepted passages are visible without changing final value approval.
    assert service.get_field(r, p, i)['status'] == 'unreviewed'
    evidence = service.get_evidence(r, p, i)[0]
    service.save_field(r, p, i, {'status': 'approved', 'value': '20 Hz', 'evidence_ids': [evidence['id']]})
    page.locator('#refresh-validation').click()
    from playwright.sync_api import expect
    expect(page.locator('#export-table tbody td')).to_have_text('20 Hz')
    expect(page.locator('#validation-summary h2')).to_have_text('0 items need attention')
    assert page.locator('#data-export').is_enabled()
    from openpyxl import load_workbook
    from io import BytesIO
    for view in ['data', 'verification', 'combined']:
        with page.expect_download() as download_event:
            page.locator(f'#{view}-export').click()
        download = download_event.value
        assert download.suggested_filename.endswith(f'-{view}.xlsx')
        book = load_workbook(BytesIO(download.path().read_bytes()))
        if view != 'verification':
            assert book['Data']['A2'].value == '20 Hz'
        if view != 'data':
            assert book['Verification']['A2'].value.startswith('Page 1, line')
    # Refresh retrieves current decisions, with no candidate/history requests.
    requests = []
    page.on('request', lambda request: requests.append(request.url))
    service.update_evidence(r, p, evidence['id'], {'decision': 'rejected'})
    page.locator('#refresh-validation').click()
    expect(page.locator('#export-table tbody td')).to_have_text('')
    assert not any('/history' in address or '/suggestions' in address for address in requests)


def test_hybrid_ranking_debug_top_ten_and_feedback(live_app, browser_page):
    from test_retrieval import page as source_page
    url, api, service = live_app
    r,p,i = seed(api)
    api.patch(f'/api/reviews/{r}/items/{i}',json={'name':'Stimulation frequency','search_mode':'fast'})
    other_item = api.post(f'/api/reviews/{r}/items', json={
        'name':'Secondary stimulation frequency','data_type':'text','keywords':['Hz'],'search_mode':'fast'
    }).json()['id']
    text = 'Methods\n' + '\n\n'.join(f'rTMS was delivered at {n} Hz in protocol {n}.' for n in range(1,13))
    service.set_pages(r,p,[source_page(text)])
    service.refresh_search(r,p)
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#suggestions .suggestion-card').first.wait_for()
    assert page.locator('#suggestions .suggestion-card').count() == 10
    card=page.locator('#suggestions .suggestion-card').first
    assert 'Relevance' in card.inner_text()
    card.get_by_text('Why this ranked here',exact=True).click()
    assert 'keyword_score' in card.inner_text()
    card.get_by_text('Reviewer feedback influence',exact=True).click()
    assert 'Feedback not measured' in card.inner_text()
    card.get_by_role('button',name='Accept passage',exact=True).click()
    page.wait_for_function("() => document.querySelector('#save-status').textContent.includes('Evidence decision saved')")
    assert service.snapshot(r)['feedback'][0]['accepted']
    assert service.get_field(r,p,i)['status']=='unreviewed'
    page.get_by_role('button',name='Show all 12 candidates',exact=True).click()
    assert page.locator('#suggestions .suggestion-card').count()==12
    page.get_by_role('button',name='Show top 10, user candidates and saved decisions',exact=True).click()
    assert page.locator('#suggestions .suggestion-card').count()==10
    page.get_by_role('button',name='Show all 12 candidates',exact=True).click()
    page.locator('#item-select').select_option(other_item)
    page.wait_for_function("() => document.querySelector('#current-item-name').textContent === 'Secondary stimulation frequency'")
    assert page.locator('#suggestions .suggestion-card').count()==10
    page.get_by_role('button',name='Show all 12 candidates',exact=True).click()
    page.locator('#item-select').select_option(i)
    page.wait_for_function("() => document.querySelector('#current-item-name').textContent === 'Stimulation frequency'")
    assert page.locator('#suggestions .suggestion-card').count()==10
    page.locator('#edit-items').click()
    page.locator('#item-form details').first.locator('summary').first.click()
    page.locator('#edit-search-mode').select_option('fast')
    page.locator('#edit-retrieval').fill('{"negative_examples":["EMG sampling"],"expected_units":["Hz"]}')
    page.locator('#save-item').click()
    page.wait_for_function("() => document.querySelector('#item-save-status').textContent.includes('Saved for all papers')")
    saved=service.list_items(r)[0]
    assert saved['negative_examples']==['EMG sampling'] and saved['search_mode']=='fast'


def test_age_ranking_refresh_decision_and_export(live_app, browser_page):
    from test_retrieval import page as source_page
    from ranking_quality_fixture import AGE_DISTRACTORS
    url, api, service = live_app
    r,p,i = seed(api)
    api.patch(f'/api/reviews/{r}/items/{i}',json={'name':'Mean age','search_mode':'fast','keywords':['age','mean age','median age']})
    service.set_pages(r,p,[source_page('\n\n'.join(AGE_DISTRACTORS) + '\n\nRecruited adults were\naged 44.3 ± 8.1 years.\n\nMean age was 52 years in the control group.')])
    service.refresh_search(r,p)
    browser_page.goto(url)
    browser_page.locator('[data-section="extract"]').click()
    cards = browser_page.locator('#suggestions .suggestion-card')
    cards.first.wait_for()
    # Both reported summaries lead discussion/eligibility; wrapped variants do
    # not consume two positions. Acceptance remains separate from value approval.
    assert 'years' in cards.first.inner_text()
    assert sum('44.3' in text for text in cards.all_inner_texts()) == 1
    card = cards.filter(has_text='44.3').first
    card.get_by_role('button',name='Accept passage',exact=True).click()
    browser_page.wait_for_function("() => document.querySelector('#save-status').textContent.includes('Evidence decision saved')")
    accepted = next(s for s in service.get_suggestions(r,p,i) if s['decision']=='accepted')
    assert service.get_field(r,p,i)['status']=='unreviewed'
    evidence = next(e for e in service.get_evidence(r,p,i) if e['suggestion_id']==accepted['id'])
    approved = service.save_field(r,p,i,{'status':'approved','value':'44.3','evidence_ids':[evidence['id']]})
    service.refresh_search(r,p)
    browser_page.reload()
    browser_page.locator('[data-section="extract"]').click()
    assert service.get_field(r,p,i)==approved
    assert next(s for s in service.get_suggestions(r,p,i) if s['id']==accepted['id'])['decision']=='accepted'
    from io import BytesIO
    from openpyxl import load_workbook
    from review_extractor.workbooks import export_workbook
    book=load_workbook(BytesIO(export_workbook(service.snapshot(r))))
    assert any('44.3' in [str(c.value) for c in row] for sheet in book for row in sheet)
    book.close()


def test_remove_button_cancel_confirm_and_last_paper(live_app, browser_page):
    url, api, _service = live_app
    r, p, i = seed(api)
    page = browser_page
    page.goto(url)
    card = page.locator('#paper-list .paper-card')
    card.get_by_role('button', name='Open paper').wait_for()
    remove = card.get_by_role('button', name='Remove', exact=True)
    assert remove.is_visible()
    page.once('dialog', lambda dialog: dialog.dismiss())
    remove.click()
    assert card.count() == 1
    page.once('dialog', lambda dialog: dialog.accept())
    remove.click()
    page.wait_for_function("() => document.querySelectorAll('#paper-list .paper-card').length === 0")
    page.wait_for_function("() => document.querySelector('#save-status').textContent.includes('Paper removed')")
    assert api.get(f'/api/reviews/{r}/papers').json() == []
    assert api.get(f'/api/reviews/{r}').json()['settings']['last_paper_id'] is None
    page.locator('[data-section="extract"]').click()
    assert page.locator('#extract-empty').is_visible()
    assert page.locator('#reading-workspace').is_hidden()
    page.reload()
    assert page.locator('#paper-list .paper-card').count() == 0


def test_automatic_keywords_remain_editable_and_saved(live_app, browser_page):
    url, api, _service = live_app
    r, _p, _i = seed(api)
    page = browser_page
    page.goto(url)
    page.locator('#setup-edit-items').click()
    page.locator('#new-item').click()
    page.locator('#edit-name').fill('Mean age')
    page.wait_for_function("() => document.querySelector('#edit-keywords').value.includes('median age')")
    page.locator('#edit-keywords').fill('mean age\nmy custom phrase')
    page.locator('#edit-definition').fill('Participant age at enrollment')
    page.locator('#save-item').click()
    page.wait_for_function("() => document.querySelector('#item-save-status').textContent.includes('Saved')")
    assert page.locator('#edit-keywords').input_value() == 'mean age\nmy custom phrase'
    page.reload()
    page.locator('#setup-edit-items').click()
    age_item = next(item for item in api.get(f'/api/reviews/{r}/items').json() if item['name'] == 'Mean age')
    page.locator(f'#editor-item-list [data-item-id="{age_item["id"]}"]').click()
    assert page.locator('#edit-keywords').input_value() == 'mean age\nmy custom phrase'
    page.locator('#suggest-keywords').click()
    page.wait_for_function("() => document.querySelector('#edit-keywords').value.includes('median age')")
    terms = page.locator('#edit-keywords').input_value().splitlines()
    assert terms.count('mean age') == 1 and 'my custom phrase' in terms
    page.locator('#edit-keywords').fill('')
    page.locator('#save-item').click()
    page.wait_for_function("() => document.querySelector('#item-save-status').textContent.includes('Saved')")
    item = next(item for item in api.get(f'/api/reviews/{r}/items').json() if item['name'] == 'Mean age')
    assert item['keywords'] == []
    page.locator('#new-item').click()
    page.locator('#edit-name').fill('Frequency')
    # Manual typing before the debounce expires must cancel automatic filling.
    page.locator('#edit-keywords').fill('user term only')
    page.locator('#save-item').click()
    page.wait_for_function("() => document.querySelector('#item-save-status').textContent.includes('Saved')")
    assert page.locator('#edit-keywords').input_value() == 'user term only'


def test_concept_candidates_explain_sources_and_keep_metadata_separate(live_app, browser_page):
    from retrieval_fixture import concept_pdf
    url, api, _service = live_app
    r = api.post('/api/reviews', json={'name': 'Concept browser test'}).json()['id']
    i = api.post(f'/api/reviews/{r}/items', json={'name': 'Authors', 'keywords': ['authors']}).json()['id']
    p = api.post(f'/api/reviews/{r}/papers', files={'files': ('concepts.pdf', concept_pdf(), 'application/pdf')}).json()['papers'][0]['id']
    api.post(f'/api/reviews/{r}/refresh')
    page = browser_page
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    byline = page.locator('.suggestion-card').filter(has=page.locator('blockquote', has_text='Alice Example, Bob Sample'))
    byline.wait_for()
    assert 'Name-like first-page byline' in byline.inner_text()
    assert 'Visible PDF text' in byline.inner_text()
    metadata = page.locator('.suggestion-card').filter(has=page.locator('blockquote', has_text='Morgan Example'))
    assert 'Document metadata' in metadata.inner_text()
    assert metadata.get_by_role('button', name='Accept passage').count() == 0
    byline.get_by_role('button', name='show source', exact=False).click()
    page.wait_for_function("() => document.querySelectorAll('.pdf-highlight').length > 0")
    assert page.locator('#page-number').input_value() == '1'
    byline.get_by_role('button', name='Accept passage').click()
    page.locator('#saved-evidence blockquote').wait_for()
    assert api.get(f'/api/reviews/{r}/papers/{p}/items/{i}/field').json()['status'] == 'unreviewed'
    folder = Path(__file__).resolve().parents[1] / 'tmp'
    folder.mkdir(exist_ok=True)
    page.screenshot(path=str(folder / 'browser-concept-retrieval.png'), full_page=True)
    assert not errors


def test_create_import_upload_and_add_item_in_reader(live_app, browser_page):
    url, api, _service = live_app
    page = browser_page
    page.goto(url)
    page.locator("#review-name").fill("UI journey")
    page.locator("#reviewer-name").fill("Researcher One")
    page.locator("#create-review-form button[type=submit]").click()
    page.wait_for_function("() => document.querySelector('#review-select').value !== ''")
    r = page.locator("#review-select").input_value()
    page.locator("#reference-template").click()
    page.locator("#template-items tr").last.wait_for()
    assert page.locator("#template-items tr").count() == 29
    page.locator("#import-items").click()
    page.wait_for_function("() => document.querySelector('#design-summary').textContent.includes('29')")
    page.locator("#pdf-files").set_input_files([str(DEMO / "study-A.pdf"), str(DEMO / "study-B.pdf")])
    page.locator("#upload-form button[type=submit]").click()
    page.wait_for_function("() => document.querySelectorAll('#paper-list .paper-card').length === 2")
    assert all(p["processing_status"] == "pending" for p in api.get(f"/api/reviews/{r}/papers").json())
    page.locator('[data-section="extract"]').click()
    page.locator("#edit-items").click()
    page.locator("#new-item").click()
    page.locator("#edit-name").fill("New frequency")
    page.locator("#edit-heading").fill("New frequency")
    page.locator("#edit-keywords").fill("20 Hz\nstimulation")
    page.locator("#save-item").click()
    page.wait_for_function("() => document.querySelector('#item-save-status').textContent.toLowerCase().includes('saved')")
    page.get_by_role("button", name="Close extraction item editor").click()
    page.locator('[data-section="setup"]').click()
    page.locator("#scan-all-papers").click()
    page.wait_for_function("() => [...document.querySelectorAll('#paper-list .pill')].every(el => el.textContent === 'Ready to review')")
    page.locator('[data-section="extract"]').click()
    page.locator("#item-select").select_option(label="New frequency")
    page.locator("#suggestions .suggestion-card").first.wait_for()
    items = api.get(f"/api/reviews/{r}/items").json()
    new_item = next(item for item in items if item["name"] == "New frequency")
    assert len(items) == 30
    for paper in api.get(f"/api/reviews/{r}/papers").json():
        assert api.get(f"/api/reviews/{r}/papers/{paper['id']}/suggestions", params={"item_id": new_item["id"]}).json()
    page.reload()
    page.locator('[data-section="extract"]').click()
    page.locator("#item-select").select_option(label="New frequency")
    page.locator("#suggestions .suggestion-card").first.wait_for()


def test_manual_scan_selected_paper_preserves_unsaved_form(live_app, browser_page):
    url, api, service = live_app
    r, p, i = seed(api)
    other = api.post(f'/api/reviews/{r}/papers', files={'files': (
        'study-B.pdf', (DEMO / 'study-B.pdf').read_bytes(), 'application/pdf')}).json()['papers'][0]['id']
    before = service.get_pages(r, p)
    page = browser_page
    requests = []
    page.on('request', lambda req: requests.append(req.url) if '/refresh' in req.url else None)
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#refresh-search').wait_for()
    assert not requests
    # A scan updates suggestions without throwing away an unfinished value.
    page.locator('#value-tab').click()
    page.locator('#field-value').fill('Unfinished reviewer value')
    page.locator('#refresh-search').click()
    page.wait_for_function("() => document.querySelector('#scan-status').textContent.includes('1 paper(s)')")
    page.wait_for_function("() => !document.querySelector('#refresh-search').disabled")
    assert len(requests) == 1 and f'paper_id={p}' in requests[0]
    assert page.locator('#field-value').input_value() == 'Unfinished reviewer value'
    assert service.get_pages(r, p) == before
    assert service.get_paper(r, other)['processing_status'] == 'pending'


def test_partial_scan_notice_names_item_page_and_rule(live_app, browser_page, monkeypatch):
    from review_extractor import service as service_module
    url, api, service = live_app
    r, p, i = seed(api)
    original = service_module.search_item

    def incomplete(*args, **kwargs):
        result = original(*args, **kwargs)
        result['errors'].append({'page_number': 1, 'rule': 'custom <pattern>',
            'code': 'regex_timeout', 'error': 'Custom pattern search timed out.'})
        return result

    monkeypatch.setattr(service_module, 'search_item', incomplete)
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#refresh-search').click()
    page.wait_for_function("() => document.querySelector('#processing-notice').textContent.includes('some searches incomplete')")
    notice = page.locator('#processing-notice').inner_text()
    assert 'Frequency · page 1 · “custom <pattern>”' in notice
    assert 'does not mean evidence is absent' in notice
    assert i not in notice  # Show item names, not UUIDs or raw error JSON.
    assert page.locator('#processing-notice pattern').count() == 0
    assert page.locator('#suggestions .suggestion-card').count() > 0
    page.reload()
    page.locator('[data-section="extract"]').click()
    page.wait_for_function("() => document.querySelector('#processing-notice').textContent.includes('some searches incomplete')")
    assert service.get_paper(r, p)['search_status'] == 'partial'
    monkeypatch.setattr(service_module, 'search_item', original)
    page.locator('#refresh-search').click()
    page.wait_for_function("() => !document.querySelector('#refresh-search').disabled && !document.querySelector('#processing-notice').textContent.includes('some searches incomplete')")
    assert service.get_paper(r, p)['search_status'] == 'current'


def test_pdf_canvas_text_selection_and_zoom_coordinates(live_app, browser_page):
    url, api, _service = live_app
    r, p, i = seed(api)
    page = browser_page
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(url)
    page.locator("#review-select").select_option(r)
    page.locator('[data-section="extract"]').click()
    page.locator("#suggestions .source-button").first.click()
    page.wait_for_function("() => document.querySelectorAll('#pdf-text-layer span').length > 0 && document.querySelectorAll('.pdf-highlight').length > 0")
    assert page.locator("#page-number").input_value() == "1"
    def geometry():
        return page.evaluate("""() => {
          const stage=document.querySelector('#pdf-stage').getBoundingClientRect();
          const box=document.querySelector('.pdf-highlight').getBoundingClientRect();
          return {width:stage.width,height:stage.height,x:(box.x-stage.x)/stage.width,
                  y:(box.y-stage.y)/stage.height,w:box.width/stage.width,h:box.height/stage.height};
        }""")
    initial = geometry()
    assert initial["w"] > 0 and initial["h"] > 0
    page.locator("#zoom-in").click()
    page.wait_for_function("previous => document.querySelector('#pdf-stage').getBoundingClientRect().width > previous", arg=initial["width"])
    page.wait_for_function("() => document.querySelectorAll('#pdf-text-layer span').length > 0 && document.querySelectorAll('.pdf-highlight').length > 0")
    zoomed = geometry()
    for key in ("x", "y", "w", "h"):
        assert zoomed[key] == pytest.approx(initial[key], abs=0.002)
    # Select actual PDF.js text and pass it through the researcher control.
    selected = page.evaluate("""() => {
      const span=[...document.querySelectorAll('#pdf-text-layer span')].find(e=>e.textContent.includes('20 Hz'));
      if(!span) throw new Error('Missing selectable source text');
      const selection=getSelection(); selection.removeAllRanges();
      const range=document.createRange(); range.selectNodeContents(span);selection.addRange(range);
      return selection.toString();
    }""")
    assert "20 Hz" in selected
    page.locator("#capture-selection").click()
    page.locator("#evidence-dialog").wait_for(state="visible")
    assert "20 Hz" in page.locator("#evidence-quote").input_value()
    page.locator("#accept-manual-evidence").click()
    page.locator("#evidence-dialog").wait_for(state="hidden")
    assert api.get(f"/api/reviews/{r}/papers/{p}/items/{i}/evidence").json()
    page.locator("#value-tab").click()
    page.locator("#field-status").select_option("approved")
    page.locator("#field-value").fill("20 Hz")
    for checkbox in page.locator("#evidence-links input[type=checkbox]").all():
        checkbox.check()
    page.locator("#save-field").click()
    page.wait_for_function("() => document.querySelector('#field-save-status').textContent.toLowerCase().includes('saved')")
    assert api.get(f"/api/reviews/{r}/papers/{p}/items/{i}/field").json()["value"] == "20 Hz"
    folder = Path(__file__).resolve().parents[1] / "tmp"
    folder.mkdir(exist_ok=True)
    page.screenshot(path=str(folder / "browser-reading.png"), full_page=True)
    page.locator('[data-section="export"]').click()
    page.locator("#export-table td").first.wait_for()
    from playwright.sync_api import expect
    expect(page.locator("#combined-export")).to_be_enabled()
    with page.expect_download() as export:
        page.locator("#combined-export").click()
    from openpyxl import load_workbook
    from io import BytesIO
    book = load_workbook(BytesIO(Path(export.value.path()).read_bytes()))
    assert book.worksheets[0]["A1"].value == "Frequency"
    assert book.worksheets[0]["A2"].value == "20 Hz"
    assert book.sheetnames == ['Data', 'Verification']
    assert 'Page 1' in book['Verification']['A2'].value
    page.screenshot(path=str(folder / "browser-export.png"), full_page=True)
    assert not errors


def test_cropped_rotated_pdf_suggestions_navigate_and_align(live_app, browser_page):
    url, api, _service = live_app
    r, p, i = seed(api, geometry=True)
    source_pages = api.get(f"/api/reviews/{r}/papers/{p}/pages").json()
    suggestions = api.get(f"/api/reviews/{r}/papers/{p}/suggestions", params={"item_id": i}).json()
    assert len(suggestions) == 4
    page = browser_page
    page.goto(url)
    page.locator("#review-select").select_option(r)
    page.locator('[data-section="extract"]').click()
    for index, (source, suggestion) in enumerate(zip(source_pages, suggestions)):
        page.locator("#suggestions .source-button").nth(index).click()
        page.wait_for_function("n => document.querySelector('#page-number').value === String(n)", arg=index + 1)
        expected_ratio = source["width"] / source["height"]
        page.wait_for_function("ratio => { const s=document.querySelector('#pdf-stage').getBoundingClientRect();return Math.abs(s.width/s.height-ratio)<0.01 && document.querySelectorAll('.pdf-highlight').length > 0; }", arg=expected_ratio)
        actual = page.evaluate("""() => {
          const stage=document.querySelector('#pdf-stage').getBoundingClientRect();
          const box=document.querySelector('.pdf-highlight').getBoundingClientRect();
          return [(box.x-stage.x)/stage.width,(box.y-stage.y)/stage.height,box.width/stage.width,box.height/stage.height];
        }""")
        box = suggestion["boxes"][0]
        expected = [box["x0"] / source["width"], box["top"] / source["height"],
                    (box["x1"] - box["x0"]) / source["width"], (box["bottom"] - box["top"]) / source["height"]]
        assert actual == pytest.approx(expected, abs=0.003)


def test_math_characters_do_not_shift_manual_evidence_offsets(live_app, browser_page):
    url, _api, _service = live_app
    browser_page.goto(url)
    result = browser_page.evaluate("""async () => {
      const {DocumentViewer}=await import('/static/viewer.js');
      const viewer=new DocumentViewer(()=>{}, error=>{throw error;});
      viewer.pages=[{page_number:1,text:'𝑥 20 Hz',words:[]}];
      viewer.renderExtracted();
      const node=document.querySelector('#extracted-text').firstChild;
      const range=document.createRange();range.setStart(node,3);range.setEnd(node,8);
      const selected=getSelection();selected.removeAllRanges();selected.addRange(range);
      return viewer.selection();
    }""")
    # Browser strings count the mathematical character as two UTF-16 units;
    # Python source offsets must count it as one Unicode code point.
    assert result["quote"] == "20 Hz"
    assert (result["start"], result["end"]) == (2, 7)


def select_extracted_quote(browser_page, quote, occurrence=0):
    """Select across text/mark nodes, preserving a particular repeated occurrence."""
    return browser_page.evaluate("""({quote,occurrence}) => {
      const root=document.querySelector('#extracted-text');
      let start=-1;
      for(let n=0;n<=occurrence;n++)start=root.textContent.indexOf(quote,start+1);
      if(start<0)throw new Error('Missing requested source occurrence');
      const end=start+quote.length,walker=document.createTreeWalker(root,NodeFilter.SHOW_TEXT);
      let node,offset=0,startNode,endNode,startOffset,endOffset;
      while(node=walker.nextNode()) {
        const next=offset+node.textContent.length;
        if(!startNode && start>=offset && start<next){startNode=node;startOffset=start-offset;}
        if(end>offset && end<=next){endNode=node;endOffset=end-offset;break;}
        offset=next;
      }
      const range=document.createRange();range.setStart(startNode,startOffset);range.setEnd(endNode,endOffset);
      const selected=getSelection();selected.removeAllRanges();selected.addRange(range);
      return {quote:selected.toString(),start,end};
    }""", {"quote": quote, "occurrence": occurrence})


def test_highlighted_candidate_pending_accept_refresh_and_repeated_source(live_app, browser_page):
    from test_retrieval import page as source_page
    from playwright.sync_api import expect
    url, api, service = live_app
    r, p, i = seed(api)
    api.patch(f'/api/reviews/{r}/items/{i}', json={'name': 'Mean age', 'search_mode': 'fast'})
    quote = 'Recruited adults were\naged 41.7 ± 9.2 years.'
    text = f'Results\n{quote}\n\nControl cohort\n{quote}\n\nDiscussion\nAge can affect recovery.'
    service.set_pages(r, p, [source_page(text)])
    service.refresh_search(r, p)
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#suggestions .suggestion-card').first.wait_for()
    page.locator('#text-toggle').click()
    page.locator('#extracted-text').wait_for(state='visible')
    selected = select_extracted_quote(page, quote, occurrence=1)
    page.evaluate("""() => {
      window.selectionRefreshes=0;
      const observer=new MutationObserver(()=>window.selectionRefreshes++);
      observer.observe(document.querySelector('#extracted-text'),{childList:true});
    }""")
    service.update_learning(r, {'enabled': True})
    page.wait_for_function('() => window.selectionRefreshes > 0')
    assert page.evaluate('() => getSelection().toString()') == quote
    page.get_by_role('button', name='Create candidate for this item', exact=True).click()
    expect(page.locator('#evidence-item')).to_contain_text('Mean age')
    assert page.locator('#evidence-quote').input_value() == quote
    assert page.locator('#evidence-quote').get_attribute('readonly') is not None
    artifacts = Path(__file__).resolve().parents[1] / 'tmp'
    artifacts.mkdir(exist_ok=True)
    page.screenshot(path=str(artifacts / 'browser-manual-candidate.png'), full_page=True)
    page.get_by_role('button', name='Save as pending candidate', exact=True).click()
    page.locator('#evidence-dialog').wait_for(state='hidden')
    candidates = service.get_suggestions(r, p, i)
    candidate = next(s for s in candidates if s.get('origin') == 'user' and s['start'] == selected['start'])
    assert candidate['quote'] == quote and candidate['end'] == selected['end']
    assert candidate['decision'] == 'pending' and candidate['boxes']
    assert candidate['document_version'] == service.get_paper(r, p)['document_version']
    assert not service.get_evidence(r, p, i)
    assert service.get_field(r, p, i)['status'] == 'unreviewed'
    identity = candidate['id']
    service.refresh_search(r, p)
    page.reload()
    page.locator('[data-section="extract"]').click()
    card = page.locator(f'[data-suggestion-id="{identity}"]')
    expect(card).to_contain_text('User-created candidate')
    page.locator('#text-toggle').click()
    page.locator('#extracted-text').wait_for(state='visible')
    select_extracted_quote(page, quote, occurrence=1)
    page.locator('#capture-selection').click()
    page.get_by_role('button', name='Accept as evidence', exact=True).click()
    page.locator('#evidence-dialog').wait_for(state='hidden')
    evidence = next(e for e in service.get_evidence(r, p, i) if e['suggestion_id'] == identity)
    assert evidence['decision'] == 'accepted' and evidence['start'] == selected['start']
    assert evidence['quote'] == text[evidence['start']:evidence['end']]
    assert service.get_field(r, p, i)['status'] == 'unreviewed'
    # Creating the same pending candidate again must not undo acceptance or fork an ID.
    select_extracted_quote(page, quote, occurrence=1)
    page.locator('#capture-selection').click()
    page.locator('#save-pending-candidate').click()
    page.locator('#evidence-dialog').wait_for(state='hidden')
    current = [s for s in service.get_suggestions(r, p, i) if s['id'] == identity]
    assert len(current) == 1 and current[0]['decision'] == 'accepted'
    assert len([e for e in service.get_evidence(r, p, i) if e['suggestion_id'] == identity]) == 1


def test_learning_controls_rejection_corrections_and_reconsideration(live_app, browser_page):
    from playwright.sync_api import expect
    url, api, service = live_app
    r, p, i = seed(api)
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    card = page.locator('#suggestions .suggestion-card').first
    card.wait_for()
    identity = card.get_attribute('data-suggestion-id')
    card = page.locator(f'[data-suggestion-id="{identity}"]')
    reason = card.get_by_label('Optional rejection reason')
    reason.select_option('wrong_statistic')
    card.get_by_role('button', name='Reject passage', exact=True).click()
    expect(card).to_have_class('suggestion-card rejected')
    assert next(s for s in service.get_suggestions(r, p, i) if s['id'] == identity)['rejection_reason'] == 'wrong_statistic'
    reason.select_option('irrelevant')
    page.wait_for_function("() => document.querySelector('#save-status').textContent === 'Rejection reason saved'")
    rejected = api.get(f'/api/reviews/{r}/learning').json()['examples']
    assert len(rejected) == 1 and rejected[0]['rejection_reason'] == 'irrelevant'
    example_id = rejected[0]['example_id']
    card.get_by_role('button', name='Reconsider', exact=True).click()
    expect(card).to_have_class('suggestion-card pending')
    pending = api.get(f'/api/reviews/{r}/learning').json()['examples']
    assert len(pending) == 1 and pending[0]['example_id'] == example_id and pending[0]['decision'] == 'pending'
    card.get_by_role('button', name='Accept passage', exact=True).click()
    expect(card).to_have_class('suggestion-card accepted')
    evidence = next(e for e in service.get_evidence(r, p, i) if e['suggestion_id'] == identity)
    page.locator('#learning-settings').click()
    expect(page.locator('#learning-status')).to_contain_text('at least two accepted source documents')
    page.locator('#learning-enabled').check()
    page.wait_for_function("() => document.querySelector('#save-status').textContent === 'Learning setting saved'")
    example = page.locator(f'[data-example-id="{example_id}"]')
    example.get_by_label('Allow this example for learning').uncheck()
    page.wait_for_function("() => document.querySelector('#save-status').textContent === 'Learning example preference saved'")
    expect(example).to_contain_text('Not used:')
    assert api.get(f'/api/reviews/{r}/learning').json()['examples'][0]['excluded']
    artifacts = Path(__file__).resolve().parents[1] / 'tmp'
    artifacts.mkdir(exist_ok=True)
    page.screenshot(path=str(artifacts / 'browser-reviewer-learning.png'), full_page=True)
    before = service.get_field(r, p, i)
    epoch = api.get(f'/api/reviews/{r}/learning').json()['config']['epoch']
    page.locator('#reset-learning').click()
    page.wait_for_function("() => document.querySelector('#save-status').textContent === 'Learned adjustments reset'")
    assert api.get(f'/api/reviews/{r}/learning').json()['config']['epoch'] == epoch + 1
    assert service.get_field(r, p, i) == before
    assert next(e for e in service.get_evidence(r, p, i) if e['id'] == evidence['id'])['decision'] == 'accepted'
    page.get_by_role('button', name='Close learning settings', exact=True).click()
    saved = page.locator(f'[data-evidence-id="{evidence["id"]}"]')
    saved.get_by_role('button', name='Reconsider evidence', exact=True).click()
    expect(saved).to_have_class('evidence-card pending')
    expect(card).to_have_class('suggestion-card pending')


def test_candidate_dialog_rejects_a_changed_context(live_app, browser_page):
    from playwright.sync_api import expect
    url, api, service = live_app
    r, p, i = seed(api)
    other = api.post(f'/api/reviews/{r}/items', json={'name': 'Different field'}).json()['id']
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#suggestions .suggestion-card').first.wait_for()
    page.locator('#text-toggle').click()
    page.locator('#extracted-text').wait_for(state='visible')
    quote = service.get_suggestions(r, p, i)[0]['quote']
    select_extracted_quote(page, quote)
    page.locator('#capture-selection').click()
    # Exercise an asynchronous context change while a captured selection is open.
    page.evaluate("""id => {
      const select=document.querySelector('#item-select');select.value=id;
      select.dispatchEvent(new Event('change',{bubbles:true}));
    }""", other)
    expect(page.locator('#current-item-name')).to_have_text('Different field')
    page.locator('#save-pending-candidate').click()
    expect(page.locator('#toast')).to_contain_text('The paper, item, or document version changed')
    assert not any(s.get('origin') == 'user' for s in service.get_suggestions(r, p, i))
    assert not any(s.get('origin') == 'user' for s in service.get_suggestions(r, p, other))


def test_feedback_explanation_navigates_to_a_different_source_paper(live_app, browser_page, tmp_path):
    from playwright.sync_api import expect
    from review_extractor.hybrid.semantic import LocalSemantic
    from test_feedback_ranker import Encoder
    from test_retrieval import page as source_page
    url, api, service = live_app
    r, first, i = seed(api)
    service.semantic = LocalSemantic(tmp_path / 'synthetic-cache', model='browser-synthetic', encoder=Encoder())
    api.patch(f'/api/reviews/{r}/items/{i}', json={'name': 'Stimulation frequency', 'search_mode': 'intelligent'})
    papers = [first]
    for filename in ('study-B.pdf', 'geometry-cropped-rotated.pdf'):
        response = api.post(f'/api/reviews/{r}/papers', files={'files': (filename, (DEMO / filename).read_bytes(), 'application/pdf')})
        papers.append(response.json()['papers'][0]['id'])
    api.post(f'/api/reviews/{r}/refresh')
    for index, paper_id in enumerate(papers):
        quote = f'rTMS was delivered at {17 + index * 3} Hz for the treatment cohort.'
        service.set_pages(r, paper_id, [source_page(quote)])
        service.refresh_search(r, paper_id)
        if index < 2:
            response = api.post(f'/api/reviews/{r}/papers/{paper_id}/items/{i}/candidates', json={
                'page_number': 1, 'quote': quote, 'start': 0, 'end': len(quote), 'decision': 'accepted',
                'document_version': service.get_paper(r, paper_id)['document_version']})
            assert response.status_code == 201, response.text
    api.patch(f'/api/reviews/{r}/learning', json={'enabled': True})
    service.refresh_search(r, papers[2])
    suggestions = service.get_suggestions(r, papers[2], i)
    target = next(row for row in suggestions if row['learning']['feedback_score'] > 0)
    influence = target['learning']['influences'][0]
    assert influence['paper_id'] != papers[2]
    api.patch(f'/api/reviews/{r}', json={'settings': {'last_paper_id': papers[2], 'last_item_id': i}})
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    card = page.locator(f'[data-suggestion-id="{target["id"]}"]')
    card.get_by_text('Reviewer feedback influence', exact=True).click()
    expect(card).to_contain_text(f'Feedback +{target["learning"]["feedback_score"]:.3f}')
    details = card.locator('.learning-explanation')
    expect(details).to_contain_text('Similarity 1.000')
    source = service.get_paper(r, influence['paper_id'])
    details.get_by_role('button', name=f'Open source: {source["title"]}', exact=True).click()
    expect(page.locator('#paper-select')).to_have_value(source['id'])
    expect(page.locator('#item-select')).to_have_value(i)
    page.wait_for_function("() => document.querySelectorAll('.pdf-highlight.focused').length > 0")


def test_extracted_selection_refresh_stays_with_the_same_source_only(live_app, browser_page):
    url, _api, _service = live_app
    browser_page.goto(url)
    result = browser_page.evaluate("""async () => {
      const {DocumentViewer}=await import('/static/viewer.js');
      const viewer=new DocumentViewer(()=>{},error=>{throw error;});
      const text='𝑥 repeated age text; 𝑥 repeated age text.';
      viewer.pages=[{page_number:1,text,words:[]},{page_number:2,text,words:[]}];
      viewer.renderExtracted();
      const quote='𝑥 repeated age text',start=text.lastIndexOf(quote),end=start+quote.length;
      let node=document.querySelector('#extracted-text').firstChild;
      const selected=getSelection();selected.setBaseAndExtent(node,end,node,start);
      viewer.setHighlights([{id:'new-hit',page_number:1,start:2,end:10}]);
      const preserved=viewer.selection();
      const backwards=selected.anchorNode===selected.focusNode ? selected.anchorOffset>selected.focusOffset :
          !!(selected.anchorNode.compareDocumentPosition(selected.focusNode)&Node.DOCUMENT_POSITION_PRECEDING);
      viewer.pageNumber=2;viewer.renderExtracted();
      const afterPageChange=selected.toString();
      node=document.querySelector('#extracted-text').firstChild;
      selected.setBaseAndExtent(node,start,node,end);
      viewer.pages=viewer.pages.map(page=>({...page}));viewer.renderExtracted();
      return {preserved,backwards,afterPageChange,afterSourceChange:selected.toString()};
    }""")
    assert result['preserved']['quote'] == '𝑥 repeated age text'
    assert result['preserved']['start'] == len('𝑥 repeated age text; ')
    assert result['backwards']
    assert result['afterPageChange'] == result['afterSourceChange'] == ''


def test_source_version_poll_reload_blocks_a_previously_captured_selection(live_app, browser_page):
    from playwright.sync_api import expect
    from test_retrieval import page as source_page
    url, api, service = live_app
    r, p, i = seed(api)
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#suggestions .suggestion-card').first.wait_for()
    page.locator('#text-toggle').click()
    quote = service.get_suggestions(r, p, i)[0]['quote']
    select_extracted_quote(page, quote)
    page.locator('#capture-selection').click()
    page.locator('#evidence-dialog').wait_for(state='visible')
    replacement = 'Treatment stimulation was administered at 31 Hz.'
    service.set_pages(r, p, [source_page(replacement)])
    service.refresh_search(r, p)
    # Processing is already ready before the next poll: version changes must
    # reload the source even when no intermediate processing state was observed.
    expect(page.locator('#extracted-text')).to_have_text(replacement, timeout=15000)
    page.locator('#save-pending-candidate').click()
    expect(page.locator('#toast')).to_contain_text('The paper, item, or document version changed')
    assert not any(s.get('origin') == 'user' for s in service.get_suggestions(r, p, i))
    page.get_by_role('button', name='Close evidence editor', exact=True).click()
    select_extracted_quote(page, replacement)
    page.locator('#capture-selection').click()
    page.locator('#save-pending-candidate').click()
    page.locator('#evidence-dialog').wait_for(state='hidden')
    candidate = next(s for s in service.get_suggestions(r, p, i) if s.get('origin') == 'user')
    assert candidate['quote'] == replacement
    assert candidate['document_version'] == service.get_paper(r, p)['document_version']
