"""Synthetic repeatable-outcome and template-faithful export integration checks."""

from copy import deepcopy
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pytest

from review_extractor.workbooks import (
    export_download_workbook, export_preview, export_workbook,
    preview_template, validate_snapshot, verification_preview,
)


META_HEADERS = ['Study', 'Outcome', 'n', 'pre-score', 'pre-sd', 'post-score', 'post-sd', 'R']


def template_bytes():
    book = Workbook()
    book.active.title = 'Systematic Review'
    book.active.append(['Citation', 'Study Design'])
    book.active.append(['NEVER_IMPORT_THIS_EXAMPLE', '=HYPERLINK("bad")'])
    sheet = book.create_sheet('Meta Analysis')
    sheet.append(META_HEADERS + [None])
    sheet.append(['EXAMPLE_STUDY', 'EXAMPLE_OUTCOME', 999, 999, 999, 999, 999, 999, None])
    for ws in book:
        for cell in ws[1]:
            cell.font = Font(name='Calibri', size=12, bold=True, color='FF112233')
            cell.fill = PatternFill('solid', fgColor='FFDDDDDD')
        ws.column_dimensions['A'].width = 34
        ws.row_dimensions[1].height = 42
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def outcome_snapshot():
    design = preview_template(template_bytes(), all_sheets=True)
    items = [dict(item, id=f'item-{index}', review_id='review-1') for index, item in enumerate(design['items'])]
    papers = [dict(id=f'paper-{index}', review_id='review-1', title=f'Study {index}', filename=f'synthetic-{index}.pdf', included=True,
                   processing_status='ready', search_status='current', document_version=1) for index in (1, 2)]
    outcomes = [dict(id=f'outcome-{index}', paper_id='paper-1', review_id='review-1', outcome='10MWT' if index == 1 else '6MWT',
                     label=f'Result {index}', group='Active', time_point='Week 4', units='m/s' if index == 1 else 'm', included=True, position=index)
                for index in (1, 2)]
    data = dict(review=dict(id='review-1', name='Synthetic review', reviewer='Researcher'), items=items, papers=papers,
                outcomes=outcomes, fields=[], outcome_fields=[], evidence=[], pages=[], history=[])
    text = 'Study one citation\nRandomized trial\n10MWT n=8 pre=0 SD=1 post=2 SD=3 r=.5\n6MWT n=7 pre=4 SD=5 post=6 SD=7 r=.4'
    data['pages'] = [dict(paper_id=paper['id'], page_number=1, status='ok', text=text) for paper in papers]
    for item in items[:2]:
        quote = text.splitlines()[item['position']]
        evidence = dict(id=f"e-{item['id']}", review_id='review-1', paper_id='paper-1', item_id=item['id'], decision='accepted',
                        quote=quote, page_number=1, start=text.index(quote), end=text.index(quote) + len(quote), document_version=1)
        data['evidence'].append(evidence)
        data['fields'].append(dict(review_id='review-1', paper_id='paper-1', item_id=item['id'], status='approved', value=quote, evidence_ids=[evidence['id']]))
    for row_index, outcome in enumerate(outcomes):
        quote = text.splitlines()[row_index + 2]
        for item, value in zip(items[4:], ([8, 0, 1, 2, 3, .5] if row_index == 0 else [7, 4, 5, 6, 7, .4])):
            evidence = dict(id=f"e-{outcome['id']}-{item['id']}", review_id='review-1', paper_id='paper-1', item_id=item['id'],
                            outcome_id=outcome['id'], decision='accepted', quote=quote, page_number=1,
                            start=text.index(quote), end=text.index(quote) + len(quote), document_version=1)
            data['evidence'].append(evidence)
            data['outcome_fields'].append(dict(review_id='review-1', paper_id='paper-1', item_id=item['id'], outcome_id=outcome['id'],
                                               status='approved', value=value, evidence_ids=[evidence['id']]))
    return data


def test_template_classifies_meta_columns_without_importing_examples_or_blank_columns():
    result = preview_template(template_bytes(), all_sheets=True)
    assert len(result['items']) == 10
    assert [item['section'] for item in result['items']] == ['study'] * 2 + ['meta_analysis'] * 8
    assert [item['export_heading'] for item in result['items'][2:]] == META_HEADERS
    assert [item['data_type'] for item in result['items'][2:]] == ['text', 'text', 'integer'] + ['number'] * 5
    assert 'EXAMPLE' not in str(result['items']) and 'value' not in result['items'][0]
    assert result['items'][0]['template_style']['column_width'] == 34


def test_repeatable_outcomes_are_optional_and_do_not_create_study_wide_meta_fields():
    data = outcome_snapshot()
    data['outcomes'] = []
    data['outcome_fields'] = []
    preview = export_preview(data, mark_unfinished=False)
    assert len(preview['rows']) == 2 and len(preview['columns']) == 2
    assert [sheet['row_scope'] for sheet in preview['worksheets']] == ['study', 'outcome']
    assert preview['worksheets'][1]['rows'] == []
    assert validate_snapshot(data)['counts']['fields'] == 4
    book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
    assert book['Meta Analysis'].max_row == 1


@pytest.mark.parametrize('view', ['data', 'verification', 'combined'])
def test_template_exports_keep_exact_columns_and_repeatable_rows_with_matching_locations(view):
    data = outcome_snapshot()
    before = deepcopy(data)
    book = load_workbook(BytesIO(export_download_workbook(data, view)))
    names = ['Systematic Review', 'Meta Analysis']
    verify_names = ['Verify - Systematic Review', 'Verify - Meta Analysis']
    assert book.sheetnames == (names if view != 'verification' else []) + (verify_names if view != 'data' else [])
    if view != 'verification':
        sheet = book['Meta Analysis']
        assert list(sheet.values)[0] == tuple(META_HEADERS)
        assert sheet.max_row == 3  # Two outcomes from one paper; no placeholder for the other.
        assert sheet['A2'].value == 'Study one citation'
        assert sheet['B2'].value == '10MWT | Active | Week 4 | m/s'
        assert sheet['B3'].value == '6MWT | Active | Week 4 | m'
        assert sheet['D2'].value == 0 and sheet['D2'].data_type == 'n'
        assert 'Outcome ID: outcome-1' in sheet['D2'].comment.text
        assert sheet['A1'].font.name == 'Calibri' and sheet['A1'].font.bold
        assert sheet['A1'].fill.fgColor.rgb == 'FFDDDDDD'
        assert sheet.column_dimensions['A'].width == 34
    if view != 'data':
        sheet = book['Verify - Meta Analysis']
        assert sheet['D2'].value == 'Page 1, line 3'
        assert sheet['D3'].value == 'Page 1, line 4'
        assert 'review metadata' in sheet['B2'].value
        assert sheet['D2'].data_type == 's'
        assert list(sheet.values)[0] == tuple(META_HEADERS)
    assert data == before


def test_outcome_sources_cannot_bleed_across_rows_papers_or_reviews():
    data = outcome_snapshot()
    field = data['outcome_fields'][0]
    field['evidence_ids'] = [data['outcome_fields'][6]['evidence_ids'][0]]
    report = validate_snapshot(data)
    assert any(issue['code'] == 'invalid_evidence' and issue['outcome_id'] == 'outcome-1' for issue in report['issues'])
    assert export_preview(data, mark_unfinished=False)['worksheets'][1]['rows'][0]['values'][2] is None
    assert verification_preview(data)['worksheets'][1]['rows'][0]['values'][2] is None
    field['status'] = 'unreviewed'
    field['evidence_ids'] = []
    own_evidence = next(record for record in data['evidence'] if record.get('outcome_id') == 'outcome-1' and record['item_id'] == field['item_id'])
    own_evidence['review_id'] = 'other-review'
    assert export_preview(data, include_accepted=True, mark_unfinished=False)['worksheets'][1]['rows'][0]['values'][2] is None


def test_selected_outcome_passages_are_literal_and_keep_source_location():
    data = outcome_snapshot()
    item = data['items'][3]  # Outcome name, a text item.
    quote = '=HYPERLINK("https://example.invalid", "literal source")'
    data['pages'][0]['text'] += '\n' + quote
    data['evidence'].append(dict(id='injection', review_id='review-1', paper_id='paper-1', item_id=item['id'], outcome_id='outcome-1',
                                 decision='accepted', page_number=1, quote=quote, start=data['pages'][0]['text'].index(quote),
                                 end=len(data['pages'][0]['text']), document_version=1))
    book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
    assert book['Meta Analysis']['B2'].value == quote
    assert book['Meta Analysis']['B2'].data_type == 's'
    assert book['Verify - Meta Analysis']['B2'].value == 'Page 1, line 5'
    assert book['Meta Analysis']['B3'].value.startswith('6MWT')
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)


def test_current_design_and_outcome_selection_drive_export_without_erasing_audit():
    data = outcome_snapshot()
    data['outcomes'][0]['included'] = False
    data['items'][4]['deleted'] = True  # n column
    data['items'][5].update(export_heading='Baseline mean', position=100)
    book = load_workbook(BytesIO(export_workbook(data, draft=False)))
    assert book.sheetnames[:4] == ['Systematic Review', 'Meta Analysis', 'Verify - Systematic Review', 'Verify - Meta Analysis']
    sheet = book['Meta Analysis']
    assert sheet.max_row == 2 and sheet['B2'].value.startswith('6MWT')
    assert list(sheet.values)[0][-1] == 'Baseline mean'
    assert 'n' not in list(sheet.values)[0]
    assert book['Outcome rows'].max_row == 3
    assert book['Outcome values'].max_row == 13
    assert 'outcome-1' in str(list(book['Evidence'].values))
    assert 'meta_analysis' in str(list(book['Data dictionary'].values))


@pytest.mark.parametrize('item_index,value', [(4, 0), (6, -1), (9, 1.1)])
def test_invalid_sample_size_sd_or_correlation_block_final_export(item_index, value):
    data = outcome_snapshot()
    field = next(field for field in data['outcome_fields'] if field['item_id'] == data['items'][item_index]['id'])
    field['value'] = value
    assert not validate_snapshot(data)['can_finalize']
    with pytest.raises(ValueError, match='blocked'):
        export_workbook(data, draft=False)


def test_outcome_withdrawn_evidence_and_stale_document_do_not_leak_into_downloads():
    for change in [dict(decision='rejected'), dict(source_stale=True), dict(document_version=5)]:
        data = outcome_snapshot()
        evidence = next(record for record in data['evidence'] if record['id'] == data['outcome_fields'][0]['evidence_ids'][0])
        evidence.update(change)
        book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
        assert book['Meta Analysis']['C2'].value is None
        assert book['Verify - Meta Analysis']['C2'].value is None


def test_foreign_and_deleted_outcome_records_are_not_exported():
    data = outcome_snapshot()
    data['outcomes'][0]['deleted'] = True
    data['outcomes'][1]['review_id'] = 'private-review'
    data['outcomes'].append(dict(id='private-outcome', review_id='review-1', paper_id='private-paper', label='PRIVATE'))
    data['outcome_fields'].append(dict(review_id='private-review', paper_id='paper-1', item_id='item-4', outcome_id='outcome-1', value='PRIVATE'))
    book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
    assert book['Meta Analysis'].max_row == 1
    assert 'PRIVATE' not in str([[list(row) for row in sheet.values] for sheet in book])


def test_invalid_nested_template_style_is_ignored_safely():
    data = outcome_snapshot()
    data['items'][2]['template_style'] = dict(column_width=float('inf'), header_height=-1,
        header=dict(font=dict(name=['bad'], size=99999, color='not-a-color'), fill=['bad'],
                    alignment=dict(horizontal={}, vertical=[], text_rotation=float('nan')),
                    border=dict(left=dict(style={}, color=123))), body=['bad'])
    data['items'][3]['template_style'] = ['not-a-style-object']
    book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
    assert book['Meta Analysis']['D2'].value == 0
    assert book['Meta Analysis'].column_dimensions['A'].width <= 255


def test_same_worksheet_name_with_mixed_row_scopes_keeps_both_tables():
    data = outcome_snapshot()
    for item in data['items']:
        item['worksheet_name'] = 'Shared'
    preview = export_preview(data)
    assert [table['row_scope'] for table in preview['worksheets']] == ['study', 'outcome']
    book = load_workbook(BytesIO(export_download_workbook(data, 'combined')))
    assert book.sheetnames == ['Shared', 'Shared (2)', 'Verify - Shared', 'Verify - Shared (2)']
    assert book['Shared'].max_column == 2
    assert book['Shared (2)'].max_column == 8
    assert book['Shared (2)']['D2'].value == 0
