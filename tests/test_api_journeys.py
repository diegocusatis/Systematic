"""End-to-end HTTP journeys, persisted through separate app instances."""

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from review_extractor.app import create_app

DEMO = Path(__file__).resolve().parents[1] / "demo"


def ok(response, expected=200):
    assert response.status_code == expected, response.text
    return response.json()


def review(client, name="Example review"):
    return ok(client.post("/api/reviews", json={"name": name, "reviewer": "Test researcher"}), 201)["id"]


def test_ui_assets_revalidate_after_local_updates(tmp_path):
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        for path in ('/', '/static/app.js', '/static/viewer.js', '/static/style.css'):
            response = client.get(path)
            assert response.status_code == 200
            assert response.headers['cache-control'] == 'no-cache, must-revalidate'
        assert '/static/app.js?v=' in client.get('/').text
        assert client.get('/api/reviews').headers['cache-control'] == 'no-store'


def upload(client, r, name="study-A.pdf"):
    with (DEMO / name).open("rb") as pdf:
        return ok(client.post(f"/api/reviews/{r}/papers", files={"files": (name, pdf, "application/pdf")}), 202)["papers"][0]


def add_item(client, r, **overrides):
    data = {"name": "Adverse events", "data_type": "integer", "required": True,
            "keywords": ["adverse events"]}
    data.update(overrides)
    return ok(client.post(f"/api/reviews/{r}/items", json=data), 201)


def test_imported_worksheet_groups_survive_restart_and_keep_same_headings_separate(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = 'Participants'
    book.active.append(['Result'])
    book.active.append(['EXAMPLE ONLY'])
    book.create_sheet('Outcomes').append(['Result'])
    stream = BytesIO(); book.save(stream)
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        r = review(client)
        preview = ok(client.post('/api/templates/preview', files={'file': ('multi.xlsx', stream.getvalue())}, data={'all_sheets': 'true'}))
        imported = ok(client.post(f'/api/reviews/{r}/items/import', json={'items': preview['items']}))['items']
        assert len({item['id'] for item in imported}) == 2
        assert [item['worksheet_name'] for item in imported] == ['Participants', 'Outcomes']
        assert client.app.state.service.export_snapshot(r)['fields'] == []
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        columns = ok(client.get(f'/api/reviews/{r}/export-preview'))['columns']
        assert [column['worksheet_name'] for column in columns] == ['Participants', 'Outcomes']
        response = client.get(f'/api/reviews/{r}/export?view=combined')
        assert response.status_code == 200
        exported = load_workbook(BytesIO(response.content))
        assert exported.sheetnames == ['Participants', 'Outcomes', 'Verify - Participants', 'Verify - Outcomes']


def test_reference_workbook_preview_preserves_all_29_columns(tmp_path):
    from openpyxl import Workbook
    workbook = Path("C:/Users/diego/Downloads/Systematic Example.xlsx")
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        reference = ok(client.get("/api/reference-template"))
        assert reference["example_included"] is False
        assert len(reference["items"]) == 29
        headings = [item.get("export_heading", item["name"]) for item in reference["items"]]
        assert headings[0] == "Citation"
        assert headings[-1] == "Follow-up timing"
        assert "Baseline / reported GMS outcome scores" in headings
        if workbook.exists():
            content = workbook.read_bytes()
        else:
            generated = Workbook()
            generated.active.append(headings)
            generated.active.append(["UNVERIFIED EXAMPLE"])
            stream = BytesIO()
            generated.save(stream)
            content = stream.getvalue()
        preview = ok(client.post("/api/templates/preview", files={"file": ("reference.xlsx", content)},
                                 data={"sheet": "Sheet1" if workbook.exists() else "Sheet", "header_rows": "1"}))
        assert [x["export_heading"] for x in preview["items"]] == headings
        r = review(client)
        imported = ok(client.post(f"/api/reviews/{r}/items/import", json={"items": preview["items"]}))
        assert len(imported["items"]) == 29
        assert ok(client.get(f"/api/reviews/{r}/papers")) == []


def test_remove_paper_is_scoped_persistent_and_restorable(tmp_path):
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        r = review(client)
        other = review(client, 'Other review')
        i = add_item(client, r)['id']
        removed = upload(client, r, 'study-C.pdf')
        remaining = upload(client, r, 'study-A.pdf')
        p = removed['id']
        service = client.app.state.service
        fields = service.store.list('fields', review_id=r, paper_id=p)
        client.patch(f'/api/reviews/{r}', json={'settings': {'last_paper_id': p}})
        assert client.delete(f'/api/reviews/{other}/papers/{p}').status_code == 404
        result = ok(client.delete(f'/api/reviews/{r}/papers/{p}'))
        assert result['removed'] is True
        assert [paper['id'] for paper in ok(client.get(f'/api/reviews/{r}/papers'))] == [remaining['id']]
        assert Path(removed['path']).exists()
        assert service.store.list('fields', review_id=r, paper_id=p) == fields
        assert service.get_review(r)['settings']['last_paper_id'] == remaining['id']
        snapshot = service.snapshot(r)
        assert all(row['paper_id'] != p for table in ('pages', 'suggestions', 'fields', 'evidence') for row in snapshot[table])
        assert not any(issue.get('paper_id') == p for issue in ok(client.get(f'/api/reviews/{r}/validation'))['issues'])
        preview = ok(client.get(f'/api/reviews/{r}/export-preview'))
        assert [row['paper_id'] for row in preview['rows']] == [remaining['id']]
        ok(client.delete(f'/api/reviews/{r}/papers/{p}'))  # Idempotent.
        assert len([event for event in snapshot['history'] if event['action'] == 'paper_removed']) == 1
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        assert len(ok(client.get(f'/api/reviews/{r}/papers'))) == 1
        restored = upload(client, r, 'study-C.pdf')
        assert restored['id'] == p and restored['restored']
        assert len(ok(client.get(f'/api/reviews/{r}/papers'))) == 2
        assert client.app.state.service.store.list('fields', review_id=r, paper_id=p) == fields


def test_review_wide_edits_evidence_and_export_survive_restart(tmp_path):
    app = create_app(tmp_path, run_background=False)
    with TestClient(app) as client:
        r = review(client)
        papers = [upload(client, r, f"study-{letter}.pdf") for letter in "ABC"]
        item = add_item(client, r)
        i = item["id"]
        ok(client.post(f"/api/reviews/{r}/refresh"))
        # Adding while reading A reaches B/C and later D, without copying values.
        for paper in papers:
            p = paper["id"]
            assert ok(client.get(f"/api/reviews/{r}/papers/{p}/items/{i}/field"))["status"] == "unreviewed"
            assert ok(client.get(f"/api/reviews/{r}/papers/{p}/suggestions", params={"item_id": i}))
        d = upload(client, r, "study-D.pdf")
        ok(client.post(f"/api/reviews/{r}/refresh", params={"paper_id": d["id"]}))
        assert ok(client.get(f"/api/reviews/{r}/papers/{d['id']}/suggestions", params={"item_id": i}))
        p = papers[0]["id"]
        suggestions = ok(client.get(f"/api/reviews/{r}/papers/{p}/suggestions", params={"item_id": i}))
        rejected = suggestions[0]
        ok(client.patch(f"/api/reviews/{r}/papers/{p}/suggestions/{rejected['id']}", json={"decision": "rejected"}))
        # A separate source passage is accepted; approving zero must retain numeric zero.
        quote = "Treatment group had 0 adverse events at week 3."
        evidence = ok(client.post(f"/api/reviews/{r}/papers/{p}/items/{i}/evidence", json={
            "page_number": 1, "quote": quote, "decision": "accepted", "notes": "Treatment group at week 3",
        }), 201)
        approved = ok(client.put(f"/api/reviews/{r}/papers/{p}/items/{i}/field", json={
            "status": "approved", "value": 0, "value_kind": "reported", "evidence_ids": [evidence["id"]],
            "observations": [{"group": "Treatment", "time_point": "week 3", "value": 0, "units": "events", "evidence_ids": [evidence["id"]]}],
        }))
        assert approved["value"] == 0
        preview = ok(client.get(f"/api/reviews/{r}/export-preview"))
        assert preview["rows"][0]["values"] == [0]
        assert all(row["values"] == [None] for row in preview["rows"][1:])
        verification = client.get(f"/api/reviews/{r}/export", params={"draft": True, "verification": True})
        assert verification.status_code == 200
        assert 'verification-DRAFT.xlsx' in verification.headers['content-disposition']
        locations = load_workbook(BytesIO(verification.content))
        assert locations.worksheets[0]['A2'].value.startswith('Page 1, line')
        assert locations.worksheets[0]['A3'].value is None
        second = add_item(client, r, name="Narrative", data_type="narrative", keywords=["completed"])
        ok(client.patch(f"/api/reviews/{r}/items/{i}", json={"name": "Events at week 3", "export_heading": "Week 3 events"}))
        ok(client.post(f"/api/reviews/{r}/items/reorder", json={"ids": [second["id"], i]}))
        assert ok(client.get(f"/api/reviews/{r}/papers/{p}/items/{i}/field"))["value"] == 0
        ok(client.delete(f"/api/reviews/{r}/items/{i}"))
        assert i not in [x["id"] for x in ok(client.get(f"/api/reviews/{r}/items"))]
        preview = ok(client.get(f"/api/reviews/{r}/export-preview"))
        assert "Week 3 events" not in [x["heading"] for x in preview["columns"]]
        ok(client.post(f"/api/reviews/{r}/items/{i}/restore"))
        assert ok(client.get(f"/api/reviews/{r}/papers/{p}/items/{i}/field"))["value"] == 0
        # Broaden rules, retaining passage rejection on A without affecting B.
        ok(client.patch(f"/api/reviews/{r}/items/{i}", json={"keywords": ["adverse events", "follow-up timing"]}))
        ok(client.post(f"/api/reviews/{r}/refresh"))
        refreshed = ok(client.get(f"/api/reviews/{r}/papers/{p}/suggestions", params={"item_id": i}))
        assert next(x for x in refreshed if x["id"] == rejected["id"])["decision"] == "rejected"
        assert any(x["page_number"] == 2 and x["active"] for x in refreshed)
        assert all(x["decision"] == "pending" for x in ok(client.get(
            f"/api/reviews/{r}/papers/{papers[1]['id']}/suggestions", params={"item_id": i})))
        assert ok(client.get(f"/api/reviews/{r}/papers/{p}"))["search_status"] == "current"
        r2 = review(client, "Independent review")
        add_item(client, r2, name="Independent item")
        assert len(ok(client.get(f"/api/reviews/{r}/items"))) == 2
        assert client.patch(f"/api/reviews/{r2}/items/{i}", json={"name": "Wrong review"}).status_code == 404
        # Draft export reflects current design and one row per included paper.
        exported = client.get(f"/api/reviews/{r}/export", params={"draft": "true"})
        assert exported.status_code == 200
        book = load_workbook(BytesIO(exported.content))
        sheet = book.worksheets[0]
        headers = [c.value for c in sheet[1]]
        assert headers == ["Narrative", "Week 3 events"]
        assert sheet.max_row == 5
        assert sheet.cell(2, 2).value == 0
        assert sheet.cell(2, 2).data_type == "n"
        assert sheet.freeze_panes == "A2"
        assert any("evidence" in s.lower() for s in book.sheetnames)
        assert any("history" in s.lower() for s in book.sheetnames)
        assert client.get(f"/api/reviews/{r}/export", params={"draft": "false"}).status_code == 409
        ok(client.patch(f"/api/reviews/{r}", json={"settings": {"last_paper_id": p, "last_item_id": i}}))
    # New application instance + new SQLite connections, same local folder.
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        assert ok(client.get(f"/api/reviews/{r}"))["settings"]["last_item_id"] == i
        field = ok(client.get(f"/api/reviews/{r}/papers/{p}/items/{i}/field"))
        assert field["value"] == 0
        assert field["evidence_ids"] == [evidence["id"]]
        suggestions = ok(client.get(f"/api/reviews/{r}/papers/{p}/suggestions", params={"item_id": i}))
        assert next(x for x in suggestions if x["id"] == rejected["id"])["decision"] == "rejected"
        assert len({x["id"] for x in suggestions}) == len(suggestions)
        # A material definition change retains value but requires reconsideration.
        ok(client.patch(f"/api/reviews/{r}/items/{i}", json={"definition": "Include only serious events at week 5."}))
        field = ok(client.get(f"/api/reviews/{r}/papers/{p}/items/{i}/field"))
        assert field["value"] == 0 and field["needs_reconsideration"]


def test_final_export_traceability_and_manual_decisions(tmp_path):
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        r = review(client)
        p = upload(client, r)["id"]
        ok(client.post(f"/api/reviews/{r}/refresh"))
        i = add_item(client, r, name="Result", data_type="narrative")["id"]
        ok(client.post(f"/api/reviews/{r}/refresh"))
        base = f"/api/reviews/{r}/papers/{p}/items/{i}"
        # Invented quotations and unsupported approvals are rejected.
        assert client.post(base + "/evidence", json={"page_number": 1, "quote": "Invented quotation", "decision": "accepted"}).status_code == 422
        assert client.put(base + "/field", json={"status": "approved", "value": "claimed fact", "evidence_ids": []}).status_code == 422
        accepted = []
        for number, quote in [(1, "Treatment group: 13 participants completed; control: 12 completed."),
                              (2, "Flow diagram: treatment 12 completed; control 13 completed.")]:
            accepted.append(ok(client.post(base + "/evidence", json={"page_number": number, "quote": quote, "decision": "accepted"}), 201)["id"])
        ok(client.put(base + "/field", json={"status": "conflicting", "value": "Counts disagree", "evidence_ids": accepted}))
        assert not ok(client.get(f"/api/reviews/{r}/validation"))["can_finalize"]
        narrative = "Text: treatment 13/control 12; diagram: treatment 12/control 13. Discrepancy retained."
        ok(client.put(base + "/field", json={"status": "approved", "value": narrative, "value_kind": "interpretation",
                                              "rationale": "Both reported counts are preserved without choosing a winner.", "evidence_ids": accepted}))
        assert ok(client.get(f"/api/reviews/{r}/validation"))["can_finalize"]
        response = client.get(f"/api/reviews/{r}/export", params={"draft": "false"})
        assert response.status_code == 200
        book = load_workbook(BytesIO(response.content))
        assert book.worksheets[0]["A2"].value == narrative
        all_values = str([[row for row in sheet.values] for sheet in book.worksheets[1:]])
        assert p in all_values and i in all_values
        assert "Flow diagram: treatment 12 completed; control 13 completed." in all_values
        assert "Test researcher" in all_values
        # Withdrawing accepted evidence invalidates approval, without erasing it.
        ok(client.patch(f"/api/reviews/{r}/papers/{p}/evidence/{accepted[0]}", json={"decision": "rejected"}))
        assert not ok(client.get(f"/api/reviews/{r}/validation"))["can_finalize"]
        assert ok(client.get(base + "/field"))["value"] == narrative


def test_duplicate_scans_errors_security_and_isolation(tmp_path):
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        r = review(client)
        scan = upload(client, r, "study-C.pdf")
        duplicate = upload(client, r, "study-C.pdf")
        assert duplicate["duplicate"] is True and duplicate["id"] == scan["id"]
        assert len(ok(client.get(f"/api/reviews/{r}/papers"))) == 1
        ok(client.post(f"/api/reviews/{r}/refresh"))
        pages = ok(client.get(f"/api/reviews/{r}/papers/{scan['id']}/pages"))
        assert pages[2]["status"] == "ocr_required"
        r2 = review(client, "Other")
        assert client.get(f"/api/reviews/{r2}/papers/{scan['id']}/pdf").status_code == 404
        assert client.post("/api/reviews", json={"name": "Attack"}, headers={"Origin": "https://unrelated.example"}).status_code == 403
        assert client.get("/api/health", headers={"Host": "unrelated.example"}).status_code == 400
        invalid = client.post(f"/api/reviews/{r}/papers", files={"files": ("bad.pdf", b"not a pdf")})
        assert invalid.status_code == 422
        broken = ok(client.post(f"/api/reviews/{r}/papers", files={"files": ("broken.pdf", b"%PDF- malformed")}), 202)["papers"][0]
        ok(client.post(f"/api/reviews/{r}/refresh", params={"paper_id": broken["id"]}))
        state = ok(client.get(f"/api/reviews/{r}/papers/{broken['id']}"))
        assert state["processing_status"] == "error" and state["processing_error"]


def test_template_import_appends_and_rolls_back_invalid_batch(tmp_path):
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        r = review(client)
        original = add_item(client, r)
        imported = ok(client.post(f"/api/reviews/{r}/items/import", json={"items": [
            {"name": "Appended first", "position": 0}, {"name": "Appended second", "position": 1},
        ]}))
        items = ok(client.get(f"/api/reviews/{r}/items"))
        assert [item["name"] for item in items] == ["Adverse events", "Appended first", "Appended second"]
        assert items[0]["id"] == original["id"]
        assert len(imported["items"]) == 2
        failed = client.post(f"/api/reviews/{r}/items/import", json={"items": [
            {"name": "Must not persist"}, {"name": "", "data_type": "nonsense"},
        ]})
        assert failed.status_code == 422
        assert len(ok(client.get(f"/api/reviews/{r}/items"))) == 3


def test_interrupted_processing_waits_for_manual_scan_on_new_app_start(tmp_path):
    from review_extractor.service import ReviewService
    service = ReviewService(tmp_path)
    r = service.create_review({"name": "Interrupted", "reviewer": "Researcher"})["id"]
    item = service.add_item(r, {"name": "Frequency", "keywords": ["20 Hz"]})
    paper = service.add_paper(r, {"filename": "study-A.pdf", "path": str(DEMO / "study-A.pdf"),
                                  "processing_status": "processing", "processing_progress": 25})
    service.store.close()
    with TestClient(create_app(tmp_path, run_background=False)) as client:
        p = paper["id"]
        processed = ok(client.get(f"/api/reviews/{r}/papers/{p}"))
        assert processed["processing_status"] == "pending"
        assert processed["processing_progress"] == 25
        assert not ok(client.get(f"/api/reviews/{r}/papers/{p}/pages"))
        ok(client.post(f"/api/reviews/{r}/refresh"))
        processed = ok(client.get(f"/api/reviews/{r}/papers/{p}"))
        assert processed["processing_status"] == "ready"
        assert processed["processing_progress"] == 100
        assert processed["search_status"] == "current"
        assert ok(client.get(f"/api/reviews/{r}/papers/{p}/suggestions", params={"item_id": item["id"]}))
