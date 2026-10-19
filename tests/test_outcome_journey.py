"""The template-to-outcome-to-verification journey uses synthetic evidence only."""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from review_extractor.app import create_app
from test_retrieval import page


def checked(response, *statuses):
    assert response.status_code in (statuses or (200,)), response.text
    return response.json()


def template_bytes():
    book = Workbook()
    book.active.title = "Systematic Review"
    book.active.append(["Citation", "Sample Size"])
    book.active.append(["EXAMPLE MUST NOT BECOME EVIDENCE", 9999])
    meta = book.create_sheet("Meta Analysis")
    meta.append(["Study", "Outcome", "n", "pre-score", "pre-sd", "post-score", "post-sd", "R"])
    meta.column_dimensions["B"].width = 30
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_template_outcomes_round_trip_independent_values_and_verification(tmp_path):
    data_dir = tmp_path / "review"
    with TestClient(create_app(data_dir, run_background=False)) as client:
        review = checked(client.post("/api/reviews", json={"name": "Synthetic outcome review", "reviewer": "Reviewer"}), 201)
        root = f'/api/reviews/{review["id"]}'
        preview = checked(client.post("/api/templates/preview", files={"file": ("synthetic.xlsx", template_bytes())}, data={"all_sheets": "true"}))
        items = checked(client.post(root + "/items/import", json={"items": preview["items"]}))["items"]
        assert [i.get("section", "study") for i in items] == ["study"] * 2 + ["meta_analysis"] * 8
        assert checked(client.get(root + "/papers")) == []
        service = client.app.state.service
        paper = service.add_paper(review["id"], {"title": "Synthetic study", "filename": "synthetic.pdf", "sha256": "synthetic-one"})
        text = "Results\nActive group: baseline 0 points, week 4 score 2 points.\nSham group: baseline 7 points, week 4 score 8 points."
        service.set_pages(review["id"], paper["id"], [page(text)])
        path = root + f'/papers/{paper["id"]}'
        outcomes = [checked(client.post(path + "/outcomes", json={"outcome": "Walking score", "group": group, "time_point": "Week 4", "units": "points"}), 200, 201) for group in ("Active", "Sham")]
        item = next(i for i in items if i["name"] == "pre-score")
        field_path = path + f'/items/{item["id"]}/field'
        evidence_ids = []
        for outcome, value, quote in zip(outcomes, (0, 7), ("baseline 0 points", "baseline 7 points")):
            params = {"outcome_id": outcome["id"]}
            start = text.index(quote)
            evidence = checked(client.post(path + f'/items/{item["id"]}/evidence', params=params, json={"page_number": 1, "start": start, "end": start + len(quote), "quote": quote, "decision": "accepted"}), 200, 201)
            evidence_ids.append(evidence["id"])
            checked(client.put(field_path, params=params, json={"status": "approved", "value": value, "evidence_ids": [evidence["id"]]}))
        # A source from another outcome must not be usable for this row.
        response = client.put(field_path, params={"outcome_id": outcomes[0]["id"]}, json={"status": "approved", "value": 7, "evidence_ids": [evidence_ids[1]]})
        assert response.status_code in {404, 422}
        assert checked(client.get(field_path, params={"outcome_id": outcomes[0]["id"]}))["value"] == 0
        table = checked(client.get(root + "/export-preview"))
        assert [(sheet["name"], len(sheet["rows"])) for sheet in table["worksheets"]] == [("Systematic Review", 1), ("Meta Analysis", 2)]
        for view, expected in (("data", ["Systematic Review", "Meta Analysis"]), ("verification", ["Verify - Systematic Review", "Verify - Meta Analysis"]), ("combined", ["Systematic Review", "Meta Analysis", "Verify - Systematic Review", "Verify - Meta Analysis"])):
            response = client.get(root + "/export", params={"view": view})
            assert response.status_code == 200, response.text
            book = load_workbook(BytesIO(response.content))
            assert book.sheetnames == expected
            if view != "verification":
                assert [book["Meta Analysis"].cell(row, 4).value for row in (2, 3)] == [0, 7]
                assert book["Meta Analysis"]["D2"].data_type == "n"
                assert book["Meta Analysis"]["B2"].value != book["Meta Analysis"]["B3"].value
            if view != "data":
                assert book["Verify - Meta Analysis"]["D2"].value == "Page 1, line 2"
                assert book["Verify - Meta Analysis"]["D3"].value == "Page 1, line 3"
        checked(client.patch(path + f'/outcomes/{outcomes[1]["id"]}', json={"included": False}))
        assert len(checked(client.get(root + "/export-preview"))["worksheets"][1]["rows"]) == 1
    # Opening a fresh service preserves independent rows, selection and numeric zero.
    with TestClient(create_app(data_dir, run_background=False)) as client:
        assert len(checked(client.get(path + "/outcomes"))) == 2
        assert checked(client.get(field_path, params={"outcome_id": outcomes[0]["id"]}))["value"] == 0
        assert checked(client.get(field_path, params={"outcome_id": outcomes[1]["id"]}))["value"] == 7
        checked(client.patch(path + f'/outcomes/{outcomes[1]["id"]}', json={"included": True}))
        assert len(checked(client.get(root + "/export-preview"))["worksheets"][1]["rows"]) == 2
