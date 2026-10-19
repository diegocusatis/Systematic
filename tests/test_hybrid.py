"""Scientific confounders, conservative proposals, cache and audit regressions."""
from copy import deepcopy
from io import BytesIO
import sys
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from reportlab.pdfgen import canvas

from review_extractor.documents import process_pdf
from review_extractor.hybrid.chunking import chunk_document, normalize
from review_extractor.hybrid.extractor import extract_value, validate_value
from review_extractor.hybrid.models import item_spec
from review_extractor.hybrid.semantic import LocalSemantic
from review_extractor.retrieval import search_item
from review_extractor.service import ReviewService
from review_extractor.workbooks import export_workbook
from test_retrieval import page, item


@pytest.mark.parametrize('name,bad,good', [
    ('stimulation_frequency', 'Surface EMG signals were sampled at 2000 Hz.', 'Repetitive TMS was administered over M1 at 10 Hz.'),
    ('Stimulation Intensity', 'The intensity of rehabilitation was progressively increased.', 'The stimulation intensity was 90% of resting motor threshold.'),
    ('sample_size', 'Participants completed 24 treatment sessions.', 'Twenty-four participants were randomized.'),
    ('LEMS SD of change', 'Baseline LEMS was 32.1 ± 4.2.', 'The change in LEMS was 4.6 ± 2.7.'),
])
def test_scientific_ranking(name, bad, good):
    rows = search_item([page('Methods\n'+bad+'\n\n'+good)], item(name, search_mode='fast'))['matches']
    best_good = max(r['rank_score'] for r in rows if r['quote'] == good)
    bad_scores = [r['rank_score'] for r in rows if r['quote'] == bad]
    assert not bad_scores or best_good > max(bad_scores) + .15
    assert all(0 <= r['rank_score'] <= 1 for r in rows)
    for row in rows:
        assert row['scores']['semantic_score'] is None


def test_source_preservation_normalization_and_neighbor_context():
    text = 'Methods\nThe protocol was prespecified.\n10-Hz rTIVIS was deliv-\nered over M1.\nPatients rested afterward.'
    source = page(text)
    source['paper_id'] = 'stable-paper'
    chunks = chunk_document([source])
    middle = next(c for c in chunks if 'rTIVIS' in c['quote'])
    assert middle['quote'] == text[middle['start']:middle['end']]
    assert middle['normalized_text'] == '10-Hz rTMS was delivered over M1.'
    assert middle['document_id'] == 'stable-paper' and middle['section'] == 'Methods'
    assert middle['previous_chunk_id'] and middle['next_chunk_id'] and middle['boxes']
    assert 'prespecified' in middle['context'] and 'afterward' in middle['context']
    assert normalize('1OMWT 10H2 4.6 +/- 2.7') == '10MWT 10 Hz 4.6 ± 2.7'
    assert normalize('4.6 + 2.7') == '4.6 + 2.7'  # '+' alone is not SD evidence


def test_fuzzy_support_and_section_preference():
    rows = search_item([page('Introduction\nStimulatlon frequency was 10 Hz.\n\nMethods\nStimulation frequency was 10 Hz.')], item('stimulation_frequency',search_mode='fast'))['matches']
    assert rows[0]['section'] == 'Methods'
    assert any(r['scores']['fuzzy_score'] > 0 for r in rows)


def test_statistical_context_and_sd_change_abstention():
    spec = item_spec(item('SD of LEMS change — intervention group', context={'time_point':'4 weeks'}))
    def extract(text): return extract_value(spec, dict(quote=text,page_number=7))
    good = extract('Active rTMS group at 4 weeks: change in LEMS mean (SD) was 4.6 (2.7).')
    assert good['value'] == '2.7' and good['statistic'] == 'SD_change'
    assert good['group'] == 'intervention' and good['outcome'] == 'LEMS' and good['timepoint'] == '4 weeks'
    assert validate_value(spec,good)['valid']
    for text in [
        'Active rTMS group at baseline: LEMS mean (SD) was 32.1 (4.2).',
        'Control group at 4 weeks: change in LEMS mean (SD) was 4.6 (2.7).',
        'Active rTMS group at 8 weeks: change in LEMS mean (SD) was 4.6 (2.7).',
        'Active rTMS group at 4 weeks: change in 10MWT mean (SD) was 4.6 (2.7).',
        'Active rTMS group at 4 weeks: change in LEMS was 4.6 ± 2.7.',
        'The change in LEMS was 4.6 ± 2.7.',
    ]:
        value = extract(text)
        assert value['value'] is None and value['ambiguous'], text


def test_equal_numbers_in_different_arms_are_still_ambiguous():
    rows = search_item([page('Active stimulation used 10 Hz; sham used 10 Hz.')],item('stimulation_frequency'))['matches']
    assert rows[0]['proposed_value'] is None and rows[0]['multiple_proposed_values']


def test_table_outcomes_never_share_nearest_sd():
    text = 'LEMS 4.6 (2.7)\n10MWT 2.1 (1.3)\nWISCI-II 1.1 (0.8)'
    source = page(text)
    source['tables'] = [dict(id='table-2',title='Table 2', headers_inferred=False,cells=[])]
    for r,line in enumerate(text.splitlines()):
        start = text.index(line)
        source['tables'][0]['cells'].append(dict(row=r,column=1,start=start,end=start+len(line),text=line.split(' ',1)[1],
            row_header=line.split()[0],column_header='Active rTMS | 4 weeks | Change mean (SD)'))
    rows = search_item([source],item('SD of LEMS change — intervention group',context={'time_point':'4 weeks'},search_mode='fast'))['matches']
    cells = [r for r in rows if r.get('table_id')]
    assert next(r for r in cells if r['row_header']=='LEMS')['proposed_value'] == '2.7'
    assert all(r['proposed_value'] is None for r in cells if r['row_header']!='LEMS')
    assert max(r['rank_score'] for r in cells if r['row_header']=='LEMS') > max(r['rank_score'] for r in cells if r['row_header']!='LEMS')


class Encoder:
    def __init__(self): self.calls=[]
    def encode(self,texts,**kwargs):
        self.calls.append((list(texts),kwargs))
        return [[1.,0.] for t in texts]


def test_embedding_cache_survives_restart_and_invalidates_on_text_model_change(tmp_path):
    encoder = Encoder()
    semantic = LocalSemantic(tmp_path,model='synthetic-model-v1',encoder=encoder)
    semantic.similarities('query',['source 1','source 2'])
    assert len(encoder.calls)==2
    semantic.similarities('query',['source 1','source 2'])
    assert len(encoder.calls)==2
    restarted = Encoder()
    LocalSemantic(tmp_path,model='synthetic-model-v1',encoder=restarted).similarities('new query',['source 1','source 2'])
    assert len(restarted.calls)==1 and restarted.calls[0][0]==['new query']
    semantic.similarities('query',['changed source'])
    assert len(encoder.calls)==3
    second = Encoder()
    LocalSemantic(tmp_path,model='synthetic-model-v2',encoder=second).similarities('query',['source 1','source 2'])
    assert len(second.calls)==2
    assert encoder.calls[0][1]['batch_size']==32


def test_document_embeddings_do_not_change_with_item_rules(tmp_path):
    encoder = Encoder()
    semantic = LocalSemantic(tmp_path,encoder=encoder)
    pages=[page('Methods\nStimulation used 10 Hz. It was administered at 90% RMT.')]
    search_item(pages,item('Stimulation frequency'),semantic=semantic)
    search_item(pages,dict(item('Stimulation intensity'),rules=[dict(kind='regex',pattern='90%')]),semantic=semantic)
    assert len(encoder.calls)==3  # one document batch and two item queries


def test_semantic_failure_is_explicit_and_fast_mode_never_loads(tmp_path,monkeypatch):
    semantic=LocalSemantic(tmp_path)
    semantic.error='Model unavailable in offline cache'
    rows=search_item([page('rTMS used 10 Hz.')],item('Stimulation frequency'),semantic=semantic)
    assert rows['warnings'] and rows['matches'][0]['search_mode_used']=='fast'
    def fail(): raise AssertionError('fast mode must not load a model')
    monkeypatch.setattr(semantic,'_load',fail)
    rows=search_item([page('rTMS used 10 Hz.')],item('Stimulation frequency',search_mode='fast'),semantic=semantic)
    assert not rows['warnings']


def test_feedback_survives_refresh_and_is_scoped_and_exported(tmp_path):
    service=ReviewService(tmp_path)
    r=service.create_review({'name':'Retrieval'})['id']
    i=service.add_item(r,{'name':'Stimulation frequency','search_mode':'fast'})['id']
    p=service.add_paper(r,{'title':'Synthetic','processing_status':'ready'})['id']
    service.set_pages(r,p,[page('Methods\nrTMS was delivered at 10 Hz.')])
    service.refresh_search(r)
    candidate=service.get_suggestions(r,p,i)[0]
    service.decide_suggestion(r,p,candidate['id'],{'decision':'accepted'})
    feedback=deepcopy(service.snapshot(r)['feedback'])
    assert feedback[0]['rank']==1 and feedback[0]['scores']
    assert service.get_field(r,p,i)['status']=='unreviewed'
    service.update_item(r,i,{'positive_examples':['10 Hz stimulation']})
    assert service.get_paper(r,p)['search_status']=='stale'
    service.refresh_search(r)
    assert service.snapshot(r)['feedback']==feedback
    assert next(s for s in service.get_suggestions(r,p,i) if s['id']==candidate['id'])['decision']=='accepted'
    other=service.create_review({'name':'Other'})['id']
    with pytest.raises(KeyError): service.decide_suggestion(other,p,candidate['id'],{'decision':'rejected'})
    book=load_workbook(BytesIO(export_workbook(service.snapshot(r))))
    assert book['Retrieval feedback'].max_row>1 and book['Retrieval audit'].max_row>1
    book.close(); service.store.close()
    reopened=ReviewService(tmp_path)
    assert reopened.snapshot(r)['feedback']==feedback
    reopened.store.close()


def test_ruled_table_ingestion_has_cell_geometry(tmp_path):
    path=tmp_path/'table.pdf'
    pdf=canvas.Canvas(str(path))
    for x in (40,200,370): pdf.line(x,600,x,690)
    for y in (600,630,660,690): pdf.line(40,y,370,y)
    for x,y,t in [(50,670,'Outcome'),(210,670,'Mean (SD)'),(50,640,'LEMS'),(210,640,'42.1 (4.3)'),(50,610,'10MWT'),(210,610,'2.1 (0.7)')]: pdf.drawString(x,y,t)
    pdf.save()
    source=process_pdf(path)[0]
    assert source['tables']
    cells=source['tables'][0]['cells']
    cell=next(c for c in cells if c['text']=='42.1 (4.3)')
    assert cell['row_header']=='LEMS' and cell['column_header']=='Mean (SD)' and cell['boxes']
    assert source['text'][cell['start']:cell['end']]=='42.1 (4.3)'


def test_optional_ocr_keeps_offsets_boxes_and_confidence(monkeypatch):
    from review_extractor.ocr import ocr_page
    monkeypatch.setenv('SYSTEMATIC_OCR','1')
    data=dict(text=['rTIVIS','10','Hz'],left=[0,20,40],top=[10,10,10],width=[15,15,15],height=[10]*3,
              block_num=[1]*3,par_num=[1]*3,line_num=[1]*3,conf=[80,90,100])
    fake=SimpleNamespace(Output=SimpleNamespace(DICT='dict'),image_to_data=lambda *a,**k:data,get_tesseract_version=lambda:'test')
    monkeypatch.setitem(sys.modules,'pytesseract',fake)
    source=SimpleNamespace(to_image=lambda **k:SimpleNamespace(original=SimpleNamespace(width=200,height=200)))
    result=dict(width=100,height=100,status='ocr_required',error='No text')
    ocr_page(source,result)
    assert result['text']=='rTIVIS 10 Hz' and result['ocr'] and result['ocr_confidence']==pytest.approx(.9)
    assert result['words'][1]['x0']==10
    for w in result['words']: assert result['text'][w['start']:w['end']]==w['text']
