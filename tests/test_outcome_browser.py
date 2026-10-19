"""Optional browser journeys for repeated, source-linked outcome extraction."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from test_browser import browser_page, live_app, pytestmark, seed, select_extracted_quote
from test_retrieval import page as source_page


def seed_outcomes(api, service):
    review, paper, study_item = seed(api)
    response = api.post(f"/api/reviews/{review}/items", json={
        "name": "Baseline mean", "worksheet_name": "Meta Analysis",
        "section": "meta_analysis", "data_type": "number", "keywords": ["baseline"],
    })
    assert response.status_code == 201, response.text
    item = response.json()["id"]
    service.set_pages(review, paper, [source_page(
        "Results\n10MWT baseline gait speed was 0.4 m/s.\n\n"
        "6MWT baseline distance was 120 m."
    )])
    service.refresh_search(review, paper)
    return review, paper, study_item, item


def add_row(page, outcome, label):
    page.locator('#add-outcome').click()
    page.locator('#outcome-outcome').fill(outcome)
    page.locator('#outcome-group').fill('Active stimulation + gait training')
    page.locator('#outcome-time-point').fill('Baseline to end of treatment')
    page.locator('#outcome-label').fill(label)
    page.locator('#save-outcome').click()
    page.locator('#outcome-dialog').wait_for(state='hidden')


def approve_selection(page, quote, value):
    from playwright.sync_api import expect
    select_extracted_quote(page, quote)
    page.locator('#capture-selection').click()
    page.wait_for_function("document.querySelector('#evidence-dialog').open || (document.querySelector('#toast').classList.contains('error') && !document.querySelector('#toast').hidden)")
    assert page.locator('#evidence-dialog').is_visible(), page.locator('#toast').inner_text()
    page.locator('#accept-manual-evidence').click()
    page.locator('#evidence-dialog').wait_for(state='hidden')
    page.locator('#value-tab').click()
    expect(page.locator('#evidence-links blockquote')).to_contain_text(quote)
    page.locator('#field-value').fill(str(value))
    page.locator('#save-field').click()
    expect(page.locator('#field-progress')).to_contain_text('Approved value')


def test_outcome_rows_keep_sources_values_and_export_locations_independent(live_app, browser_page, tmp_path):
    from playwright.sync_api import expect
    url, api, service = live_app
    review, paper, _study_item, item = seed_outcomes(api, service)
    path = f'/api/reviews/{review}/papers/{paper}'
    page = browser_page
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#meta-analysis-tab').click()
    expect(page.locator('#outcome-summary')).to_contain_text('This section is optional')
    expect(page.locator('#reading-progress-summary')).to_have_text('No outcome row selected')
    expect(page.locator('#manual-evidence')).to_be_disabled()
    add_row(page, '10MWT', 'Walking speed')
    first = api.get(f'{path}/outcomes').json()[0]['id']
    expect(page.locator('#item-select')).to_have_value(item)
    page.locator('#text-toggle').click()
    approve_selection(page, '10MWT baseline gait speed was 0.4 m/s.', 0.4)

    add_row(page, '6MWT', 'Walking distance')
    second = next(row['id'] for row in api.get(f'{path}/outcomes').json() if row['id'] != first)
    expect(page.locator('#field-value')).to_have_value('')
    expect(page.locator('#evidence-links blockquote')).to_have_count(0)
    approve_selection(page, '6MWT baseline distance was 120 m.', 120)
    page.screenshot(path=str(tmp_path / 'outcome-workspace.png'), full_page=True)
    first_field = api.get(f'{path}/items/{item}/field', params={'outcome_id': first}).json()
    second_field = api.get(f'{path}/items/{item}/field', params={'outcome_id': second}).json()
    assert first_field['value'] == 0.4
    assert second_field['value'] == 120
    assert set(first_field['evidence_ids']).isdisjoint(second_field['evidence_ids'])

    # Exclusion is optional and keeps the saved decision available for later reuse.
    page.locator('#outcome-included').uncheck()
    expect(page.locator('#outcome-summary')).to_contain_text('1 selected')
    page.locator('[data-section="export"]').click()
    page.get_by_role('tab', name='Meta Analysis', exact=True).click()
    expect(page.locator('#export-table tbody tr')).to_have_count(1)
    page.locator('[data-section="extract"]').click()
    page.locator('#outcome-included').check()
    expect(page.locator('#outcome-summary')).to_contain_text('2 selected')
    page.locator('#outcome-select').select_option(first)
    expect(page.locator('#field-value')).to_have_value('0.4')
    expect(page.locator('#evidence-links blockquote')).to_have_text('10MWT baseline gait speed was 0.4 m/s.')
    page.reload()
    page.locator('[data-section="extract"]').click()
    expect(page.locator('#meta-analysis-tab')).to_have_attribute('aria-selected', 'true')
    expect(page.locator('#outcome-select')).to_have_value(first)
    page.locator('#value-tab').click()
    expect(page.locator('#field-value')).to_have_value('0.4')
    page.locator('[data-section="export"]').click()
    page.get_by_role('tab', name='Meta Analysis', exact=True).click()
    expect(page.locator('#export-table tbody tr')).to_have_count(2)
    with page.expect_download() as download:
        page.locator('#combined-export').click()
    workbook = load_workbook(BytesIO(download.value.path().read_bytes()))
    assert workbook.sheetnames == ['Data', 'Meta Analysis', 'Verify - Data', 'Verify - Meta Analysis']
    sheet = workbook['Meta Analysis']
    heading = [cell.value for cell in sheet[1]].index('Baseline mean') + 1
    assert [sheet.cell(row, heading).value for row in (2, 3)] == [0.4, 120]
    verification = workbook['Verify - Meta Analysis']
    assert all('Page 1' in verification.cell(row, heading).value for row in (2, 3))
    assert verification.cell(2, heading).value != verification.cell(3, heading).value
    assert not errors


def test_unsaved_outcome_values_are_guarded_and_removed_rows_restore(live_app, browser_page):
    from playwright.sync_api import expect
    url, api, service = live_app
    review, paper, _study_item, item = seed_outcomes(api, service)
    path = f'/api/reviews/{review}/papers/{paper}/outcomes'
    first = api.post(path, json={'outcome': '10MWT', 'label': 'Speed'}).json()['id']
    second = api.post(path, json={'outcome': '6MWT', 'label': 'Distance'}).json()['id']
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#meta-analysis-tab').click()
    page.locator('#value-tab').click()
    page.locator('#field-value').fill('0.5')
    page.once('dialog', lambda dialog: dialog.dismiss())
    page.locator('#outcome-select').select_option(second)
    expect(page.locator('#outcome-select')).to_have_value(first)
    expect(page.locator('#field-value')).to_have_value('0.5')
    page.once('dialog', lambda dialog: dialog.dismiss())
    page.locator('#study-extraction-tab').click()
    expect(page.locator('#meta-analysis-tab')).to_have_attribute('aria-selected', 'true')
    expect(page.locator('#field-value')).to_have_value('0.5')
    page.once('dialog', lambda dialog: dialog.accept())
    page.locator('#outcome-select').select_option(second)
    expect(page.locator('#field-value')).to_have_value('')
    page.locator('#edit-outcome').click()
    page.once('dialog', lambda dialog: dialog.accept())
    page.locator('#delete-outcome').click()
    page.locator('#outcome-dialog').wait_for(state='hidden')
    expect(page.locator('#outcome-select option')).to_have_count(1)
    page.locator('#deleted-outcomes-details summary').click()
    page.get_by_role('button', name='Restore Distance', exact=True).click()
    expect(page.locator('#outcome-select option')).to_have_count(2)
    expect(page.locator('#outcome-select')).to_have_value(second)
    assert api.get(f'/api/reviews/{review}/papers/{paper}/items/{item}/field', params={'outcome_id': first}).json()['status'] == 'unreviewed'


def test_template_preview_classifies_meta_analysis_and_preserves_study_items(live_app, browser_page, tmp_path):
    from playwright.sync_api import expect
    url, api, _service = live_app
    review = api.post('/api/reviews', json={'name': 'Outcome template'}).json()['id']
    workbook = Workbook()
    workbook.active.title = 'Systematic Review'
    workbook.active.append(['Study', 'Protocol'])
    workbook.active.append(['EXAMPLE', 'DO NOT IMPORT'])
    workbook.create_sheet('Meta Analysis').append(['Study', 'Outcome', 'Baseline mean'])
    path = tmp_path / 'outcome-template.xlsx'
    workbook.save(path)
    page = browser_page
    page.goto(url)
    page.locator('#template-file').set_input_files(path)
    page.locator('#template-form button[type=submit]').click()
    expect(page.locator('#template-items tr')).to_have_count(5)
    assert page.locator('#template-items [data-property="section"]').evaluate_all('(nodes)=>nodes.map(node=>node.value)') == ['study', 'study', 'meta_analysis', 'meta_analysis', 'meta_analysis']
    page.locator('#import-items').click()
    expect(page.locator('#template-preview')).to_be_hidden()
    items = api.get(f'/api/reviews/{review}/items').json()
    assert [item['section'] for item in items] == ['study', 'study', 'meta_analysis', 'meta_analysis', 'meta_analysis']
    assert _service.export_snapshot(review)['fields'] == []


def test_outcome_switch_blocks_stale_controls_and_ignores_previous_save(live_app, browser_page):
    from playwright.sync_api import expect
    url, api, service = live_app
    review, paper, _study_item, item = seed_outcomes(api, service)
    path = f'/api/reviews/{review}/papers/{paper}'
    first = api.post(path + '/outcomes', json={'outcome': '10MWT', 'label': 'Speed'}).json()['id']
    second = api.post(path + '/outcomes', json={'outcome': '6MWT', 'label': 'Distance'}).json()['id']
    field_path = path + f'/items/{item}/field'
    page = browser_page
    page.goto(url)
    page.locator('[data-section="extract"]').click()
    page.locator('#meta-analysis-tab').click()
    expect(page.locator('.evidence-panel')).not_to_have_attribute('aria-busy', 'true')
    page.locator('#value-tab').click()
    expect(page.locator('#field-value')).to_have_value('')

    pending = []
    def delay_second_field(route):
        if route.request.method == 'GET' and f'outcome_id={second}' in route.request.url:
            pending.append(route)
        else:
            route.continue_()
    page.route('**/items/*/field?*', delay_second_field)
    page.locator('#outcome-select').select_option(second)
    expect(page.locator('.evidence-panel')).to_have_attribute('aria-busy', 'true')
    expect(page.locator('#field-form')).to_be_hidden()
    expect(page.locator('#suggestions .decision-actions button')).to_have_count(0)
    expect(page.locator('#manual-evidence')).to_be_disabled()
    page.wait_for_timeout(100)
    assert pending
    pending.pop().continue_()
    expect(page.locator('.evidence-panel')).not_to_have_attribute('aria-busy', 'true')
    page.unroute('**/items/*/field?*', delay_second_field)

    # Send an old-row save, move to the other row, then deliver its delayed response.
    page.locator('#outcome-select').select_option(first)
    expect(page.locator('.evidence-panel')).not_to_have_attribute('aria-busy', 'true')
    page.locator('#field-status').select_option('not_reported')
    page.locator('#field-rationale').fill('Synthetic fixture reviewed; baseline unavailable.')
    saved_responses = []
    def delay_save(route):
        if route.request.method == 'PUT':
            saved_responses.append((route, route.fetch()))
        else:
            route.continue_()
    page.route('**/items/*/field?*', delay_save)
    page.locator('#save-field').click()
    expect(page.locator('#save-field')).to_be_disabled()
    page.once('dialog', lambda dialog: dialog.accept())
    page.locator('#outcome-select').select_option(second)
    expect(page.locator('.evidence-panel')).not_to_have_attribute('aria-busy', 'true')
    expect(page.locator('#field-status')).to_have_value('unreviewed')
    assert saved_responses
    route, response = saved_responses.pop()
    route.fulfill(response=response)
    page.wait_for_timeout(150)
    expect(page.locator('#outcome-select')).to_have_value(second)
    expect(page.locator('#field-status')).to_have_value('unreviewed')
    expect(page.locator('#field-rationale')).to_have_value('')
    assert api.get(field_path, params={'outcome_id': first}).json()['status'] == 'not_reported'
    assert api.get(field_path, params={'outcome_id': second}).json()['status'] == 'unreviewed'


def test_same_worksheet_name_keeps_study_and_outcome_previews_accessible(live_app, browser_page):
    from playwright.sync_api import expect
    url, api, service = live_app
    review, paper, study_item, item = seed_outcomes(api, service)
    api.patch(f'/api/reviews/{review}/items/{study_item}', json={'worksheet_name': 'Shared'})
    api.patch(f'/api/reviews/{review}/items/{item}', json={'worksheet_name': 'Shared'})
    for label in ('Speed', 'Distance'):
        api.post(f'/api/reviews/{review}/papers/{paper}/outcomes', json={'outcome': label, 'label': label})
    page = browser_page
    page.goto(url)
    page.locator('[data-section="export"]').click()
    page.get_by_role('tab', name='Shared (studies)', exact=True).click()
    expect(page.locator('#export-table tbody tr')).to_have_count(1)
    page.get_by_role('tab', name='Shared (outcomes)', exact=True).click()
    expect(page.locator('#export-table tbody tr')).to_have_count(2)
    expect(page.locator('#export-table th')).to_have_text('Baseline mean')
