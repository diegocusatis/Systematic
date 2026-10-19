"""Synthetic backend lifecycle checks for manual candidates and reviewer learning."""
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from review_extractor.hybrid.semantic import LocalSemantic
from review_extractor.service import ReviewService
from review_extractor.workbooks import export_workbook
from test_hybrid import Encoder
from test_retrieval import item, page


TEXT = "Participants were recruited and\naged 48.2 ± 9.6 years."


@pytest.fixture
def project(tmp_path):
    service = ReviewService(tmp_path)
    service.semantic = LocalSemantic(tmp_path / "synthetic-cache", encoder=Encoder())
    review = service.create_review({"name": "Synthetic feedback", "reviewer": "Synthetic reviewer"})["id"]
    field = service.add_item(review, item("Mean age"))["id"]
    paper = add_paper(service, review, TEXT)
    yield service, review, field, paper
    service.store.close()


def add_paper(service, review, text, **data):
    paper = service.add_paper(review, dict(title="Synthetic document", processing_status="ready", **data))["id"]
    service.set_pages(review, paper, [page(text)])
    return paper


def selection(service, review, paper, quote=TEXT, start=0, **data):
    return dict(page_number=1, quote=quote, start=start, end=start + len(quote),
                document_version=service.get_paper(review, paper)["document_version"], **data)


def candidate(service, review, field, paper, **data):
    return service.add_candidate(review, paper, field, selection(service, review, paper, **data))


def example(service, review, identity):
    return next(e for e in service.get_learning(review)["examples"] if e["example_id"] == identity)


def test_pending_candidate_has_exact_source_and_no_evidence_or_approval(project):
    service, review, field, paper = project
    before_field = deepcopy(service.get_field(review, paper, field))
    row = candidate(service, review, field, paper)
    assert row["decision"] == "pending" and row["origin"] == "user"
    assert row["document_version"] == 1 and row["boxes"]
    assert row["quote"] == TEXT and row["start"] == 0 and row["end"] == len(TEXT)
    assert len({b["top"] for b in row["boxes"]}) == 2
    assert service.get_evidence(review, paper, field) == []
    assert service.get_field(review, paper, field) == before_field
    assert not example(service, review, "s:" + row["id"])["eligible"]
    service.refresh_search(review, paper)
    kept = next(s for s in service.get_suggestions(review, paper, field) if s["id"] == row["id"])
    assert kept["origin"] == "user" and kept["active"] and kept["decision"] == "pending"


def test_exact_algorithm_overlap_keeps_reference_and_single_effective_example(project):
    service, review, field, paper = project
    service.refresh_search(review, paper)
    original = next(s for s in service.get_suggestions(review, paper, field) if s["quote"] == TEXT)
    row = candidate(service, review, field, paper, decision="accepted")
    assert row["id"] == original["id"] and row["origin"] == "user"
    assert service.get_field(review, paper, field)["status"] == "unreviewed"
    evidence = service.get_evidence(review, paper, field)
    assert len(evidence) == 1 and evidence[0]["suggestion_id"] == original["id"]
    before_events = deepcopy(service.snapshot(review)["feedback"])
    for _ in range(3):
        assert candidate(service, review, field, paper, decision="accepted")["id"] == row["id"]
        service.decide_suggestion(review, paper, row["id"], {"decision": "accepted"})
    assert service.snapshot(review)["feedback"] == before_events
    assert len(service.get_learning(review)["examples"]) == 1
    assert example(service, review, "s:" + row["id"])["eligible"]
    service.refresh_search(review, paper)
    assert service.snapshot(review)["feedback"] == before_events
    assert next(s for s in service.get_suggestions(review, paper, field) if s["id"] == row["id"])["decision"] == "accepted"


def test_overlapping_manual_spans_preserve_distinct_stable_references(project):
    service, review, field, paper = project
    whole = candidate(service, review, field, paper, decision="accepted")
    quote = "aged 48.2 ± 9.6 years."
    partial = service.add_candidate(review, paper, field, selection(service, review, paper, quote, TEXT.index(quote)))
    assert partial["id"] != whole["id"]
    assert whole["id"] in partial["overlapping_candidate_ids"]
    service.refresh_search(review, paper)
    rows = {s["id"]: s for s in service.get_suggestions(review, paper, field)}
    assert rows[whole["id"]]["decision"] == "accepted"
    assert rows[partial["id"]]["origin"] == "user"
    assert service.get_evidence(review, paper, field)[0]["suggestion_id"] == whole["id"]


def test_decision_reversal_reason_and_pending_derive_one_current_label(project):
    service, review, field, paper = project
    row = candidate(service, review, field, paper, decision="accepted")
    identity = "s:" + row["id"]
    for decision, reason in [("rejected", "wrong_statistic"), ("pending", ""), ("accepted", "")]:
        changed = service.decide_suggestion(review, paper, row["id"], dict(decision=decision, rejection_reason=reason))
        current = example(service, review, identity)
        assert current["decision"] == decision and current["rejection_reason"] == reason
        assert current["eligible"] == (decision != "pending")
        assert changed["evidence_id"] == row["evidence_id"]
    assert [e["decision"] for e in service.snapshot(review)["feedback"]] == ["accepted", "rejected", "pending", "accepted"]
    assert len(service.get_learning(review)["examples"]) == 1
    assert service.get_field(review, paper, field)["status"] == "unreviewed"
    before = deepcopy(service.snapshot(review))
    with pytest.raises(ValueError, match="rejection reason"):
        service.decide_suggestion(review, paper, row["id"], {"decision": "rejected", "rejection_reason": "invented"})
    assert service.snapshot(review) == before


def test_standalone_evidence_uses_same_history_without_automatic_approval(project):
    service, review, field, paper = project
    row = service.add_evidence(review, paper, field, selection(service, review, paper))
    identity = "e:" + row["id"]
    assert row["suggestion_id"] is None and example(service, review, identity)["eligible"]
    service.update_evidence(review, paper, row["id"], {"decision": "rejected", "rejection_reason": "irrelevant"})
    assert example(service, review, identity)["decision"] == "rejected"
    service.update_evidence(review, paper, row["id"], {"decision": "pending"})
    assert not example(service, review, identity)["eligible"]
    service.update_evidence(review, paper, row["id"], {"decision": "accepted"})
    assert example(service, review, identity)["eligible"]
    assert len(service.get_learning(review)["examples"]) == 1
    assert service.get_field(review, paper, field)["status"] == "unreviewed"


def test_legacy_manual_import_does_not_treat_notes_as_a_new_acceptance(project, monkeypatch):
    service, review, field, paper = project
    # Earlier application versions saved standalone evidence without feedback.
    with monkeypatch.context() as legacy:
        legacy.setattr(service, "_save_feedback", lambda *args, **kwargs: None)
        evidence = service.add_evidence(review, paper, field, selection(service, review, paper))
    assert service.snapshot(review)["feedback"] == []
    service.update_item(review, field, {"definition": "Only the intervention arm at baseline"})
    service.update_evidence(review, paper, evidence["id"], {"notes": "An annotation does not confirm this new definition"})
    service.update_learning(review, {"enabled": True})
    current = example(service, review, "e:" + evidence["id"])
    assert not current["eligible"]
    assert "Item meaning changed" in current["excluded_reason"]


@pytest.mark.parametrize("standalone", [False, True])
def test_item_meaning_changes_require_explicit_reconsideration_not_notes(project, standalone):
    service, review, field, paper = project
    if standalone:
        evidence = service.add_evidence(review, paper, field, selection(service, review, paper))
        identity = "e:" + evidence["id"]
    else:
        row = candidate(service, review, field, paper, decision="accepted")
        evidence = next(e for e in service.get_evidence(review, paper, field) if e["id"] == row["evidence_id"])
        identity = "s:" + row["id"]
    before = deepcopy(service.snapshot(review)["feedback"])
    service.update_item(review, field, {"definition": "Participant age in the intervention arm"})
    assert "Item meaning changed" in example(service, review, identity)["excluded_reason"]
    service.update_evidence(review, paper, evidence["id"], {"notes": "Clarifying annotation only"})
    assert service.snapshot(review)["feedback"] == before
    assert not example(service, review, identity)["eligible"]
    service.update_evidence(review, paper, evidence["id"], {"decision": "accepted"})
    assert example(service, review, identity)["eligible"]
    assert len(service.snapshot(review)["feedback"]) == len(before) + 1
    # Returning to an older definition is a further meaning revision, not a way
    # to resurrect previously invalidated examples implicitly.
    service.update_item(review, field, {"definition": ""})
    assert not example(service, review, identity)["eligible"]


def test_reset_exclude_restore_toggle_preserve_decisions_values_and_history(project):
    service, review, field, paper = project
    row = candidate(service, review, field, paper, decision="accepted")
    identity = "s:" + row["id"]
    service.save_field(review, paper, field, {"status": "approved", "value": "48.2", "evidence_ids": [row["evidence_id"]]})
    before = deepcopy(service.snapshot(review))
    enabled = service.update_learning(review, {"enabled": True})
    assert enabled["config"]["enabled"] and enabled["config"]["version"] > 1
    version = enabled["config"]["version"]
    excluded = service.exclude_learning_example(review, identity, {"excluded": True})
    assert excluded["config"]["version"] > version
    assert not example(service, review, identity)["eligible"]
    service.exclude_learning_example(review, identity, {"excluded": False})
    assert example(service, review, identity)["eligible"]
    old_epoch = service.get_learning(review)["config"]["epoch"]
    service.update_learning(review, {"reset": True})
    assert service.get_learning(review)["config"]["epoch"] == old_epoch + 1
    assert "reset" in example(service, review, identity)["excluded_reason"]
    service.update_learning(review, {"enabled": False})
    service.update_learning(review, {"enabled": True})
    assert not example(service, review, identity)["eligible"]
    after = service.snapshot(review)
    for key in ("fields", "evidence", "feedback"):
        assert after[key] == before[key]
    assert {h["id"] for h in before["history"]} <= {h["id"] for h in after["history"]}
    service.decide_suggestion(review, paper, row["id"], {"decision": "accepted"})
    assert example(service, review, identity)["eligible"]


def test_source_reprocessing_preserves_old_reference_and_excludes_learning(project):
    service, review, field, paper = project
    row = candidate(service, review, field, paper, decision="accepted")
    identity = "s:" + row["id"]
    service.save_field(review, paper, field, {"status": "approved", "value": "48.2", "evidence_ids": [row["evidence_id"]]})
    service.update_learning(review, {"enabled": True})
    old_pages = deepcopy(service.get_pages(review, paper))
    old_events = deepcopy(service.snapshot(review)["feedback"])
    service.set_pages(review, paper, [page("Reprocessed source\n" + TEXT)])
    assert not example(service, review, identity)["eligible"]
    assert "Source version changed" in example(service, review, identity)["excluded_reason"]
    saved = service.get_field(review, paper, field)
    assert saved["status"] == "approved" and saved["value"] == "48.2" and saved["needs_reconsideration"]
    assert service.get_evidence(review, paper, field)[0]["source_stale"]
    with pytest.raises(ValueError, match="reprocessed"):
        service.decide_suggestion(review, paper, row["id"], {"decision": "accepted"})
    service.refresh_search(review, paper)
    historical = next(s for s in service.get_suggestions(review, paper, field) if s["id"] == row["id"])
    assert historical["decision"] == "accepted" and not historical["active"]
    assert historical["document_version"] == 1 and historical["quote"] == TEXT
    assert service.snapshot(review)["feedback"] == old_events
    revision = next(h for h in service.snapshot(review)["history"] if h["action"] == "document_source_changed")
    assert revision["before"]["pages"] == old_pages


def test_repeated_quotation_needs_offsets_and_current_document_version(project):
    service, review, field, paper = project
    text = TEXT + "\n\n" + TEXT
    service.set_pages(review, paper, [page(text)])
    version = service.get_paper(review, paper)["document_version"]
    with pytest.raises(ValueError, match="more than once"):
        service.add_candidate(review, paper, field, dict(page_number=1, quote=TEXT, document_version=version))
    second = service.add_candidate(review, paper, field, selection(service, review, paper, TEXT, text.rfind(TEXT)))
    first = candidate(service, review, field, paper)
    assert first["id"] != second["id"] and first["start"] == 0 and second["start"] == text.rfind(TEXT)
    assert second["quote"] == text[second["start"]:second["end"]]
    for invalid in (None, True, version - 1):
        with pytest.raises(ValueError, match="current version"):
            service.add_candidate(review, paper, field, dict(selection(service, review, paper), document_version=invalid))
    with pytest.raises(ValueError, match="exactly match"):
        service.add_candidate(review, paper, field, dict(selection(service, review, paper), start=1))


def test_reviews_and_items_isolate_selection_feedback_and_learning(project):
    service, review, field, paper = project
    row = candidate(service, review, field, paper, decision="accepted")
    other_review = service.create_review({"name": "Another synthetic review"})["id"]
    other_item = service.add_item(other_review, item("Mean age"))["id"]
    other_paper = add_paper(service, other_review, TEXT)
    same_review_item = service.add_item(review, item("Separate age field"))["id"]
    assert service.get_evidence(review, paper, same_review_item) == []
    assert service.get_learning(other_review)["examples"] == []
    for action in [
        lambda: service.add_candidate(other_review, paper, other_item, selection(service, review, paper)),
        lambda: service.add_candidate(review, paper, other_item, selection(service, review, paper)),
        lambda: service.decide_suggestion(other_review, other_paper, row["id"], {"decision": "accepted"}),
        lambda: service.exclude_learning_example(other_review, "s:" + row["id"], {"excluded": True}),
    ]:
        with pytest.raises(KeyError): action()


def test_service_learning_uses_other_documents_and_only_the_same_item(project):
    service, review, field, paper = project
    candidate(service, review, field, paper, decision="accepted")
    text2 = "The enrolled participants had a mean age of 39 years."
    paper2 = add_paper(service, review, text2)
    service.add_candidate(review, paper2, field, selection(service, review, paper2, text2))
    accepted2 = service.get_suggestions(review, paper2, field)[0]
    service.decide_suggestion(review, paper2, accepted2["id"], {"decision": "accepted"})
    another_item = service.add_item(review, item("Median age"))["id"]
    target = add_paper(service, review, "The mean age was 52 years.\n\nThe median age was 50 years.")
    service.update_learning(review, {"enabled": True})
    service.refresh_search(review, target)
    assert any(s["scores"]["feedback_score"] > 0 for s in service.get_suggestions(review, target, field))
    assert all(s["learning"]["status"] == "sparse_feedback" for s in service.get_suggestions(review, target, another_item))
    service.refresh_search(review, paper)
    assert all(s.get("learning", {}).get("status") == "sparse_feedback" for s in service.get_suggestions(review, paper, field) if "learning" in s)
    duplicate = add_paper(service, review, TEXT, sha256="different-container-bytes")
    service.refresh_search(review, duplicate)
    assert all(s["learning"]["status"] == "sparse_feedback" for s in service.get_suggestions(review, duplicate, field))


def test_manual_transcription_remains_evidence_but_cannot_train(project):
    service, review, field, paper = project
    service.set_pages(review, paper, [dict(page(""), status="ocr_required")])
    data = dict(page_number=1, quote="A reviewer transcribed the participant age.", manual=True, rationale="Unreadable scan")
    evidence = service.add_evidence(review, paper, field, data)
    current = example(service, review, "e:" + evidence["id"])
    assert evidence["decision"] == "accepted" and not current["eligible"]
    assert "No verified searchable source span" == current["excluded_reason"]


def test_export_keeps_current_design_and_text_safe_feedback_audit(project):
    service, review, field, paper = project
    row = candidate(service, review, field, paper, decision="accepted", notes="=untrusted note")
    service.save_field(review, paper, field, {"status": "approved", "value": "48.2", "evidence_ids": [row["evidence_id"]]})
    service.update_item(review, field, {"export_heading": "Participant age"})
    service.update_learning(review, {"enabled": True})
    snapshot = deepcopy(service.snapshot(review))
    book = load_workbook(BytesIO(export_workbook(snapshot)))
    assert "Participant age" in [cell.value for cell in book["Extraction DRAFT"][1]]
    assert book["Retrieval feedback"].max_row > 1
    feedback_headers = [cell.value for cell in book["Retrieval feedback"][1]]
    assert {"Decision", "Rejection reason", "Learning epoch", "Feedback sequence", "Item meaning version"} <= set(feedback_headers)
    audit_headers = [cell.value for cell in book["Retrieval audit"][1]]
    assert {"Origin", "Reviewer feedback adjustment"} <= set(audit_headers)
    assert book["Learning examples"].max_row == len(snapshot["learning_examples"]) + 1
    assert "Learning configuration" in [row[0].value for row in book["Export metadata"]]
    assert any("=untrusted note" in str(cell.value) for sheet in book for row in sheet for cell in row)
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row if "=untrusted note" in str(cell.value))
    assert service.snapshot(review) == snapshot
    book.close()


def test_export_feedback_quotes_are_text_and_foreign_records_are_filtered(project):
    service, review, field, _ = project
    text = '=SUM(1,2) The participants had a mean age of 44 years.'
    paper = add_paper(service, review, text)
    service.add_candidate(review, paper, field, selection(service, review, paper, text, decision="accepted"))
    snapshot = service.snapshot(review)
    for table, text_key in [("feedback", "candidate_text"), ("learning_examples", "quote")]:
        base = deepcopy(snapshot[table][0])
        for scope in ("review", "paper", "item"):
            foreign = dict(base, **{scope + "_id": "outside-scope"})
            foreign[text_key] = f"FOREIGN-{table}-{scope}"
            snapshot[table].append(foreign)
    book = load_workbook(BytesIO(export_workbook(snapshot)))
    for sheet_name, header in [("Retrieval feedback", "Candidate text"), ("Learning examples", "Quote")]:
        sheet = book[sheet_name]
        column = [c.value for c in sheet[1]].index(header) + 1
        cells = [sheet.cell(row, column) for row in range(2, sheet.max_row + 1)]
        assert len(cells) == 1 and cells[0].value == text and cells[0].data_type == "s"
    assert all("FOREIGN-" not in str(cell.value) for sheet in book for row in sheet for cell in row)
    book.close()
