"""Data-preservation and failure tests use synthetic, isolated review folders."""
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3

from openpyxl import load_workbook
import pytest

from review_extractor.storage import Store
from review_extractor.service import ReviewService
from review_extractor.workbooks import export_download_workbook
from test_meta_workbooks import outcome_snapshot, template_bytes
from windows import installer as setup
from windows.validate import validate


@pytest.fixture(autouse=True)
def isolated_system_dependencies(tmp_path, monkeypatch):
    # Discovery tests must not find OCR installed on the developer's computer.
    for variable in ("LOCALAPPDATA", "ProgramFiles", "ProgramFiles(x86)"):
        monkeypatch.setenv(variable, str(tmp_path / variable))
    monkeypatch.delenv("SYSTEMATIC_TESSERACT", raising=False)


def seeded(folder):
    folder.mkdir(parents=True)
    sample = outcome_snapshot()
    sample["review"].update(created_at="2026-01-01T00:00:00Z",
                            updated_at="2026-01-01T00:00:00Z",
                            schema_version=1, rule_version=1, settings={})
    store = Store(folder / setup.DB)
    store.put("reviews", sample["review"])
    for table in setup.TABLES[1:]:
        for record in sample.get(table, []):
            record = deepcopy(record)
            if table == "items":
                record["keywords_initialized"] = True
            if table == "papers":
                pdf = folder / "pdfs" / record["review_id"] / record["filename"]
                pdf.parent.mkdir(parents=True, exist_ok=True)
                pdf.write_bytes(b"%PDF-1.4\n% synthetic fixture " + record["id"].encode())
                record.update(path=str(pdf), sha256=setup.digest(pdf), created_at="2026-01-01T00:00:00Z",
                              citation="", url="", processing_error="", manual_review_ack=False,
                              processing_progress=100, page_count=1, search_version=1,
                              pdf_metadata=None, retrieval_engine_version="", removed=False)
            if table == "pages":
                record["review_id"] = sample["review"]["id"]
            store.put(table, record)
    store.put("history", dict(id="h-1", review_id="review-1", action="synthetic", quote="Exact source quote"))
    store.close()
    (folder / "original-template.xlsx").write_bytes(template_bytes())
    return sample


def exported(folder):
    service = ReviewService(folder)
    try:
        result = []
        for view in ("data", "verification", "combined"):
            book = load_workbook(BytesIO(export_download_workbook(service.export_snapshot("review-1"), view)))
            result.append([(sheet.title, list(sheet.values),
                            [(cell.coordinate, cell.comment.text if cell.comment else None,
                              str(cell.font), str(cell.fill)) for row in sheet for cell in row],
                            [(key, dim.width) for key, dim in sheet.column_dimensions.items()])
                           for sheet in book])
        return result
    finally:
        service.store.close()


def test_transfer_preserves_records_pdfs_templates_and_all_export_sheets(tmp_path):
    source = tmp_path / "old"
    seeded(source)
    before = setup.records(source / setup.DB)
    before_exports = exported(source)
    final = tmp_path / "appdata"
    work = tmp_path / "work"
    work.mkdir()
    extra = tmp_path / "external export.xlsx"
    extra.write_bytes(template_bytes())
    stage, report = setup.prepare_data(final, source, work, final, [extra])
    assert report["reviews"] == 1 and report["pdfs"] == 2
    after = setup.records(stage / setup.DB)
    for table in setup.TABLES:
        if table != "papers":
            assert before[table] == after[table]
    for row in after["papers"]:
        paper = json.loads(row[-1])
        assert Path(paper["path"]).is_relative_to(final)
        copied = stage / Path(paper["path"]).relative_to(final)
        assert setup.digest(copied) == paper["sha256"]
    assert setup.records(source / setup.DB) == before
    assert list(stage.rglob("original-template.xlsx"))[0].read_bytes() == (source / "original-template.xlsx").read_bytes()
    assert list((stage / "original-workbooks").rglob("*.xlsx"))[0].read_bytes() == extra.read_bytes()
    assert exported(stage) == before_exports
    validate(stage, final)
    assert setup.records(stage / setup.DB) == after


def test_missing_pdf_stops_without_touching_source_or_existing_data(tmp_path):
    source = tmp_path / "old"
    seeded(source)
    list((source / "pdfs").rglob("*.pdf"))[0].unlink()
    before = setup.digest(source / setup.DB)
    with pytest.raises(ValueError, match="Missing or changed PDF"):
        setup.prepare_data(tmp_path / "target", source, tmp_path / "work", tmp_path / "target")
    assert setup.digest(source / setup.DB) == before
    assert not (tmp_path / "target").exists()


def test_pdf_relocated_from_other_computer_by_stored_hash(tmp_path):
    source = tmp_path / "old"
    seeded(source)
    with setup.connect_write(source / setup.DB) as db:
        for rid, text in db.execute("SELECT id,data FROM papers").fetchall():
            data = json.loads(text)
            data["path"] = "Z:\\AnotherComputer\\.review-data\\pdfs\\review-1\\" + data["filename"]
            db.execute("UPDATE papers SET data=? WHERE id=?", (json.dumps(data), rid))
    final = tmp_path / "final"
    stage, report = setup.prepare_data(final, source, tmp_path / "work", final)
    assert report["pdfs"] == 2


def test_conflicting_approved_work_never_overwrites_destination(tmp_path):
    source, current = tmp_path / "old", tmp_path / "current"
    seeded(source)
    seeded(current)
    with setup.connect_write(current / setup.DB) as db:
        rid, text = db.execute("SELECT id,data FROM fields LIMIT 1").fetchone()
        data = json.loads(text)
        data["value"] = "Different approved work"
        db.execute("UPDATE fields SET data=? WHERE id=?", (json.dumps(data), rid))
    before = setup.digest(current / setup.DB)
    with pytest.raises(ValueError, match="different saved work"):
        setup.prepare_data(current, source, tmp_path / "work", current)
    assert setup.digest(current / setup.DB) == before


def test_identical_duplicate_records_import_once(tmp_path):
    source, current = tmp_path / "old", tmp_path / "current"
    seeded(source)
    seeded(current)
    stage, report = setup.prepare_data(current, source, tmp_path / "work", current)
    assert len(setup.records(stage / setup.DB)["reviews"]) == 1
    assert report["pdfs"] == 2


def test_activation_rolls_back_after_failed_state_write(tmp_path, monkeypatch):
    home, data, stage = tmp_path / "program", tmp_path / "data", tmp_path / "stage"
    home.mkdir()
    data.mkdir()
    stage.mkdir()
    (data / "original").write_text("approved work")
    (stage / "new").write_text("new version")
    setup.atomic_json(home / "current.json", {"installation_id": "old"})
    original = setup.atomic_json
    def fail_current(path, value):
        if Path(path).name == "current.json":
            raise OSError("Simulated disk failure")
        original(path, value)
    monkeypatch.setattr(setup, "atomic_json", fail_current)
    with pytest.raises(OSError, match="disk failure"):
        setup.activate(home, data, stage, {"installation_id": "new"})
    assert (data / "original").read_text() == "approved work"
    assert setup.read_json(home / "current.json")["installation_id"] == "old"
    assert not (home / "pending.json").exists()


def test_successful_activation_keeps_backup_and_handles_commit_interruption(tmp_path):
    home, data, stage = tmp_path / "program", tmp_path / "data", tmp_path / "stage"
    home.mkdir()
    data.mkdir()
    stage.mkdir()
    (data / "original").write_text("old")
    (stage / "new").write_text("new")
    state = {"installation_id": "release2"}
    backup = setup.activate(home, data, stage, state)
    assert (backup / "data" / "original").exists()
    assert (data / "new").exists()
    setup.atomic_json(home / "pending.json", dict(data=str(data), expected_data=str(data),
                      backup=str(backup / "data"), had_data=True, installation_id="release2"))
    setup.recover(home)
    assert (data / "new").exists()
    assert (backup / "data" / "original").exists()


def test_running_installation_is_locked(tmp_path):
    with setup.installation_lock(tmp_path):
        with pytest.raises(RuntimeError, match="running"):
            with setup.installation_lock(tmp_path):
                pass


@pytest.mark.skipif(setup.os.name != "nt", reason="Windows legacy database handle check")
def test_open_legacy_database_is_detected(tmp_path):
    seeded(tmp_path / "old")
    db = sqlite3.connect(tmp_path / "old" / setup.DB)
    db.execute("SELECT * FROM reviews").fetchall()
    try:
        with pytest.raises(RuntimeError, match="Close Systematic"):
            setup.require_closed(tmp_path / "old")
    finally:
        db.close()
    setup.require_closed(tmp_path / "old")


def test_unsupported_python_stops_before_creating_installation(tmp_path, monkeypatch):
    payload = tmp_path / "payload"
    payload.mkdir()
    setup.atomic_json(payload / "release.json", {"supported_python": ["2.7"], "architecture": "64bit"})
    with pytest.raises(ValueError, match="supported"):
        setup.install(payload, tmp_path / "home", tmp_path / "data")
    assert not (tmp_path / "home").exists()
    assert not (tmp_path / "data").exists()


def test_dependency_failure_preserves_active_version(tmp_path, monkeypatch):
    source = tmp_path / "source"
    seeded(source)
    home, data, payload = tmp_path / "program", tmp_path / "data", tmp_path / "payload"
    home.mkdir()
    payload.mkdir()
    setup.atomic_json(home / "current.json", {"installation_id": "old"})
    setup.atomic_json(payload / "release.json", {"supported_python": ["3.12", "3.11"],
                      "architecture": "64bit", "version": "0.1.0", "data_format": 1})
    def fail(*args, **kwargs):
        raise RuntimeError("Simulated download failure")
    monkeypatch.setattr(setup, "execute", fail)
    with pytest.raises(RuntimeError, match="download failure"):
        setup.install(payload, home, data, source)
    assert setup.read_json(home / "current.json") == {"installation_id": "old"}
    assert not data.exists()
    assert setup.records(source / setup.DB)["reviews"]


def test_disjoint_reviews_are_imported_together(tmp_path):
    source, current = tmp_path / "old", tmp_path / "current"
    seeded(source)
    service = ReviewService(current)
    review = service.create_review({"name": "Existing separate review"})
    service.store.close()
    stage, report = setup.prepare_data(current, source, tmp_path / "work", current)
    assert report["reviews"] == 2
    assert {r[0] for r in setup.records(stage / setup.DB)["reviews"]} == {"review-1", review["id"]}


def test_blank_installation_prepares_and_validates_empty_database(tmp_path):
    final = tmp_path / "data"
    stage, report = setup.prepare_data(final, None, tmp_path / "work", final)
    validate(stage, final)
    setup.db_check(stage / setup.DB)
    assert setup.records(stage / setup.DB)["reviews"] == []
    assert report["reviews"] == report["pdfs"] == 0
    assert not final.exists()


def test_preservation_fixture_is_a_valid_startable_application(tmp_path):
    from review_extractor.app import create_app
    seeded(tmp_path / "reviews")
    app = create_app(tmp_path / "reviews", run_background=False)
    try:
        app.state.scans.recover()
        assert len(app.state.service.list_reviews()) == 1
        for paper in app.state.service.list_papers("review-1"):
            app.state.service._validate_paper(paper)
    finally:
        app.state.scans.close()
        app.state.service.store.close()


def test_new_history_on_duplicate_review_is_a_conflict(tmp_path):
    source, current = tmp_path / "old", tmp_path / "current"
    seeded(source)
    seeded(current)
    store = Store(source / setup.DB)
    store.put("history", dict(id="new-event", review_id="review-1", action="changed"))
    store.close()
    with pytest.raises(ValueError, match="different saved work"):
        setup.prepare_data(current, source, tmp_path / "work", current)


def test_data_changed_during_package_install_stops_activation(tmp_path, monkeypatch):
    source = tmp_path / "old"
    seeded(source)
    payload, home, data = tmp_path / "payload", tmp_path / "program", tmp_path / "data"
    payload.mkdir()
    setup.atomic_json(payload / "release.json", dict(supported_python=["3.12", "3.11"],
                      architecture="64bit", version="0.1.0", data_format=1, payload_id="test"))
    for name in ("runner.py", "installer.py", "validate.py", "models.py"):
        (payload / name).write_text("# fake package fixture")
    def change_after_copy(args, **kwargs):
        if any(str(value).endswith("validate.py") for value in args):
            store = Store(source / setup.DB)
            store.put("history", dict(id="new-during-install", review_id="review-1", action="new-approved-work"))
            store.close()
    monkeypatch.setattr(setup, "execute", change_after_copy)
    monkeypatch.setattr(setup, "install_packages", lambda *args: None)
    with pytest.raises(ValueError, match="changed during setup"):
        setup.install(payload, home, data, source)
    assert not data.exists()
    assert not (home / "current.json").exists()
    assert any(row[0] == "new-during-install" for row in setup.records(source / setup.DB)["history"])


def package_project(folder):
    folder.mkdir()
    (folder / "requirements.txt").write_text("example==1\n")
    (folder / "pyproject.toml").write_text(
        '[build-system]\nrequires = ["setuptools==84.0.0"]\n'
        '[project]\nversion = "1.0"\ndependencies = ["example==1"]\n'
        '[project.optional-dependencies]\nsemantic = ["semantic-example==2"]\n')
    return folder


def test_dependency_cache_tracks_requirements_python_and_features_not_app_version(tmp_path, monkeypatch):
    project = package_project(tmp_path / "project")
    original = setup.package_spec(project, False)[1]
    metadata = project / "pyproject.toml"
    metadata.write_text(metadata.read_text().replace('version = "1.0"', 'version = "2.0"'))
    assert setup.package_spec(project, False)[1] == original
    assert setup.package_spec(project, True)[1] != original
    (project / "requirements.txt").write_text("example==2\n")
    assert setup.package_spec(project, False)[1] != original
    (project / "requirements.txt").write_text("example==1\n")
    monkeypatch.setattr(setup.sys, "version", "a different Python runtime")
    assert setup.package_spec(project, False)[1] != original


def test_cached_packages_reused_and_corruption_requires_new_download(tmp_path, monkeypatch, capsys):
    project = package_project(tmp_path / "project")
    calls = []
    def fake_execute(args, **kwargs):
        calls.append([str(a) for a in args])
        if "wheel" in args:
            output = Path(args[args.index("--wheel-dir") + 1])
            (output / "example.whl").write_bytes(b"synthetic wheel")
    monkeypatch.setattr(setup, "execute", fake_execute)
    home = tmp_path / "home"
    setup.install_packages("python", project, home, False)
    assert sum("wheel" in call for call in calls) == 1
    calls.clear()
    setup.install_packages("python", project, home, False)
    assert not any("wheel" in call for call in calls)
    assert all("--no-index" in call for call in calls if "install" in call)
    assert "Reusing verified package downloads" in capsys.readouterr().out
    cache = home / "package-cache" / setup.package_spec(project, False)[1]
    (cache / "example.whl").write_bytes(b"damaged download")
    calls.clear()
    setup.install_packages("python", project, home, False)
    assert sum("wheel" in call for call in calls) == 1
    assert setup.verified_wheels(cache)


def test_interrupted_package_download_never_creates_reusable_cache(tmp_path, monkeypatch):
    project = package_project(tmp_path / "project")
    home = tmp_path / "home"
    def interrupted(args, **kwargs):
        output = Path(args[args.index("--wheel-dir") + 1])
        (output / "partial.whl").write_bytes(b"incomplete")
        raise RuntimeError("connection lost")
    monkeypatch.setattr(setup, "execute", interrupted)
    with pytest.raises(RuntimeError, match="connection lost"):
        setup.install_packages("python", project, home, False)
    assert not setup.verified_wheels(home / "package-cache" / setup.package_spec(project, False)[1])


def _touch_db(folder, when):
    folder.mkdir(parents=True, exist_ok=True)
    seeded_path = folder / setup.DB
    os.utime(seeded_path, (when, when))
    return folder.resolve()


def test_discovery_finds_previous_copies_newest_first_and_skips_unrelated_folders(tmp_path):
    documents, downloads, data = tmp_path / "Documents", tmp_path / "Downloads", tmp_path / "AppData" / "ReviewExtractor"
    # A GitHub Desktop clone and a ZIP extracted twice into Downloads.
    clone = documents / "GitHub" / "Systematic" / ".review-data"
    zipped = downloads / "Systematic-main" / "Systematic-main" / ".review-data"
    seeded(clone)
    seeded(zipped)
    _touch_db(clone, 1_700_000_000)
    _touch_db(zipped, 1_800_000_000)
    # Unrelated folders and the installed data are never reported.
    seeded(documents / "Taxes" / "Systematic" / ".review-data")
    seeded(data)
    found = setup.discover([documents, downloads, data.parent], exclude=[data], installed_data=data)
    assert [Path(c["path"]) for c in found] == [zipped.resolve(), clone.resolve()]
    newest = found[0]
    assert newest["usable"] and not newest["already_imported"]
    summary = newest["summary"]
    assert summary["reviews"] == 1 and summary["pdfs"] == 2 and summary["template_columns"] > 0
    assert summary["accepted_evidence"] + summary["rejected_evidence"] > 0
    assert summary["history_entries"] >= 1
    assert "review" in newest["description"] and "template column" in newest["description"]


def test_discovery_survives_unreadable_folders_and_ignores_links(tmp_path, monkeypatch):
    root = tmp_path / "Documents"
    target = root / "Systematic" / ".review-data"
    seeded(target)
    elsewhere = tmp_path / "elsewhere" / "Systematic"
    seeded(elsewhere / ".review-data")
    try:
        (root / "Systematic-link").symlink_to(elsewhere, target_is_directory=True)
    except OSError:
        pass
    blocked = root / "Systematic" / "systematic-locked"
    blocked.mkdir()
    real_is_dir = Path.is_dir

    def is_dir(self):
        if self.name == "systematic-locked":
            raise PermissionError("denied")
        return real_is_dir(self)

    monkeypatch.setattr(Path, "is_dir", is_dir)
    found = setup.discover([root])
    assert [Path(c["path"]) for c in found] == [target.resolve()]


def test_discovery_marks_a_folder_that_was_already_imported(tmp_path):
    source = tmp_path / "Documents" / "Systematic" / ".review-data"
    seeded(source)
    final, work = tmp_path / "appdata", tmp_path / "work"
    work.mkdir()
    stage, _ = setup.prepare_data(final, source, work, final)
    stage.rename(final)
    found = setup.discover([tmp_path / "Documents"], exclude=[final], installed_data=final)
    assert found[0]["already_imported"] is True


def test_transfer_reports_and_proves_everything_carried_over(tmp_path):
    source = tmp_path / "old"
    seeded(source)
    final, work = tmp_path / "appdata", tmp_path / "work"
    work.mkdir()
    expected = setup.carryover_summary(source / setup.DB)
    stage, report = setup.prepare_data(final, source, work, final)
    assert report["carried_over"] == expected
    setup.verify_carryover(source / setup.DB, stage / setup.DB)


@pytest.mark.parametrize("damage", ["missing", "changed"])
def test_carryover_check_rejects_lost_or_altered_records(tmp_path, damage):
    source = tmp_path / "old"
    seeded(source)
    copy = tmp_path / "copy.sqlite3"
    with setup.connect_readonly(source / setup.DB) as src, setup.connect_write(copy) as dst:
        src.backup(dst)
    with setup.connect_write(copy) as db:
        row_id, data = db.execute("SELECT id,data FROM evidence LIMIT 1").fetchone()
        if damage == "missing":
            db.execute("DELETE FROM evidence WHERE id=?", (row_id,))
        else:
            record = json.loads(data)
            record["decision"] = "rejected" if record.get("decision") == "accepted" else "accepted"
            db.execute("UPDATE evidence SET data=? WHERE id=?", (json.dumps(record), row_id))
    with pytest.raises(ValueError, match="evidence"):
        setup.verify_carryover(source / setup.DB, copy)


def test_discover_command_prints_candidates_for_the_setup_window(tmp_path, capsys):
    project = tmp_path / "Desktop" / "Systematic"
    seeded(project / ".review-data")
    roots = tmp_path / "roots.json"
    roots.write_text(json.dumps([str(tmp_path / "Desktop")]), encoding="utf-8")
    code = setup.discover_main(["--discover", "--roots-json", str(roots),
                                "--home", str(tmp_path / "home"), "--data", str(tmp_path / "data")])
    assert code == 0
    line = capsys.readouterr().out.strip()
    assert line.startswith("SYSTEMATIC_DISCOVERY=")
    payload = json.loads(line.split("=", 1)[1])
    assert payload["installed"] is None
    assert Path(payload["candidates"][0]["path"]) == (project / ".review-data").resolve()


def test_remove_old_versions_keeps_active_version_reviews_and_backups(tmp_path):
    home = tmp_path / "Systematic"
    active, old = home / "versions" / "0.2-new", home / "versions" / "0.1-old"
    for folder in (active, old, home / "transactions" / "t1", home / "backups" / "b1" / "data",
                   home / "package-cache" / "good", home / "package-cache" / "x.partial"):
        folder.mkdir(parents=True)
        (folder / "file.bin").write_bytes(b"0" * 1024)
    (active / "runner.py").write_text("")
    result = setup.remove_old_versions(home, {"runner": str(active / "runner.py")})
    assert result["removed"] == 3 and result["freed_bytes"] == 3 * 1024
    assert active.is_dir() and not old.exists() and not (home / "transactions" / "t1").exists()
    assert (home / "backups" / "b1" / "data" / "file.bin").is_file()
    assert (home / "package-cache" / "good").is_dir() and not (home / "package-cache" / "x.partial").exists()


def test_remove_old_versions_skips_the_version_running_setup(tmp_path, monkeypatch):
    home = tmp_path / "Systematic"
    active, running = home / "versions" / "new", home / "versions" / "old"
    for folder in (active, running / ".venv" / "Scripts"):
        folder.mkdir(parents=True)
    (running / ".venv" / "Scripts" / "python.exe").write_bytes(b"")
    monkeypatch.setattr(setup.sys, "executable", str(running / ".venv" / "Scripts" / "python.exe"))
    result = setup.remove_old_versions(home, {"runner": str(active / "runner.py")})
    assert result["removed"] == 0 and running.is_dir() and active.is_dir()


def _cache(folder, rows):
    folder.mkdir(parents=True, exist_ok=True)
    with setup.connect_write(folder / "embeddings.sqlite3") as db:
        db.execute("CREATE TABLE IF NOT EXISTS text_embeddings (key TEXT PRIMARY KEY, vector TEXT NOT NULL)")
        db.executemany("INSERT OR IGNORE INTO text_embeddings VALUES (?,?)", rows)


def _cache_keys(folder):
    with setup.connect_readonly(folder / "embeddings.sqlite3") as db:
        return {r[0] for r in db.execute("SELECT key FROM text_embeddings")}


def test_import_carries_search_cache_into_live_storage_not_archive(tmp_path):
    source = tmp_path / "old"
    seeded(source)
    _cache(source / "retrieval-cache", [("a", "[1]"), ("b", "[2]")])
    final, work = tmp_path / "appdata", tmp_path / "work"
    work.mkdir()
    stage, _ = setup.prepare_data(final, source, work, final)
    assert _cache_keys(stage / "retrieval-cache") == {"a", "b"}
    assert not list((stage / "imported-assets").glob("*/retrieval-cache"))


def test_update_merges_previously_archived_cache_without_losing_new_rows(tmp_path):
    current = tmp_path / "appdata"
    seeded(current)
    _cache(current / "retrieval-cache", [("new", "[9]")])
    _cache(current / "imported-assets" / "abc" / "retrieval-cache", [("old", "[1]"), ("new", "[0]")])
    work = tmp_path / "work"
    work.mkdir()
    stage, _ = setup.prepare_data(current, None, work, current)
    assert _cache_keys(stage / "retrieval-cache") == {"old", "new"}
    with setup.connect_readonly(stage / "retrieval-cache" / "embeddings.sqlite3") as db:
        assert db.execute("SELECT vector FROM text_embeddings WHERE key='new'").fetchone()[0] == "[9]"


TESSERACT = {"version": "5.4.0", "url": "https://example.test/tesseract-setup.exe",
             "sha256": __import__("hashlib").sha256(b"genuine installer").hexdigest()}


def test_existing_tesseract_is_reused_without_downloading(tmp_path, monkeypatch):
    engine = tmp_path / "Tesseract-OCR" / setup.TESSERACT_EXE
    engine.parent.mkdir()
    engine.write_bytes(b"")
    monkeypatch.setattr(setup, "tesseract_works", lambda path: Path(path) == engine)
    monkeypatch.setattr(setup, "download", lambda *a: pytest.fail("must not download"))
    assert setup.ensure_tesseract(tmp_path / "home", TESSERACT, hint=str(engine), windows=True) == str(engine.resolve())


def test_missing_tesseract_is_downloaded_verified_and_installed(tmp_path, monkeypatch):
    home = tmp_path / "home"
    ran = []
    monkeypatch.setattr(setup.shutil, "which", lambda name: None)
    monkeypatch.setattr(setup, "tesseract_works", lambda path: Path(path).is_file())
    monkeypatch.setattr(setup, "download", lambda url, dest: Path(dest).write_bytes(b"genuine installer"))

    def install(installer, target):
        ran.append(installer.read_bytes())
        (target).mkdir(parents=True)
        (target / setup.TESSERACT_EXE).write_bytes(b"")
    monkeypatch.setattr(setup, "run_tesseract_setup", install)
    found = setup.ensure_tesseract(home, TESSERACT, windows=True)
    assert found == str((home / "tesseract" / setup.TESSERACT_EXE).resolve())
    assert ran == [b"genuine installer"]
    assert not list((home / "downloads").iterdir())  # installer removed afterwards


def test_tampered_tesseract_download_never_runs_and_setup_continues(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(setup.shutil, "which", lambda name: None)
    monkeypatch.setattr(setup, "tesseract_works", lambda path: False)
    monkeypatch.setattr(setup, "download", lambda url, dest: Path(dest).write_bytes(b"tampered"))
    monkeypatch.setattr(setup, "run_tesseract_setup", lambda *a: pytest.fail("must not run"))
    assert setup.ensure_tesseract(tmp_path / "home", TESSERACT, windows=True) is None
    assert "checksum" in capsys.readouterr().out
    assert not list((tmp_path / "home" / "downloads").iterdir())


def test_offline_tesseract_install_disables_ocr_without_failing(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(setup.shutil, "which", lambda name: None)
    monkeypatch.setattr(setup, "tesseract_works", lambda path: False)

    def offline(url, dest):
        raise OSError("network unreachable")
    monkeypatch.setattr(setup, "download", offline)
    assert setup.ensure_tesseract(tmp_path / "home", TESSERACT, windows=True) is None
    assert "Systematic works without it" in capsys.readouterr().out


def test_ocr_bridge_is_installed_with_application_packages(tmp_path):
    project = package_project(tmp_path / "project")
    metadata = project / "pyproject.toml"
    metadata.write_text(metadata.read_text() + 'ocr = ["pytesseract>=0.3.13,<0.4"]\n')
    assert "pytesseract>=0.3.13,<0.4" in setup.package_spec(project, False)[0]


def test_runner_enables_ocr_only_for_a_verified_engine(tmp_path, monkeypatch):
    monkeypatch.syspath_prepend(str(Path(setup.__file__).parent))
    from windows import runner
    engine = tmp_path / setup.TESSERACT_EXE
    monkeypatch.delenv("SYSTEMATIC_OCR", raising=False)
    monkeypatch.delenv("SYSTEMATIC_TESSERACT", raising=False)
    runner.enable_ocr({"tesseract": str(engine)})
    assert "SYSTEMATIC_OCR" not in os.environ
    engine.write_bytes(b"")
    runner.enable_ocr({"tesseract": str(engine)})
    assert os.environ["SYSTEMATIC_OCR"] == "1" and os.environ["SYSTEMATIC_TESSERACT"] == str(engine)
    monkeypatch.setenv("SYSTEMATIC_OCR", "0")
    runner.enable_ocr({"tesseract": str(engine)})
    assert os.environ["SYSTEMATIC_OCR"] == "0"


def test_tesseract_falls_back_to_one_admin_install_when_per_user_is_refused(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("SYSTEMATIC_TESSERACT", raising=False)
    monkeypatch.delenv("SYSTEMATIC_TESSERACT", raising=False)
    program_files = tmp_path / "Program Files"
    monkeypatch.setenv("ProgramFiles", str(program_files))
    monkeypatch.setattr(setup.shutil, "which", lambda name: None)
    monkeypatch.setattr(setup, "tesseract_works", lambda path: Path(path).is_file())
    monkeypatch.setattr(setup, "download", lambda url, dest: Path(dest).write_bytes(b"genuine installer"))

    def refused(installer, target):
        raise OSError("[WinError 740] The requested operation requires elevation")
    elevated = []

    def admin(installer, arguments):
        elevated.append(arguments)
        engine = program_files / "Tesseract-OCR" / setup.TESSERACT_EXE
        engine.parent.mkdir(parents=True)
        engine.write_bytes(b"")
        return 0
    monkeypatch.setattr(setup, "run_tesseract_setup", refused)
    monkeypatch.setattr(setup, "run_elevated", admin)
    found = setup.ensure_tesseract(tmp_path / "home", TESSERACT, windows=True)
    assert found == str((program_files / "Tesseract-OCR" / setup.TESSERACT_EXE).resolve())
    assert elevated == ["/S"]
    assert "administrator permission" in capsys.readouterr().out


def test_declined_admin_prompt_disables_ocr_without_failing(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("SYSTEMATIC_TESSERACT", raising=False)
    monkeypatch.delenv("SYSTEMATIC_TESSERACT", raising=False)
    monkeypatch.setenv("ProgramFiles", str(tmp_path / "none"))
    monkeypatch.setattr(setup.shutil, "which", lambda name: None)
    monkeypatch.setattr(setup, "tesseract_works", lambda path: False)
    monkeypatch.setattr(setup, "download", lambda url, dest: Path(dest).write_bytes(b"genuine installer"))
    monkeypatch.setattr(setup, "run_tesseract_setup", lambda *a: None)

    def declined(*args):
        raise PermissionError("administrator permission was declined")
    monkeypatch.setattr(setup, "run_elevated", declined)
    assert setup.ensure_tesseract(tmp_path / "home", TESSERACT, windows=True) is None
    assert "declined" in capsys.readouterr().out
    assert not list((tmp_path / "home" / "downloads").iterdir())


MODEL = {"repo": "org/model", "revision": "main"}
VC = {"url": "https://example.test/vc_redist.x64.exe", "signer": "Microsoft Corporation"}


def fake_model_runs(monkeypatch, outputs):
    calls = []

    def run(args, **kwargs):
        calls.append(args)
        return setup.subprocess.CompletedProcess(args, 0, *outputs.pop(0))
    monkeypatch.setattr(setup.subprocess, "run", run)
    return calls


def test_semantic_model_is_downloaded_and_verified(tmp_path, monkeypatch):
    found = {"repo": "org/model", "revision": "abc123", "cache": str(tmp_path / "models")}
    calls = fake_model_runs(monkeypatch, [("SYSTEMATIC_MODEL=" + json.dumps(found) + "\n", "")])
    assert setup.ensure_semantic_model("python", "models.py", tmp_path, MODEL, windows=False) == found
    assert calls[0][-2:] == ["org/model", "main"]


def test_semantic_model_update_keeps_the_resolved_commit(tmp_path, monkeypatch):
    calls = fake_model_runs(monkeypatch, [("", "offline")])
    setup.ensure_semantic_model("python", "models.py", tmp_path, MODEL,
                                previous={"repo": "org/model", "revision": "abc123"}, windows=False)
    assert calls[0][-1] == "abc123"


def test_missing_cpp_runtime_is_installed_with_admin_then_model_retried(tmp_path, monkeypatch):
    found = {"repo": "org/model", "revision": "abc123", "cache": "c"}
    calls = fake_model_runs(monkeypatch, [("", "OSError: [WinError 126] ... DLL load failed while importing torch"),
                                          ("SYSTEMATIC_MODEL=" + json.dumps(found), "")])
    runtime = []
    monkeypatch.setattr(setup, "ensure_vc_runtime", lambda home, spec: runtime.append(spec))
    assert setup.ensure_semantic_model("python", "models.py", tmp_path, MODEL, vc_spec=VC, windows=True) == found
    assert runtime == [VC] and len(calls) == 2


def test_cpp_runtime_requires_microsoft_signature_before_admin_run(tmp_path, monkeypatch):
    monkeypatch.setattr(setup, "download", lambda url, dest: Path(dest).write_bytes(b"unsigned"))
    monkeypatch.setattr(setup, "signed_by", lambda path, signer: False)
    monkeypatch.setattr(setup, "run_elevated", lambda *a: pytest.fail("must not run"))
    with pytest.raises(ValueError, match="not signed"):
        setup.ensure_vc_runtime(tmp_path, VC)
    assert not list((tmp_path / "downloads").iterdir())


def test_semantic_model_failure_keeps_setup_running(tmp_path, monkeypatch, capsys):
    fake_model_runs(monkeypatch, [("", "ConnectionError: huggingface.co unreachable")])
    assert setup.ensure_semantic_model("python", "models.py", tmp_path, MODEL, windows=False) is None
    assert "fast retrieval" in capsys.readouterr().out


def test_dependency_report_lists_each_feature(capsys):
    result = setup.report_dependencies(setup.sys.executable, None, True,
                                       {"repo": "org/model", "revision": "abc123", "cache": "c"})
    assert result == {"Python": True, "Application packages": True,
                      "OCR for scanned PDFs": False, "Semantic search model": True}
    assert "[MISSING] OCR for scanned PDFs" in capsys.readouterr().out


def test_extras_lock_constrains_installs_and_selects_new_cache(tmp_path, monkeypatch):
    project = package_project(tmp_path / "project")
    before = setup.package_spec(project, False)[1]
    (project / setup.EXTRAS_LOCK).write_text("semantic-example==2\n")
    assert setup.package_spec(project, False)[1] != before
    calls = []
    def fake_execute(args, **kwargs):
        calls.append([str(a) for a in args])
        if "wheel" in args:
            (Path(args[args.index("--wheel-dir") + 1]) / "example.whl").write_bytes(b"wheel")
    monkeypatch.setattr(setup, "execute", fake_execute)
    setup.install_packages("python", project, tmp_path / "home", True)
    pip_calls = [c for c in calls if "wheel" in c or "--find-links" in c]
    assert len(pip_calls) == 2
    assert all(str(project / setup.EXTRAS_LOCK) in c and "-c" in c for c in pip_calls)


def test_runner_uses_only_the_verified_offline_model(tmp_path, monkeypatch):
    monkeypatch.syspath_prepend(str(Path(setup.__file__).parent))
    from windows import runner
    for name in ("HF_HUB_OFFLINE", "HF_HUB_CACHE", "SENTENCE_TRANSFORMERS_HOME",
                 "SYSTEMATIC_EMBEDDING_MODEL", "SYSTEMATIC_EMBEDDING_REVISION"):
        monkeypatch.delenv(name, raising=False)
    model = {"repo": "org/model", "revision": "abc123", "cache": str(tmp_path / "models")}
    runner.enable_semantic({"semantic_model": model})
    assert "HF_HUB_CACHE" not in os.environ  # cache folder missing
    (tmp_path / "models").mkdir()
    runner.enable_semantic({"semantic_model": model})
    assert os.environ["HF_HUB_OFFLINE"] == "1" and os.environ["HF_HUB_CACHE"] == model["cache"]
    assert os.environ["SYSTEMATIC_EMBEDDING_REVISION"] == "abc123"
